"""Warehouse module — counterparties, products, specifications, stock, movements, batches, orders."""
import base64
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal

import csv
import hashlib
import io
import re
import unicodedata
import uuid

import openpyxl
from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, Query, Response, UploadFile, status
from app.api.ws import broadcast_warehouse
from pydantic import BaseModel
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app.api.deps import get_current_org, require_roles, require_warehouse_full
from app.core.db import get_db
from app.models.organization import Organization
from app.models.user import User, UserRole
from app.models.warehouse import (
    AssemblySession,
    BankAccount, BankAccountStatus,
    BatchStatus, CashTransaction, CashTxType, CashTxCategory,
    CellMoveKind, CellMovement, CellStock, Counterparty, LabelTemplate, MovementType, Order, OrderItem, OrderPayment,
    OrderStatus, ProductCategory, ProductImage, ProductionBatch, SpecComponent, SpecOperation,
    SpecOpType, Specification, StockEntry, Warehouse, WarehouseCell, WarehouseMovement,
    WarehouseType, WarehouseZone, Product,
    PurchaseOrder, PurchaseOrderItem, PurchaseOrderStatus,
)
from app.services.label_templates import BUILTIN_TEMPLATES
from app.schemas.warehouse import (
    AssemblySessionOut, AssemblySessionUpdate, BatchAssignRequest,
    WorkerStatsOut, WorkerProductStats,
    BatchClose, BatchComponentOut, BatchCreate, BatchOut, BatchUpdate,
    BankAccountCreate, BankAccountOut, BankAccountUpdate,
    CashFlowSummary, CashRegisterReceipt, CashRegisterReceiptItem, CashRegisterSell, CashTxCreate, CashTxOut,
    CellAssign, CellDetailOut, CellMovementOut, CellNotesUpdate, CellOut, CellStockOut, CellStockSet,
    ScanAction, ScanActionRequest, ScanActionResult, ScanResult,
    CostBreakdown, CounterpartyBalanceAdjust, CounterpartyCreate,
    CounterpartyOut, CounterpartyUpdate,
    DashboardLowStockOut, DashboardSummaryOut,
    MovementCreate, MovementListOut, MovementOut, _MOVEMENT_DIRECTION,
    OrderCreate, OrderItemOut, OrderOut, OrderPaymentCreate, OrderPaymentOut, OrderUpdate,
    ProductCategoryCreate, ProductCategoryOut, ProductCategoryUpdate,
    ProductCreate, ProductImageOut, ProductOptionOut, ProductOut, ProductUpdate,
    ProductCellLocationOut, ProductLocationsOut, ProductWarehouseLocationOut,
    PutawayRequest, RelocateRequest, ReserveRequest,
    ShipPick, ShipRequest,
    SpecComponentCreate, SpecCreate, SpecDefaultSummaryOut, SpecOperationCreate, SpecOut,
    StockEntryOut, StockSummaryOut, UnassignedItemOut, WarehouseCreate, WarehouseOut, WarehouseUpdate,
    ZoneCreate, ZoneOut, ZoneOverviewOut, ZoneUpdate, ZoneWithCellsOut,
    PurchaseOrderCreate, PurchaseOrderItemOut, PurchaseOrderOut,
)

router = APIRouter(prefix="/warehouse", tags=["warehouse"])

# All routes that require Starter plan or above (full warehouse access).
# Free plan can only access /products and /categories.
_full = APIRouter(dependencies=[Depends(require_warehouse_full)])

_ELECTRICITY_RATE  = Decimal("4.5")   # ₴/кВт·год
_LABOR_RATE        = Decimal("150")   # ₴/год
_PRINTER_WATTS     = 200              # Вт
_LOW_STOCK_THRESHOLD = Decimal("10")  # finished-goods units below this surface on the dashboard
_IMAGE_PREFIX      = "product-images"
_IMAGE_MIME_EXT    = {"image/jpeg": "jpg", "image/png": "png", "image/webp": "webp"}
_IMAGE_MAX_BYTES   = 8 * 1024 * 1024  # 8 MB


def _image_url(key: str, product_id: int, org_id: int) -> str:
    from app.services import storage as storage_svc
    url = storage_svc.presigned_url(key, org_id, prefix=_IMAGE_PREFIX, expires=86400)
    return url or f"/api/warehouse/products/{product_id}/images/{key}"


def _product_image_url(p: Product, org_id: int) -> str | None:
    if not p.image_key:
        return None
    return _image_url(p.image_key, p.id, org_id)


def _make_product_out(p: Product, org_id: int) -> ProductOut:
    out = ProductOut.model_validate(p)
    out.image_url = _product_image_url(p, org_id)
    return out


# ── Product categories ────────────────────────────────────────────────────────

@_full.get("/scan", response_model=ScanResult)
def scan(
    q:   str,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> ScanResult:
    """Universal QR/barcode lookup. Accepts CELL:{id}, PROD:{id}, barcode, or SKU."""
    q = q.strip()

    if q.upper().startswith("CELL:"):
        try:
            cell_id = int(q.split(":", 1)[1])
        except ValueError:
            raise HTTPException(status_code=400, detail="Невірний формат CELL")
        cell = (
            db.query(WarehouseCell)
            .join(WarehouseZone, WarehouseZone.id == WarehouseCell.zone_id)
            .filter(WarehouseCell.id == cell_id, WarehouseZone.organization_id == org.id)
            .first()
        )
        if not cell:
            raise HTTPException(status_code=404, detail="Комірку не знайдено")
        zone = db.get(WarehouseZone, cell.zone_id)
        wh   = db.get(Warehouse, zone.warehouse_id)
        stocks = db.query(CellStock).filter_by(cell_id=cell.id).all()
        stock_out = []
        for cs in stocks:
            p = db.get(Product, cs.product_id)
            if p:
                stock_out.append(CellStockOut(
                    product_id=cs.product_id, product_name=p.name,
                    product_sku=p.sku, product_unit=p.unit, quantity=cs.quantity,
                    image_url=_product_image_url(p, org.id),
                ))
        detail = CellDetailOut(
            cell_id=cell.id, cell_code=cell.code, cell_notes=cell.notes,
            zone_id=zone.id, zone_name=zone.name,
            warehouse_id=wh.id, warehouse_name=wh.name,
            stock=stock_out,
        )
        return ScanResult(type="cell", cell=detail)

    # PROD:{id} or barcode or sku
    pid: int | None = None
    if q.upper().startswith("PROD:"):
        try:
            pid = int(q.split(":", 1)[1])
        except ValueError:
            pass
    p = None
    if pid:
        p = db.query(Product).filter_by(id=pid, organization_id=org.id).first()
    if not p:
        p = db.query(Product).filter_by(barcode=q, organization_id=org.id).first()
    if not p:
        p = db.query(Product).filter_by(sku=q, organization_id=org.id).first()
    if not p:
        raise HTTPException(status_code=404, detail="Не знайдено")
    return ScanResult(type="product", product=_make_product_out(p, org.id))


@router.get("/categories", response_model=list[ProductCategoryOut])
def list_categories(
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[ProductCategoryOut]:
    rows = (
        db.query(ProductCategory)
        .filter(ProductCategory.organization_id == org.id)
        .order_by(ProductCategory.sort_order, ProductCategory.name)
        .all()
    )
    return [ProductCategoryOut.model_validate(r) for r in rows]


@router.post("/categories", response_model=ProductCategoryOut, status_code=status.HTTP_201_CREATED)
def create_category(
    payload: ProductCategoryCreate,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    _:    User         = Depends(require_roles(UserRole.admin)),
) -> ProductCategoryOut:
    cat = ProductCategory(**payload.model_dump(), organization_id=org.id)
    db.add(cat)
    db.commit()
    db.refresh(cat)
    return ProductCategoryOut.model_validate(cat)


@router.patch("/categories/{cat_id}", response_model=ProductCategoryOut)
def update_category(
    cat_id:  int,
    payload: ProductCategoryUpdate,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    _:    User         = Depends(require_roles(UserRole.admin)),
) -> ProductCategoryOut:
    cat = db.query(ProductCategory).filter_by(id=cat_id, organization_id=org.id).first()
    if not cat:
        raise HTTPException(status_code=404, detail="Category not found")
    for k, v in payload.model_dump(exclude_unset=True).items():
        setattr(cat, k, v)
    db.commit()
    db.refresh(cat)
    return ProductCategoryOut.model_validate(cat)


@router.delete("/categories/{cat_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_category(
    cat_id: int,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    _:    User         = Depends(require_roles(UserRole.admin)),
) -> None:
    cat = db.query(ProductCategory).filter_by(id=cat_id, organization_id=org.id).first()
    if not cat:
        raise HTTPException(status_code=404, detail="Category not found")
    db.delete(cat)
    db.commit()


# ── Helpers ───────────────────────────────────────────────────────────────────

def _get_product(product_id: int, org: Organization, db: Session) -> Product:
    p = db.query(Product).filter(Product.id == product_id, Product.organization_id == org.id).first()
    if not p:
        raise HTTPException(status_code=404, detail="Product not found")
    return p


def _get_warehouse(wh_id: int, org: Organization, db: Session) -> Warehouse:
    wh = db.query(Warehouse).filter(Warehouse.id == wh_id, Warehouse.organization_id == org.id).first()
    if not wh:
        raise HTTPException(status_code=404, detail="Warehouse not found")
    return wh


def _get_counterparty(cp_id: int, org: Organization, db: Session) -> Counterparty:
    cp = db.query(Counterparty).filter(Counterparty.id == cp_id, Counterparty.organization_id == org.id).first()
    if not cp:
        raise HTTPException(status_code=404, detail="Counterparty not found")
    return cp


def _get_spec(spec_id: int, org: Organization, db: Session) -> Specification:
    spec = (
        db.query(Specification)
        .join(Product, Product.id == Specification.product_id)
        .filter(Specification.id == spec_id, Product.organization_id == org.id)
        .first()
    )
    if not spec:
        raise HTTPException(status_code=404, detail="Specification not found")
    return spec


def _get_spec_for_product(
    spec_id: int | None,
    product_id: int,
    org: Organization,
    db: Session,
) -> Specification | None:
    if spec_id is None:
        return None
    spec = _get_spec(spec_id, org, db)
    if spec.product_id != product_id:
        raise HTTPException(status_code=400, detail="Specification does not belong to product")
    return spec


def _norm_lookup(value: str | None) -> str:
    return " ".join((value or "").casefold().split())


def _component_sku_from_name(name: str) -> str:
    ascii_name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^A-Za-z0-9]+", "-", ascii_name).strip("-").upper()
    if len(slug) < 3:
        slug = "COMP"
    digest = hashlib.sha1(name.casefold().encode("utf-8")).hexdigest()[:8].upper()
    return f"COMP-{slug[:24]}-{digest}"


def _unique_component_sku(
    db: Session,
    org_id: int,
    name: str,
    preferred_sku: str | None = None,
) -> str:
    base = (preferred_sku or "").strip() or _component_sku_from_name(name)
    base = base[:80] or _component_sku_from_name(name)
    candidate = base
    suffix = 2
    while db.query(Product.id).filter(Product.organization_id == org_id, Product.sku == candidate).first():
        tail = f"-{suffix}"
        candidate = f"{base[:80 - len(tail)]}{tail}"
        suffix += 1
    return candidate


def _ensure_component_product(
    *,
    db: Session,
    org: Organization,
    user: User | None,
    name: str,
    sku: str | None,
    unit: str | None,
    unit_price: Decimal | None,
    product_by_sku: dict[str, Product] | None = None,
    product_by_name: dict[str, Product] | None = None,
) -> Product:
    clean_name = name.strip()
    if not clean_name:
        raise HTTPException(status_code=400, detail="Component name is required")

    product = product_by_sku.get(_norm_lookup(sku)) if product_by_sku and sku else None
    if product is None and product_by_name is not None:
        product = product_by_name.get(_norm_lookup(clean_name))
    if product is None and product_by_sku is None and sku:
        product = db.query(Product).filter(Product.organization_id == org.id, Product.sku == sku.strip()).first()
    if product is None and product_by_name is None:
        product = next(
            (
                existing
                for existing in db.query(Product).filter(Product.organization_id == org.id).all()
                if _norm_lookup(existing.name) == _norm_lookup(clean_name)
            ),
            None,
        )
    if product is not None:
        return product

    _check_product_limit(org, db)
    product = Product(
        organization_id=org.id,
        created_by_id=user.id if user else None,
        name=clean_name,
        sku=_unique_component_sku(db, org.id, clean_name, preferred_sku=sku),
        categories=[],
        unit=(unit or "").strip() or "шт",
        cost_price=unit_price,
        is_active=True,
    )
    db.add(product)
    db.flush()
    if product_by_sku is not None:
        product_by_sku[_norm_lookup(product.sku)] = product
    if product_by_name is not None:
        product_by_name[_norm_lookup(product.name)] = product
    return product


def _component_payloads(
    components: list[SpecComponent],
    db: Session,
    org_id: int,
    product_names: dict[int, str] | None = None,
) -> list[dict]:
    if product_names is None:
        product_ids = {c.product_id for c in components if c.product_id is not None}
        product_names = {}
        if product_ids:
            product_names = {
                product_id: name
                for product_id, name in (
                    db.query(Product.id, Product.name)
                    .filter(Product.organization_id == org_id, Product.id.in_(product_ids))
                    .all()
                )
            }
    return [
        {
            "id": c.id,
            "name": c.name,
            "product_id": c.product_id,
            "product_name": product_names.get(c.product_id),
            "material_id": c.material_id,
            "quantity": c.quantity,
            "unit": c.unit,
            "unit_price": c.unit_price,
            "waste_pct": c.waste_pct,
            "sort_order": c.sort_order,
        }
        for c in components
    ]


def _spec_to_out(spec: Specification, db: Session, org_id: int) -> SpecOut:
    components = db.query(SpecComponent).filter(SpecComponent.specification_id == spec.id).order_by(SpecComponent.sort_order).all()
    operations = db.query(SpecOperation).filter(SpecOperation.specification_id == spec.id).order_by(SpecOperation.sort_order).all()
    return SpecOut.model_validate({
        **spec.__dict__,
        "components": _component_payloads(components, db, org_id),
        "operations": operations,
    })


def _specs_to_out(specs: list[Specification], db: Session, org_id: int) -> list[SpecOut]:
    spec_ids = [s.id for s in specs]
    if not spec_ids:
        return []

    components_by_spec: dict[int, list[SpecComponent]] = {sid: [] for sid in spec_ids}
    operations_by_spec: dict[int, list[SpecOperation]] = {sid: [] for sid in spec_ids}

    components = (
        db.query(SpecComponent)
        .filter(SpecComponent.specification_id.in_(spec_ids))
        .order_by(SpecComponent.specification_id, SpecComponent.sort_order)
        .all()
    )
    for component in components:
        components_by_spec.setdefault(component.specification_id, []).append(component)

    component_product_ids = {c.product_id for c in components if c.product_id is not None}
    component_product_names: dict[int, str] = {}
    if component_product_ids:
        component_product_names = {
            product_id: name
            for product_id, name in (
                db.query(Product.id, Product.name)
                .filter(Product.organization_id == org_id, Product.id.in_(component_product_ids))
                .all()
            )
        }

    operations = (
        db.query(SpecOperation)
        .filter(SpecOperation.specification_id.in_(spec_ids))
        .order_by(SpecOperation.specification_id, SpecOperation.sort_order)
        .all()
    )
    for operation in operations:
        operations_by_spec.setdefault(operation.specification_id, []).append(operation)

    return [
        SpecOut.model_validate({
            **spec.__dict__,
            "components": _component_payloads(
                components_by_spec.get(spec.id, []),
                db,
                org_id,
                product_names=component_product_names,
            ),
            "operations": operations_by_spec.get(spec.id, []),
        })
        for spec in specs
    ]


def _calc_cost(
    spec: Specification,
    db: Session,
    electricity_rate: Decimal = _ELECTRICITY_RATE,
    labor_rate: Decimal = _LABOR_RATE,
) -> CostBreakdown:
    components = db.query(SpecComponent).filter(SpecComponent.specification_id == spec.id).all()
    operations = db.query(SpecOperation).filter(SpecOperation.specification_id == spec.id).all()

    material_cost = Decimal("0")
    for c in components:
        if c.unit_price is not None:
            waste = Decimal("1") + (c.waste_pct or Decimal("0")) / 100
            material_cost += c.quantity * c.unit_price * waste

    electricity_cost = Decimal("0")
    labor_cost       = Decimal("0")
    other_cost       = Decimal("0")
    print_time_min   = Decimal("0")

    for op in operations:
        if op.type.value == "print" and op.print_time_min:
            hours = op.print_time_min / 60
            kwh   = Decimal(op.power_watts or _PRINTER_WATTS) * hours / 1000
            electricity_cost += kwh * electricity_rate
            print_time_min   += op.print_time_min
        if op.labor_minutes:
            rate       = op.labor_rate_per_hour or labor_rate
            labor_cost += (op.labor_minutes / 60) * rate
        if op.explicit_cost:
            other_cost += op.explicit_cost

    total = material_cost + electricity_cost + labor_cost + other_cost

    product = db.get(Product, spec.product_id)
    margin = None
    if product and product.sale_price and total > 0:
        margin = ((product.sale_price - total) / product.sale_price * 100).quantize(Decimal("0.01"))

    return CostBreakdown(
        material_cost=material_cost.quantize(Decimal("0.0001")),
        electricity_cost=electricity_cost.quantize(Decimal("0.0001")),
        labor_cost=labor_cost.quantize(Decimal("0.0001")),
        other_cost=other_cost.quantize(Decimal("0.0001")),
        total=total.quantize(Decimal("0.0001")),
        print_time_min=print_time_min,
        margin_pct=margin,
    )


def _next_order_number(org: Organization, db: Session) -> str:
    count = db.query(Order).filter(Order.organization_id == org.id).count()
    return f"#ORD-{count + 1:04d}"


def _update_avco(product_id: int, incoming_qty: Decimal, incoming_cost: Decimal, db: Session) -> None:
    """Recalculate product.cost_price using Average Cost (AVCO) method.

    Must be called BEFORE the new stock is added to StockEntry so that
    current_qty reflects the quantity already on hand.
    """
    product = db.query(Product).with_for_update().filter(Product.id == product_id).first()
    if product is None:
        return

    current_qty = db.query(func.sum(StockEntry.quantity)).filter(
        StockEntry.organization_id == product.organization_id,
        StockEntry.product_id == product_id,
    ).scalar() or Decimal("0")

    current_cost = product.cost_price or Decimal("0")

    if current_qty > 0 and current_cost > 0:
        new_cost = (current_qty * current_cost + incoming_qty * incoming_cost) / (current_qty + incoming_qty)
    else:
        new_cost = incoming_cost

    product.cost_price = new_cost.quantize(Decimal("0.0001"))


# ── Bin (cell) reconciliation ───────────────────────────────────────────────
#
# StockEntry stays the source of truth for totals. Cells are a physical
# allocation underneath it, bounded by the invariant:
#
#     for each (product, warehouse):  sum(CellStock) <= StockEntry.quantity
#
# The difference is "unassigned" (floor) stock. Every stock decrease runs
# through _clamp_cells_to_stock, so cell allocations can never drift above the
# real total — even for callers (orders, batches, auto-replenish) that know
# nothing about cells.

def _get_cell(cell_id: int, org: Organization, db: Session, warehouse_id: int | None = None) -> WarehouseCell:
    cell = (
        db.query(WarehouseCell)
        .join(WarehouseZone, WarehouseZone.id == WarehouseCell.zone_id)
        .filter(WarehouseCell.id == cell_id, WarehouseZone.organization_id == org.id)
        .first()
    )
    if not cell:
        raise HTTPException(status_code=404, detail="Cell not found")
    if warehouse_id is not None:
        zone = db.get(WarehouseZone, cell.zone_id)
        if not zone or zone.warehouse_id != warehouse_id:
            raise HTTPException(status_code=400, detail="Комірка не належить цьому складу")
    return cell


def _cell_warehouse_id(cell: WarehouseCell, db: Session) -> int:
    zone = db.get(WarehouseZone, cell.zone_id)
    return zone.warehouse_id  # type: ignore[union-attr]


def _cells_in_warehouse(product_id: int, warehouse_id: int, db: Session) -> list[CellStock]:
    """CellStock rows for a product in one warehouse, oldest first (FIFO)."""
    return (
        db.query(CellStock)
        .join(WarehouseCell, WarehouseCell.id == CellStock.cell_id)
        .join(WarehouseZone, WarehouseZone.id == WarehouseCell.zone_id)
        .filter(
            WarehouseZone.warehouse_id == warehouse_id,
            CellStock.product_id == product_id,
            CellStock.quantity > 0,
        )
        .order_by(CellStock.updated_at, CellStock.id)
        .all()
    )


def _cell_assigned(product_id: int, warehouse_id: int, db: Session) -> Decimal:
    return sum((cs.quantity for cs in _cells_in_warehouse(product_id, warehouse_id, db)), Decimal("0"))


def _unassigned_qty(product_id: int, warehouse_id: int, org_id: int, db: Session) -> Decimal:
    entry = db.query(StockEntry).filter_by(
        organization_id=org_id,
        product_id=product_id,
        warehouse_id=warehouse_id,
    ).first()
    total = entry.quantity if entry else Decimal("0")
    return max(Decimal("0"), total - _cell_assigned(product_id, warehouse_id, db))


def _lock_stock_row(product_id: int, warehouse_id: int, org_id: int, db: Session) -> None:
    """Row-lock the (product, warehouse) stock entry so concurrent bin edits on
    the same product+warehouse serialize (prevents over-assigning the pool)."""
    db.query(StockEntry).filter_by(
        organization_id=org_id, product_id=product_id, warehouse_id=warehouse_id
    ).with_for_update().first()


def _log_cell_move(
    db: Session, *, org_id: int, product_id: int, quantity: Decimal, kind: CellMoveKind,
    cell_from_id: int | None = None, cell_to_id: int | None = None,
    movement_id: int | None = None, created_by_id: int | None = None,
) -> None:
    if quantity <= 0:
        return
    db.add(CellMovement(
        organization_id=org_id, product_id=product_id, quantity=quantity, kind=kind,
        cell_from_id=cell_from_id, cell_to_id=cell_to_id,
        movement_id=movement_id, created_by_id=created_by_id,
    ))


def _putaway(
    cell: WarehouseCell, product_id: int, qty: Decimal, org_id: int, db: Session,
    *, kind: CellMoveKind = CellMoveKind.putaway, movement_id: int | None = None, created_by_id: int | None = None,
) -> None:
    """Assign qty from the unassigned pool into a cell. Caller must cap qty to availability."""
    if qty <= 0:
        return
    cs = db.query(CellStock).filter_by(cell_id=cell.id, product_id=product_id).first()
    if cs:
        cs.quantity += qty
    else:
        db.add(CellStock(cell_id=cell.id, product_id=product_id, quantity=qty))
    _log_cell_move(db, org_id=org_id, product_id=product_id, quantity=qty, kind=kind,
                   cell_to_id=cell.id, movement_id=movement_id, created_by_id=created_by_id)


def _pick_from_cell(
    cell: WarehouseCell, product_id: int, qty: Decimal, org_id: int, db: Session,
    *, kind: CellMoveKind = CellMoveKind.pick, movement_id: int | None = None,
    created_by_id: int | None = None, keep_row: bool = False,
) -> Decimal:
    """Remove up to qty from a cell. Returns the amount actually removed.

    keep_row=True preserves the CellStock row at 0 (write-off: goods destroyed in place).
    """
    cs = db.query(CellStock).filter_by(cell_id=cell.id, product_id=product_id).first()
    take = min(cs.quantity, qty) if cs else Decimal("0")
    if take > 0:
        cs.quantity -= take
        _log_cell_move(db, org_id=org_id, product_id=product_id, quantity=take, kind=kind,
                       cell_from_id=cell.id, movement_id=movement_id, created_by_id=created_by_id)
        if cs.quantity <= 0 and not keep_row:
            db.delete(cs)
    return take


def _clamp_cells_to_stock(
    product_id: int, warehouse_id: int, org_id: int, db: Session,
    *, movement_id: int | None = None, created_by_id: int | None = None,
    keep_rows: bool = False,
) -> None:
    """Reduce cell allocations FIFO until sum(cells) <= StockEntry.quantity."""
    entry = db.query(StockEntry).filter_by(
        organization_id=org_id,
        product_id=product_id,
        warehouse_id=warehouse_id,
    ).first()
    total = entry.quantity if entry else Decimal("0")
    cells = _cells_in_warehouse(product_id, warehouse_id, db)
    excess = sum((c.quantity for c in cells), Decimal("0")) - total
    if excess <= 0:
        return
    for cs in cells:
        if excess <= 0:
            break
        cell = db.get(WarehouseCell, cs.cell_id)
        take = _pick_from_cell(cell, product_id, min(cs.quantity, excess), org_id, db,
                               movement_id=movement_id, created_by_id=created_by_id,
                               keep_row=keep_rows)
        excess -= take


def _apply_movement(movement: WarehouseMovement, db: Session) -> None:
    """Update StockEntry rows to reflect a committed movement.

    For PURCHASE_IN, updates AVCO on the product BEFORE adding stock so
    the formula uses the quantity currently on hand. After applying, clamps
    cell allocations in any touched warehouse back within the new totals.
    """
    q   = movement.quantity
    mt  = movement.type
    pid = movement.product_id

    def _entry(product_id: int, warehouse_id: int) -> StockEntry:
        from decimal import Decimal
        row = db.query(StockEntry).filter_by(
            organization_id=movement.organization_id,
            product_id=product_id,
            warehouse_id=warehouse_id,
        ).first()
        if not row:
            row = StockEntry(
                organization_id=movement.organization_id,
                product_id=product_id,
                warehouse_id=warehouse_id,
                quantity=Decimal("0"),
                reserved_qty=Decimal("0"),
            )
            db.add(row)
        return row

    def _decrease(product_id: int, warehouse_id: int, amount: Decimal) -> StockEntry:
        row = _entry(product_id, warehouse_id)
        if row.quantity < amount:
            raise HTTPException(status_code=422, detail="Недостатньо товару на складі")
        row.quantity -= amount
        return row

    if mt == MovementType.PURCHASE_IN and movement.warehouse_to_id:
        if movement.unit_cost:
            _update_avco(pid, q, movement.unit_cost, db)
        _entry(pid, movement.warehouse_to_id).quantity += q

    elif mt in (MovementType.PRODUCTION_IN, MovementType.RETURN_IN) and movement.warehouse_to_id:
        if mt == MovementType.PRODUCTION_IN and movement.unit_cost:
            _update_avco(pid, q, movement.unit_cost, db)
        _entry(pid, movement.warehouse_to_id).quantity += q

    elif mt == MovementType.DEFECT:
        if movement.warehouse_from_id:
            _decrease(pid, movement.warehouse_from_id, q)
        if movement.warehouse_to_id:
            _entry(pid, movement.warehouse_to_id).quantity += q

    elif mt in (MovementType.SALE_OUT, MovementType.PRODUCTION_OUT, MovementType.WRITE_OFF) and movement.warehouse_from_id:
        row = _decrease(pid, movement.warehouse_from_id, q)
        if mt == MovementType.SALE_OUT and movement.order_id:
            row.reserved_qty -= min(row.reserved_qty, q)

    elif mt == MovementType.ADJUSTMENT:
        if movement.warehouse_to_id:
            _entry(pid, movement.warehouse_to_id).quantity += q
        elif movement.warehouse_from_id:
            _decrease(pid, movement.warehouse_from_id, q)

    elif mt == MovementType.TRANSFER and movement.warehouse_from_id and movement.warehouse_to_id:
        _decrease(pid, movement.warehouse_from_id, q)
        _entry(pid, movement.warehouse_to_id).quantity   += q

    # Keep cell allocations within the new totals for every touched warehouse.
    # Clamping a warehouse whose stock only increased is a harmless no-op.
    # For WRITE_OFF, goods are destroyed in place — preserve CellStock rows at 0
    # so the cell assignment remains visible (keep_rows=True skips db.delete).
    db.flush()
    keep = mt == MovementType.WRITE_OFF
    for wh_id in {movement.warehouse_from_id, movement.warehouse_to_id}:
        if wh_id:
            _clamp_cells_to_stock(pid, wh_id, movement.organization_id, db,
                                  movement_id=movement.id, created_by_id=movement.created_by_id,
                                  keep_rows=keep)


def _check_and_auto_replenish(pid: int, org_id: int, db: Session) -> None:
    from sqlalchemy import func
    from app.models.warehouse import Product, StockEntry, ProductionBatch, Specification, BatchStatus

    p = db.query(Product).filter(Product.id == pid, Product.organization_id == org_id).first()
    if not p or p.min_stock is None:
        return

    total_qty = db.query(func.sum(StockEntry.quantity)).filter_by(organization_id=org_id, product_id=pid).scalar() or 0
    total_res = db.query(func.sum(StockEntry.reserved_qty)).filter_by(organization_id=org_id, product_id=pid).scalar() or 0
    available = float(total_qty - total_res)

    if available >= p.min_stock:
        return

    active_batch = db.query(ProductionBatch).filter(
        ProductionBatch.organization_id == org_id,
        ProductionBatch.product_id == pid,
        ProductionBatch.status.in_([BatchStatus.draft, BatchStatus.active])
    ).first()

    if active_batch:
        return

    target_qty = (p.desired_stock - available) if p.desired_stock is not None else (p.min_stock - available)
    if target_qty <= 0:
        target_qty = 10

    target_qty = int(target_qty)

    spec = (
        db.query(Specification)
        .join(Product, Product.id == Specification.product_id)
        .filter(
            Product.organization_id == org_id,
            Specification.product_id == pid,
            Specification.is_default.is_(True),
        )
        .first()
    )

    batch = ProductionBatch(
        organization_id=org_id,
        product_id=pid,
        specification_id=spec.id if spec else None,
        target_qty=target_qty,
        status=BatchStatus.draft,
        notes="Автоматичне поповнення запасів"
    )
    db.add(batch)


def _order_to_out(
    o: Order,
    db: Session,
    *,
    items: list[OrderItem] | None = None,
    product_map: dict[int, Product] | None = None,
    cp_map: dict[int, Counterparty] | None = None,
) -> OrderOut:
    if items is None:
        items = db.query(OrderItem).filter_by(order_id=o.id).all()
    item_outs = []
    for it in items:
        p = product_map.get(it.product_id) if product_map is not None else db.get(Product, it.product_id)
        item_outs.append(OrderItemOut(
            id=it.id,
            product_id=it.product_id,
            product_name=p.name if p else "",  # type: ignore[union-attr]
            product_sku=p.sku if p else None,
            image_url=_product_image_url(p, o.organization_id) if p else None,
            warehouse_id=it.warehouse_id,
            quantity=it.quantity,
            unit_price=it.unit_price,
            total_price=it.total_price,
        ))

    counterparty_name: str | None = None
    customer_phone: str | None = None
    customer_email: str | None = None
    delivery_address: str | None = None
    if o.counterparty_id:
        cp = cp_map.get(o.counterparty_id) if cp_map is not None else db.get(Counterparty, o.counterparty_id)
        if cp:
            counterparty_name = cp.name
            customer_phone = cp.phone
            customer_email = cp.email
            delivery_address = cp.address

    payload = o.external_payload or {}
    delivery_type = payload.get("delivery_type") if isinstance(payload, dict) else None
    payment_type = payload.get("payment_type") if isinstance(payload, dict) else None
    delivery_service = delivery_type.get("title") if isinstance(delivery_type, dict) else None
    payment_method = payment_type.get("title") if isinstance(payment_type, dict) else None
    delivery_city = None
    if isinstance(payload, dict):
        customer_phone = payload.get("delivery_phone") or customer_phone
        customer_email = payload.get("delivery_email") or customer_email
        delivery_city = payload.get("delivery_city_stable") or payload.get("delivery_city")
        delivery_address = payload.get("delivery_address") or delivery_address

    total       = o.total_amount or Decimal("0")
    paid        = o.paid_amount  or Decimal("0")
    outstanding = max(total - paid, Decimal("0"))

    if paid <= 0:
        payment_status = "unpaid"
    elif total > 0 and paid >= total:
        payment_status = "paid"
    else:
        payment_status = "partial"

    return OrderOut(
        id=o.id,
        order_number=o.order_number,
        counterparty_id=o.counterparty_id,
        counterparty_name=counterparty_name,
        customer_name=o.customer_name,
        customer_phone=customer_phone,
        customer_email=customer_email,
        delivery_city=delivery_city,
        delivery_service=delivery_service,
        delivery_address=delivery_address,
        payment_method=payment_method,
        source=o.source,
        status=o.status,
        total_amount=o.total_amount,
        paid_amount=paid,
        outstanding=outstanding,
        payment_status=payment_status,
        currency=o.currency,
        due_date=o.due_date,
        notes=o.notes,
        items=item_outs,
        created_at=o.created_at,
    )


# ── Warehouses ────────────────────────────────────────────────────────────────

@_full.get("/warehouses", response_model=list[WarehouseOut])
def list_warehouses(db: Session = Depends(get_db), org: Organization = Depends(get_current_org)) -> list[WarehouseOut]:
    rows = db.query(Warehouse).filter(Warehouse.organization_id == org.id).order_by(Warehouse.name).all()
    return [WarehouseOut.model_validate(r) for r in rows]


@_full.post("/warehouses", response_model=WarehouseOut, status_code=status.HTTP_201_CREATED)
def create_warehouse(
    payload: WarehouseCreate,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> WarehouseOut:
    wh = Warehouse(**payload.model_dump(), organization_id=org.id)
    db.add(wh)
    db.commit()
    db.refresh(wh)
    return WarehouseOut.model_validate(wh)


@_full.patch("/warehouses/{wh_id}", response_model=WarehouseOut)
def update_warehouse(
    wh_id:   int,
    payload: WarehouseUpdate,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> WarehouseOut:
    wh = _get_warehouse(wh_id, org, db)
    for k, v in payload.model_dump(exclude_unset=True).items():
        setattr(wh, k, v)
    db.commit()
    db.refresh(wh)
    return WarehouseOut.model_validate(wh)


@_full.delete("/warehouses/{wh_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_warehouse(
    wh_id: int,
    force: bool       = False,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> None:
    wh = _get_warehouse(wh_id, org, db)
    has_stock = db.query(func.count(StockEntry.id)).filter(
        StockEntry.warehouse_id == wh_id,
        StockEntry.quantity > 0,
    ).scalar() or 0
    if has_stock:
        raise HTTPException(
            status_code=400,
            detail=f"Не можна видалити склад — {has_stock} позицій мають ненульовий залишок. Спочатку перемістіть або спишіть товари.",
        )
    has_movements = db.query(func.count(WarehouseMovement.id)).filter(
        (WarehouseMovement.warehouse_from_id == wh_id) | (WarehouseMovement.warehouse_to_id == wh_id)
    ).scalar() or 0
    if has_movements and not force:
        raise HTTPException(
            status_code=400,
            detail=f"Є {has_movements} рухів пов'язаних з цим складом. Щоб видалити разом з історією, підтвердьте примусове видалення.",
        )
    if force:
        # nullify movement references so FK won't block deletion
        db.query(WarehouseMovement).filter(WarehouseMovement.warehouse_from_id == wh_id).update({"warehouse_from_id": None})
        db.query(WarehouseMovement).filter(WarehouseMovement.warehouse_to_id   == wh_id).update({"warehouse_to_id":   None})
        db.query(StockEntry).filter(StockEntry.warehouse_id == wh_id).delete()
    db.delete(wh)
    db.commit()


# ── Zones ─────────────────────────────────────────────────────────────────────

def _col_letter(n: int) -> str:
    """0-based column index → Excel-style letter(s): 0→A, 25→Z, 26→AA …"""
    result = ""
    n += 1
    while n:
        n, r = divmod(n - 1, 26)
        result = chr(65 + r) + result
    return result


def _generate_cells(zone: WarehouseZone) -> list[WarehouseCell]:
    """Generate cell codes as row-letter + column-number: row A → A1, A2, …;
    row B → B1, B2, …  Appended row-major so cell id order matches the grid."""
    cells = []
    for r in range(zone.rows):
        for c in range(zone.cols):
            code = f"{_col_letter(r)}{c + 1}"
            cells.append(WarehouseCell(zone_id=zone.id, code=code))
    return cells


def _zone_out(zone: WarehouseZone, db: Session) -> ZoneOut:
    count = db.query(WarehouseCell).filter(WarehouseCell.zone_id == zone.id).count()
    return ZoneOut(
        id=zone.id, name=zone.name, rows=zone.rows, cols=zone.cols,
        sort_order=zone.sort_order, cell_count=count, created_at=zone.created_at,
    )


@_full.get("/zones", response_model=list[ZoneOverviewOut])
def list_all_zones(
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[ZoneOverviewOut]:
    """Flat overview of every zone across all warehouses (for the Стелажі page)."""
    rows = (
        db.query(WarehouseZone, Warehouse)
        .join(Warehouse, Warehouse.id == WarehouseZone.warehouse_id)
        .filter(WarehouseZone.organization_id == org.id)
        .order_by(Warehouse.name, WarehouseZone.sort_order, WarehouseZone.id)
        .all()
    )
    # One grouped query each for total and filled cell counts (avoids N+1).
    counts = dict(
        db.query(WarehouseCell.zone_id, func.count(WarehouseCell.id))
        .group_by(WarehouseCell.zone_id).all()
    )
    filled = dict(
        db.query(WarehouseCell.zone_id, func.count(func.distinct(WarehouseCell.id)))
        .join(CellStock, CellStock.cell_id == WarehouseCell.id)
        .filter(CellStock.quantity > 0)
        .group_by(WarehouseCell.zone_id).all()
    )
    return [
        ZoneOverviewOut(
            id=z.id, name=z.name, rows=z.rows, cols=z.cols, sort_order=z.sort_order,
            cell_count=counts.get(z.id, 0), created_at=z.created_at,
            warehouse_id=wh.id, warehouse_name=wh.name, filled_cells=filled.get(z.id, 0),
        )
        for z, wh in rows
    ]


@_full.get("/warehouses/{wh_id}/zones", response_model=list[ZoneOut])
def list_zones(
    wh_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[ZoneOut]:
    _get_warehouse(wh_id, org, db)
    zones = (db.query(WarehouseZone)
               .filter(WarehouseZone.warehouse_id == wh_id, WarehouseZone.organization_id == org.id)
               .order_by(WarehouseZone.sort_order, WarehouseZone.id)
               .all())
    zone_ids = [z.id for z in zones]
    counts = dict(
        db.query(WarehouseCell.zone_id, func.count(WarehouseCell.id))
        .filter(WarehouseCell.zone_id.in_(zone_ids))
        .group_by(WarehouseCell.zone_id)
        .all()
    ) if zone_ids else {}
    return [
        ZoneOut(
            id=z.id,
            name=z.name,
            rows=z.rows,
            cols=z.cols,
            sort_order=z.sort_order,
            cell_count=counts.get(z.id, 0),
            created_at=z.created_at,
        )
        for z in zones
    ]


@_full.post("/warehouses/{wh_id}/zones", response_model=ZoneOut, status_code=status.HTTP_201_CREATED)
def create_zone(
    wh_id:   int,
    payload: ZoneCreate,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> ZoneOut:
    _get_warehouse(wh_id, org, db)
    zone = WarehouseZone(**payload.model_dump(), warehouse_id=wh_id, organization_id=org.id)
    db.add(zone)
    db.flush()
    for cell in _generate_cells(zone):
        db.add(cell)
    db.commit()
    db.refresh(zone)
    return _zone_out(zone, db)


@_full.patch("/warehouses/{wh_id}/zones/{zone_id}", response_model=ZoneOut)
def update_zone(
    wh_id:   int,
    zone_id: int,
    payload: ZoneUpdate,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> ZoneOut:
    _get_warehouse(wh_id, org, db)
    zone = db.query(WarehouseZone).filter(
        WarehouseZone.id == zone_id, WarehouseZone.warehouse_id == wh_id,
        WarehouseZone.organization_id == org.id,
    ).first()
    if not zone:
        raise HTTPException(status_code=404, detail="Zone not found")
    data = payload.model_dump(exclude_unset=True)
    # Only a real change in dimensions counts as a resize; renaming a zone whose
    # rows/cols are re-sent unchanged must not wipe and regenerate its cells.
    resize = data.get("rows", zone.rows) != zone.rows or data.get("cols", zone.cols) != zone.cols
    if resize:
        occupied = (
            db.query(func.count(CellStock.id))
            .join(WarehouseCell, CellStock.cell_id == WarehouseCell.id)
            .filter(WarehouseCell.zone_id == zone.id)
            .scalar() or 0
        )
        if occupied:
            raise HTTPException(
                status_code=400,
                detail=f"Не можна змінити розмір стелажу — {occupied} комірок містять товари. Спочатку очистіть їх.",
            )
    for k, v in data.items():
        setattr(zone, k, v)
    if resize:
        db.query(WarehouseCell).filter(WarehouseCell.zone_id == zone.id).delete()
        for cell in _generate_cells(zone):
            db.add(cell)
    db.commit()
    db.refresh(zone)
    return _zone_out(zone, db)


@_full.delete("/warehouses/{wh_id}/zones/{zone_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_zone(
    wh_id:   int,
    zone_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> None:
    _get_warehouse(wh_id, org, db)
    zone = db.query(WarehouseZone).filter(
        WarehouseZone.id == zone_id, WarehouseZone.warehouse_id == wh_id,
        WarehouseZone.organization_id == org.id,
    ).first()
    if not zone:
        raise HTTPException(status_code=404, detail="Zone not found")
    occupied = (
        db.query(func.count(CellStock.id))
        .join(WarehouseCell, CellStock.cell_id == WarehouseCell.id)
        .filter(WarehouseCell.zone_id == zone.id)
        .scalar() or 0
    )
    if occupied:
        raise HTTPException(
            status_code=400,
            detail=f"Не можна видалити стелаж — {occupied} комірок містять товари. Спочатку очистіть їх.",
        )
    db.delete(zone)
    db.commit()


# ── Cells ─────────────────────────────────────────────────────────────────────

def _cell_stock_out(cs: CellStock, org_id: int) -> CellStockOut:
    return CellStockOut(
        product_id=cs.product_id,
        product_name=cs.product.name,
        product_sku=cs.product.sku,
        product_unit=cs.product.unit,
        quantity=cs.quantity,
        image_url=_product_image_url(cs.product, org_id),
    )


def _cells_out_for_zones(zone_ids: list[int], org_id: int, db: Session) -> dict[int, list[CellOut]]:
    """Cells + stock for many zones in two queries (avoids per-cell N+1)."""
    if not zone_ids:
        return {}
    # Order by id = generation (row-major) order, so the CSS grid lays cells out
    # correctly regardless of code string sorting (e.g. A10 vs A2).
    cells = (db.query(WarehouseCell)
               .filter(WarehouseCell.zone_id.in_(zone_ids))
               .order_by(WarehouseCell.id)
               .all())
    stock_by_cell: dict[int, list[CellStock]] = {}
    if cells:
        stock_rows = (db.query(CellStock)
                        .filter(CellStock.cell_id.in_([c.id for c in cells]))
                        .all())
        for cs in stock_rows:
            stock_by_cell.setdefault(cs.cell_id, []).append(cs)
    out: dict[int, list[CellOut]] = {zid: [] for zid in zone_ids}
    for cell in cells:
        out[cell.zone_id].append(CellOut(
            id=cell.id, code=cell.code, notes=cell.notes,
            stock=[_cell_stock_out(cs, org_id) for cs in stock_by_cell.get(cell.id, [])],
        ))
    return out


@_full.get("/zones/{zone_id}/cells", response_model=ZoneWithCellsOut)
def get_zone_cells(
    zone_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> ZoneWithCellsOut:
    zone = db.query(WarehouseZone).filter(
        WarehouseZone.id == zone_id, WarehouseZone.organization_id == org.id,
    ).first()
    if not zone:
        raise HTTPException(status_code=404, detail="Zone not found")
    cells_out = _cells_out_for_zones([zone_id], org.id, db)[zone_id]
    return ZoneWithCellsOut(
        id=zone.id, name=zone.name, rows=zone.rows, cols=zone.cols,
        sort_order=zone.sort_order, cell_count=len(cells_out), created_at=zone.created_at,
        cells=cells_out,
    )


@_full.get("/warehouses/{wh_id}/zones-with-cells", response_model=list[ZoneWithCellsOut])
def list_zones_with_cells(
    wh_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[ZoneWithCellsOut]:
    """Return all zones + their cells for a warehouse in one call."""
    _get_warehouse(wh_id, org, db)
    zones = (db.query(WarehouseZone)
               .filter(WarehouseZone.warehouse_id == wh_id, WarehouseZone.organization_id == org.id)
               .order_by(WarehouseZone.sort_order, WarehouseZone.id)
               .all())
    cells_by_zone = _cells_out_for_zones([z.id for z in zones], org.id, db)
    return [
        ZoneWithCellsOut(
            id=zone.id, name=zone.name, rows=zone.rows, cols=zone.cols,
            sort_order=zone.sort_order, cell_count=len(cells_by_zone[zone.id]),
            created_at=zone.created_at, cells=cells_by_zone[zone.id],
        )
        for zone in zones
    ]


@_full.put("/cells/{cell_id}/stock", response_model=CellStockOut)
def set_cell_stock(
    cell_id: int,
    payload: CellStockSet,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> CellStockOut:
    """Set the absolute quantity of a product in a cell.

    The delta is drawn from / returned to the warehouse's unassigned pool so the
    invariant sum(cells) <= StockEntry.quantity always holds. Raising a cell
    above the unassigned amount is rejected.
    """
    cell = _get_cell(cell_id, org, db)
    product = db.query(Product).filter(
        Product.id == payload.product_id, Product.organization_id == org.id,
    ).first()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")

    wh_id   = _cell_warehouse_id(cell, db)
    _lock_stock_row(payload.product_id, wh_id, org.id, db)
    cs      = db.query(CellStock).filter_by(cell_id=cell_id, product_id=payload.product_id).first()
    current = cs.quantity if cs else Decimal("0")
    delta   = payload.quantity - current

    if delta > 0:
        avail = _unassigned_qty(payload.product_id, wh_id, org.id, db)
        if delta > avail:
            raise HTTPException(
                status_code=400,
                detail=f"Нерозкладено лише {avail} {product.unit} на цьому складі",
            )
        _putaway(cell, payload.product_id, delta, org.id, db,
                 kind=CellMoveKind.adjust, created_by_id=user.id)
    elif delta < 0:
        _pick_from_cell(cell, payload.product_id, -delta, org.id, db,
                        kind=CellMoveKind.adjust, created_by_id=user.id)

    # Keep the CellStock row even when quantity reaches 0 so the product
    # stays registered in the cell.
    cs = db.query(CellStock).filter_by(cell_id=cell_id, product_id=payload.product_id).first()
    if cs is None:
        db.add(CellStock(cell_id=cell_id, product_id=payload.product_id, quantity=Decimal("0")))

    db.commit()
    return CellStockOut(
        product_id=payload.product_id, product_name=product.name,
        product_sku=product.sku, product_unit=product.unit, quantity=payload.quantity,
        image_url=_product_image_url(product, org.id),
    )


@_full.delete("/cells/{cell_id}/stock/{product_id}", status_code=status.HTTP_204_NO_CONTENT)
def remove_cell_stock(
    cell_id:    int,
    product_id: int,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> None:
    """Empty a cell — the quantity returns to the warehouse's unassigned pool."""
    cell = _get_cell(cell_id, org, db)
    cs = db.query(CellStock).filter_by(cell_id=cell_id, product_id=product_id).first()
    if cs:
        took = _pick_from_cell(cell, product_id, cs.quantity, org.id, db,
                               kind=CellMoveKind.adjust, created_by_id=user.id)
        if took == 0:
            db.delete(cs)
    db.commit()


@_full.post("/cells/{cell_id}/assign", response_model=CellStockOut, status_code=status.HTTP_201_CREATED)
def assign_cell_product(
    cell_id: int,
    payload: CellAssign,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> CellStockOut:
    """Assign a product to a cell. qty=0 → ADJUSTMENT (register only); qty>0 → PURCHASE_IN."""
    cell = _get_cell(cell_id, org, db)
    product = db.query(Product).filter_by(id=payload.product_id, organization_id=org.id).first()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    if payload.quantity < 0:
        raise HTTPException(status_code=400, detail="Кількість не може бути від'ємною")

    wh_id = _cell_warehouse_id(cell, db)
    if payload.quantity > 0:
        _lock_stock_row(payload.product_id, wh_id, org.id, db)
    mtype = MovementType.PURCHASE_IN if payload.quantity > 0 else MovementType.ADJUSTMENT
    reason = (
        f"Отримання в комірку {cell.code}"
        if payload.quantity > 0
        else f"Реєстрація в комірку {cell.code}"
    )
    m = WarehouseMovement(
        organization_id=org.id,
        type=mtype,
        product_id=payload.product_id,
        warehouse_to_id=wh_id,
        quantity=payload.quantity,
        reason=reason,
        created_by_id=user.id,
    )
    db.add(m)
    db.flush()
    _apply_movement(m, db)
    if payload.quantity > 0:
        _putaway(cell, payload.product_id, payload.quantity, org.id, db,
                 kind=CellMoveKind.putaway, movement_id=m.id, created_by_id=user.id)
    else:
        cs_existing = db.query(CellStock).filter_by(cell_id=cell.id, product_id=payload.product_id).first()
        if not cs_existing:
            db.add(CellStock(cell_id=cell.id, product_id=payload.product_id, quantity=Decimal("0")))
    db.commit()

    cs = db.query(CellStock).filter_by(cell_id=cell_id, product_id=payload.product_id).first()
    return CellStockOut(
        product_id=payload.product_id,
        product_name=product.name,
        product_sku=product.sku,
        product_unit=product.unit,
        quantity=cs.quantity if cs else Decimal("0"),
        image_url=_product_image_url(product, org.id),
    )


@_full.patch("/cells/{cell_id}", response_model=CellOut)
def update_cell(
    cell_id: int,
    payload: CellNotesUpdate,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    _:    User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> CellOut:
    """Edit a cell's free-text note (e.g. «верхня полиця»)."""
    cell = _get_cell(cell_id, org, db)
    cell.notes = (payload.notes or "").strip() or None
    db.commit()
    stock_rows = db.query(CellStock).filter(CellStock.cell_id == cell.id).all()
    return CellOut(
        id=cell.id, code=cell.code, notes=cell.notes,
        stock=[_cell_stock_out(cs, org.id) for cs in stock_rows],
    )


# ── Putaway / relocate / locations ────────────────────────────────────────────

@_full.get("/warehouses/{wh_id}/unassigned", response_model=list[UnassignedItemOut])
def list_unassigned(
    wh_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[UnassignedItemOut]:
    """Products with floor stock not yet put away into a cell in this warehouse."""
    _get_warehouse(wh_id, org, db)
    entries = (
        db.query(StockEntry, Product)
        .join(Product, Product.id == StockEntry.product_id)
        .filter(StockEntry.organization_id == org.id,
                StockEntry.warehouse_id == wh_id,
                StockEntry.quantity > 0)
        .order_by(Product.name)
        .all()
    )
    # One grouped query for all cell allocations in this warehouse (avoids N+1).
    assigned_map = dict(
        db.query(CellStock.product_id, func.sum(CellStock.quantity))
        .join(WarehouseCell, WarehouseCell.id == CellStock.cell_id)
        .join(WarehouseZone, WarehouseZone.id == WarehouseCell.zone_id)
        .filter(WarehouseZone.warehouse_id == wh_id, CellStock.quantity > 0)
        .group_by(CellStock.product_id)
        .all()
    )
    out: list[UnassignedItemOut] = []
    for e, p in entries:
        unassigned = e.quantity - assigned_map.get(p.id, Decimal("0"))
        if unassigned > 0:
            out.append(UnassignedItemOut(
                product_id=p.id, product_name=p.name, product_sku=p.sku,
                unit=p.unit, unassigned=unassigned,
            ))
    return out


@_full.post("/cells/{cell_id}/putaway", response_model=CellStockOut)
def putaway_to_cell(
    cell_id: int,
    payload: PutawayRequest,
    bg:   BackgroundTasks,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> CellStockOut:
    """Move quantity from the unassigned pool into a cell."""
    cell = _get_cell(cell_id, org, db)
    product = db.query(Product).filter_by(id=payload.product_id, organization_id=org.id).first()
    if not product:
        raise HTTPException(status_code=404, detail="Product not found")
    if payload.quantity <= 0:
        raise HTTPException(status_code=400, detail="Кількість має бути більшою за 0")

    wh_id = _cell_warehouse_id(cell, db)
    _lock_stock_row(payload.product_id, wh_id, org.id, db)
    avail = _unassigned_qty(payload.product_id, wh_id, org.id, db)
    if payload.quantity > avail:
        raise HTTPException(
            status_code=400,
            detail=f"Нерозкладено лише {avail} {product.unit} на цьому складі",
        )
    _putaway(cell, payload.product_id, payload.quantity, org.id, db, created_by_id=user.id)
    db.commit()
    bg.add_task(broadcast_warehouse, org.id, "stock")

    cs = db.query(CellStock).filter_by(cell_id=cell_id, product_id=payload.product_id).first()
    return CellStockOut(
        product_id=payload.product_id, product_name=product.name,
        product_sku=product.sku, product_unit=product.unit, quantity=cs.quantity if cs else Decimal("0"),
        image_url=_product_image_url(product, org.id),
    )


@_full.post("/cells/relocate", status_code=status.HTTP_204_NO_CONTENT)
def relocate_between_cells(
    payload: RelocateRequest,
    bg:   BackgroundTasks,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> None:
    """Move quantity from one cell to another within the same warehouse."""
    if payload.quantity <= 0:
        raise HTTPException(status_code=400, detail="Кількість має бути більшою за 0")
    if payload.from_cell_id == payload.to_cell_id:
        raise HTTPException(status_code=400, detail="Комірки збігаються")
    src = _get_cell(payload.from_cell_id, org, db)
    dst = _get_cell(payload.to_cell_id, org, db)
    if _cell_warehouse_id(src, db) != _cell_warehouse_id(dst, db):
        raise HTTPException(status_code=400, detail="Переміщення можливе лише в межах одного складу")

    _lock_stock_row(payload.product_id, _cell_warehouse_id(src, db), org.id, db)
    src_cs = db.query(CellStock).filter_by(cell_id=src.id, product_id=payload.product_id).first()
    if not src_cs or src_cs.quantity < payload.quantity:
        have = src_cs.quantity if src_cs else Decimal("0")
        raise HTTPException(status_code=400, detail=f"У комірці лише {have}")

    src_cs.quantity -= payload.quantity
    if src_cs.quantity <= 0:
        db.delete(src_cs)
    dst_cs = db.query(CellStock).filter_by(cell_id=dst.id, product_id=payload.product_id).first()
    if dst_cs:
        dst_cs.quantity += payload.quantity
    else:
        db.add(CellStock(cell_id=dst.id, product_id=payload.product_id, quantity=payload.quantity))
    _log_cell_move(db, org_id=org.id, product_id=payload.product_id, quantity=payload.quantity,
                   kind=CellMoveKind.relocate, cell_from_id=src.id, cell_to_id=dst.id,
                   created_by_id=user.id)
    db.commit()
    bg.add_task(broadcast_warehouse, org.id, "stock")


@_full.post("/scan-action", response_model=ScanActionResult)
def scan_action(
    payload: ScanActionRequest,
    bg:   BackgroundTasks,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> ScanActionResult:
    """Execute a warehouse operation chosen via a functional ACTION QR.

    write_off      — remove qty from a cell and the warehouse ledger (WRITE_OFF).
    transfer       — cell → cell: relocate within a warehouse, or TRANSFER across.
    receive        — book qty into a cell (PURCHASE_IN); unit_cost optional → AVCO.
    stocktake      — set a cell to the counted qty and correct the warehouse total.
    sale_out       — ship qty from a cell (SALE_OUT); unit_price optional → revenue.
    defect         — mark qty defective, remove from a cell (DEFECT).
    production_in  — book finished goods into a cell (PRODUCTION_IN).
    production_out — issue components to production from a cell (PRODUCTION_OUT).
    """
    product = _get_product(payload.product_id, org, db)
    if payload.quantity < 0 or (payload.quantity == 0 and payload.action != ScanAction.stocktake):
        raise HTTPException(status_code=400, detail="Кількість має бути > 0")

    cell  = _get_cell(payload.cell_id, org, db)
    wh_id = _cell_warehouse_id(cell, db)
    _lock_stock_row(payload.product_id, wh_id, org.id, db)

    def _outbound(mtype: MovementType, reason: str, *,
                  unit_price: Decimal | None = None, replenish: bool = False) -> None:
        """Pick qty from the cell and decrement the warehouse via an outbound ledger move."""
        cs   = db.query(CellStock).filter_by(cell_id=cell.id, product_id=payload.product_id).first()
        have = cs.quantity if cs else Decimal("0")
        if have < payload.quantity:
            raise HTTPException(status_code=400, detail=f"У комірці лише {have}")
        revenue = (payload.quantity * unit_price) if unit_price else None
        m = WarehouseMovement(
            organization_id=org.id, type=mtype,
            product_id=payload.product_id, warehouse_from_id=wh_id,
            quantity=payload.quantity, unit_price=unit_price, total_revenue=revenue,
            reason=reason, created_by_id=user.id,
        )
        db.add(m)
        db.flush()
        # Pick from the cell BEFORE _apply_movement so the FIFO clamp leaves it alone.
        _pick_from_cell(cell, payload.product_id, payload.quantity, org.id, db,
                        movement_id=m.id, created_by_id=user.id)
        _apply_movement(m, db)
        if replenish:
            _check_and_auto_replenish(payload.product_id, org.id, db)
        db.commit()

    def _inbound(mtype: MovementType, reason: str, *, unit_cost: Decimal | None = None) -> None:
        """Add qty to the warehouse via an inbound ledger move and put it away into the cell."""
        total = (payload.quantity * unit_cost) if unit_cost else None
        m = WarehouseMovement(
            organization_id=org.id, type=mtype,
            product_id=payload.product_id, warehouse_to_id=wh_id,
            quantity=payload.quantity, unit_cost=unit_cost, total_cost=total,
            reason=reason, created_by_id=user.id,
        )
        db.add(m)
        db.flush()
        _apply_movement(m, db)   # AVCO before stock add when unit_cost is set
        avail = _unassigned_qty(payload.product_id, wh_id, org.id, db)
        _putaway(cell, payload.product_id, min(payload.quantity, avail), org.id, db,
                 movement_id=m.id, created_by_id=user.id)
        db.commit()

    # ── Переміщення: cell → cell (relocate), or warehouse → warehouse (TRANSFER) ─
    if payload.action == ScanAction.transfer:
        if not payload.to_cell_id:
            raise HTTPException(status_code=400, detail="Не вказано комірку призначення")
        if payload.to_cell_id == cell.id:
            raise HTTPException(status_code=400, detail="Комірки збігаються")
        dst    = _get_cell(payload.to_cell_id, org, db)
        dst_wh = _cell_warehouse_id(dst, db)
        src_cs = db.query(CellStock).filter_by(cell_id=cell.id, product_id=payload.product_id).first()
        if not src_cs or src_cs.quantity < payload.quantity:
            have = src_cs.quantity if src_cs else Decimal("0")
            raise HTTPException(status_code=400, detail=f"У комірці лише {have}")

        if dst_wh == wh_id:
            # Same warehouse → pure relocate; warehouse total unchanged.
            src_cs.quantity -= payload.quantity
            if src_cs.quantity <= 0:
                db.delete(src_cs)
            dst_cs = db.query(CellStock).filter_by(cell_id=dst.id, product_id=payload.product_id).first()
            if dst_cs:
                dst_cs.quantity += payload.quantity
            else:
                db.add(CellStock(cell_id=dst.id, product_id=payload.product_id, quantity=payload.quantity))
            _log_cell_move(db, org_id=org.id, product_id=payload.product_id, quantity=payload.quantity,
                           kind=CellMoveKind.relocate, cell_from_id=cell.id, cell_to_id=dst.id,
                           created_by_id=user.id)
            db.commit()
            bg.add_task(broadcast_warehouse, org.id, "stock")
            return ScanActionResult(message=f"Переміщено {payload.quantity} {product.unit}: {cell.code} → {dst.code}")

        # Cross-warehouse → real TRANSFER ledger movement (decrements src wh, adds dst wh).
        _lock_stock_row(payload.product_id, dst_wh, org.id, db)
        m = WarehouseMovement(
            organization_id=org.id, type=MovementType.TRANSFER,
            product_id=payload.product_id, warehouse_from_id=wh_id, warehouse_to_id=dst_wh,
            quantity=payload.quantity, reason=f"Переміщення між складами (скан) {cell.code} → {dst.code}",
            created_by_id=user.id,
        )
        db.add(m)
        db.flush()
        _pick_from_cell(cell, payload.product_id, payload.quantity, org.id, db,
                        movement_id=m.id, created_by_id=user.id)
        _apply_movement(m, db)   # src wh -= qty, dst wh += qty, clamps cells in both
        avail = _unassigned_qty(payload.product_id, dst_wh, org.id, db)
        _putaway(dst, payload.product_id, min(payload.quantity, avail), org.id, db,
                 movement_id=m.id, created_by_id=user.id)
        db.commit()
        bg.add_task(broadcast_warehouse, org.id, "stock")
        return ScanActionResult(message=f"Переміщено між складами {payload.quantity} {product.unit}: {cell.code} → {dst.code}")

    # ── Списання (WRITE_OFF) ───────────────────────────────────────────────────
    if payload.action == ScanAction.write_off:
        _outbound(MovementType.WRITE_OFF, f"Списання (скан) з {cell.code}")
        bg.add_task(broadcast_warehouse, org.id, "stock")
        return ScanActionResult(message=f"Списано {payload.quantity} {product.unit} з {cell.code}")

    # ── Відвантаження (SALE_OUT); unit_price optional → revenue ────────────────
    if payload.action == ScanAction.sale_out:
        _outbound(MovementType.SALE_OUT, f"Відвантаження (скан) з {cell.code}",
                  unit_price=payload.unit_price, replenish=True)
        bg.add_task(broadcast_warehouse, org.id, "stock")
        return ScanActionResult(message=f"Відвантажено {payload.quantity} {product.unit} з {cell.code}")

    # ── Брак (DEFECT) ──────────────────────────────────────────────────────────
    if payload.action == ScanAction.defect:
        _outbound(MovementType.DEFECT, f"Брак (скан) з {cell.code}", replenish=True)
        bg.add_task(broadcast_warehouse, org.id, "stock")
        return ScanActionResult(message=f"Брак {payload.quantity} {product.unit} з {cell.code}")

    # ── Видача у виробництво (PRODUCTION_OUT) ──────────────────────────────────
    if payload.action == ScanAction.production_out:
        _outbound(MovementType.PRODUCTION_OUT, f"Видача у виробництво (скан) з {cell.code}", replenish=True)
        bg.add_task(broadcast_warehouse, org.id, "stock")
        return ScanActionResult(message=f"Видано у виробництво {payload.quantity} {product.unit} з {cell.code}")

    # ── Прийом (PURCHASE_IN); unit_cost optional → AVCO ────────────────────────
    if payload.action == ScanAction.receive:
        _inbound(MovementType.PURCHASE_IN, f"Прийом (скан) у {cell.code}", unit_cost=payload.unit_cost)
        bg.add_task(broadcast_warehouse, org.id, "stock")
        return ScanActionResult(message=f"Прийнято {payload.quantity} {product.unit} у {cell.code}")

    # ── Оприбуткування з виробництва (PRODUCTION_IN) ───────────────────────────
    if payload.action == ScanAction.production_in:
        _inbound(MovementType.PRODUCTION_IN, f"Оприбуткування з виробництва (скан) у {cell.code}")
        bg.add_task(broadcast_warehouse, org.id, "stock")
        return ScanActionResult(message=f"Оприбутковано {payload.quantity} {product.unit} у {cell.code}")

    # ── Інвентаризація: set cell to counted qty, correct warehouse total ───────
    cs      = db.query(CellStock).filter_by(cell_id=cell.id, product_id=payload.product_id).first()
    current = cs.quantity if cs else Decimal("0")
    delta   = payload.quantity - current
    if delta == 0:
        if cs is None:
            db.add(CellStock(cell_id=cell.id, product_id=payload.product_id, quantity=Decimal("0")))
            db.commit()
        return ScanActionResult(message=f"{cell.code}: без змін ({current} {product.unit})")
    if delta > 0:
        m = WarehouseMovement(
            organization_id=org.id, type=MovementType.ADJUSTMENT,
            product_id=payload.product_id, warehouse_to_id=wh_id,
            quantity=delta, reason=f"Інвентаризація {cell.code} (надлишок)",
            created_by_id=user.id,
        )
        db.add(m)
        db.flush()
        _apply_movement(m, db)   # total += delta
        _putaway(cell, payload.product_id, delta, org.id, db,
                 movement_id=m.id, created_by_id=user.id)
    else:
        short = -delta
        m = WarehouseMovement(
            organization_id=org.id, type=MovementType.WRITE_OFF,
            product_id=payload.product_id, warehouse_from_id=wh_id,
            quantity=short, reason=f"Інвентаризація {cell.code} (нестача)",
            created_by_id=user.id,
        )
        db.add(m)
        db.flush()
        _pick_from_cell(cell, payload.product_id, short, org.id, db,
                        movement_id=m.id, created_by_id=user.id)
        _apply_movement(m, db)   # total -= short
    db.commit()
    bg.add_task(broadcast_warehouse, org.id, "stock")
    return ScanActionResult(message=f"Інвентаризація {cell.code}: {current} → {payload.quantity} {product.unit}")


@_full.get("/products/{product_id}/locations", response_model=ProductLocationsOut)
def product_locations(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> ProductLocationsOut:
    """Where a product physically sits: cells + unassigned pool, per warehouse."""
    _get_product(product_id, org, db)
    entries = (
        db.query(StockEntry, Warehouse)
        .join(Warehouse, Warehouse.id == StockEntry.warehouse_id)
        .filter(StockEntry.organization_id == org.id,
                StockEntry.product_id == product_id,
                StockEntry.quantity > 0)
        .order_by(Warehouse.name)
        .all()
    )
    wh_ids = [wh.id for _, wh in entries]
    cells_by_warehouse: dict[int, list[ProductCellLocationOut]] = {wh_id: [] for wh_id in wh_ids}
    if wh_ids:
        cell_rows = (
            db.query(CellStock, WarehouseCell, WarehouseZone)
            .join(WarehouseCell, WarehouseCell.id == CellStock.cell_id)
            .join(WarehouseZone, WarehouseZone.id == WarehouseCell.zone_id)
            .filter(
                WarehouseZone.organization_id == org.id,
                WarehouseZone.warehouse_id.in_(wh_ids),
                CellStock.product_id == product_id,
                CellStock.quantity > 0,
            )
            .order_by(WarehouseZone.name, WarehouseCell.code)
            .all()
        )
        for cs, c, z in cell_rows:
            cells_by_warehouse.setdefault(z.warehouse_id, []).append(
                ProductCellLocationOut(cell_id=c.id, zone_name=z.name, code=c.code, quantity=cs.quantity)
            )

    warehouses: list[ProductWarehouseLocationOut] = []
    for e, wh in entries:
        cells = cells_by_warehouse.get(wh.id, [])
        assigned = sum((c.quantity for c in cells), Decimal("0"))
        warehouses.append(ProductWarehouseLocationOut(
            warehouse_id=wh.id, warehouse_name=wh.name, cells=cells,
            unassigned=e.quantity - assigned, total=e.quantity,
        ))
    return ProductLocationsOut(product_id=product_id, warehouses=warehouses)


@_full.get("/products/{product_id}/cell-history", response_model=list[CellMovementOut])
def product_cell_history(
    product_id: int,
    limit: int = Query(50, le=200),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[CellMovementOut]:
    """Audit trail of bin relocations for a product (newest first)."""
    _get_product(product_id, org, db)
    rows = (
        db.query(CellMovement)
        .filter(CellMovement.organization_id == org.id, CellMovement.product_id == product_id)
        .order_by(CellMovement.created_at.desc(), CellMovement.id.desc())
        .limit(limit)
        .all()
    )

    cell_ids = {cid for r in rows for cid in (r.cell_from_id, r.cell_to_id) if cid}
    labels: dict[int, str] = {}
    if cell_ids:
        for cell_id, zone_name, cell_code in (
            db.query(WarehouseCell.id, WarehouseZone.name, WarehouseCell.code)
            .join(WarehouseZone, WarehouseZone.id == WarehouseCell.zone_id)
            .filter(WarehouseCell.id.in_(cell_ids), WarehouseZone.organization_id == org.id)
            .all()
        ):
            labels[cell_id] = f"{zone_name} {cell_code}"

    p = _get_product(product_id, org, db)
    return [
        CellMovementOut(
            id=r.id, product_id=r.product_id, product_name=p.name,
            quantity=r.quantity, kind=r.kind.value,
            cell_from=labels.get(r.cell_from_id) if r.cell_from_id else None,
            cell_to=labels.get(r.cell_to_id) if r.cell_to_id else None,
            created_at=r.created_at,
        )
        for r in rows
    ]


# ── Counterparties ────────────────────────────────────────────────────────────

@_full.get("/counterparties", response_model=list[CounterpartyOut])
def list_counterparties(
    cp_type: str | None = Query(None, alias="type"),
    search:  str | None = Query(None),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[CounterpartyOut]:
    q = db.query(Counterparty).filter(Counterparty.organization_id == org.id)
    if cp_type:
        q = q.filter(Counterparty.type == cp_type)
    if search:
        q = q.filter(Counterparty.name.ilike(f"%{search}%") | Counterparty.email.ilike(f"%{search}%"))
    rows = q.order_by(Counterparty.name).all()
    return [CounterpartyOut.model_validate(r) for r in rows]


@_full.post("/counterparties", response_model=CounterpartyOut, status_code=status.HTTP_201_CREATED)
def create_counterparty(
    payload: CounterpartyCreate,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    _:    User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> CounterpartyOut:
    cp = Counterparty(**payload.model_dump(), organization_id=org.id)
    db.add(cp)
    db.commit()
    db.refresh(cp)
    return CounterpartyOut.model_validate(cp)


@_full.get("/counterparties/{cp_id}", response_model=CounterpartyOut)
def get_counterparty(
    cp_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> CounterpartyOut:
    return CounterpartyOut.model_validate(_get_counterparty(cp_id, org, db))


@_full.patch("/counterparties/{cp_id}", response_model=CounterpartyOut)
def update_counterparty(
    cp_id:   int,
    payload: CounterpartyUpdate,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> CounterpartyOut:
    cp = _get_counterparty(cp_id, org, db)
    for k, v in payload.model_dump(exclude_unset=True).items():
        setattr(cp, k, v)
    db.commit()
    db.refresh(cp)
    return CounterpartyOut.model_validate(cp)


@_full.delete("/counterparties/{cp_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_counterparty(
    cp_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> None:
    cp = _get_counterparty(cp_id, org, db)
    # Nullify references before deletion so existing orders are not orphaned
    db.query(Order).filter_by(organization_id=org.id, counterparty_id=cp.id).update({"counterparty_id": None})
    db.delete(cp)
    db.commit()


@_full.post("/counterparties/{cp_id}/adjust-balance", response_model=CounterpartyOut)
def adjust_balance(
    cp_id:   int,
    payload: CounterpartyBalanceAdjust,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> CounterpartyOut:
    """Manually adjust counterparty balance (record a payment)."""
    cp = _get_counterparty(cp_id, org, db)
    cp.balance -= payload.delta   # delta positive = they paid us → reduces their balance
    db.commit()
    db.refresh(cp)
    return CounterpartyOut.model_validate(cp)


# ── Products ──────────────────────────────────────────────────────────────────

@router.get("/products", response_model=list[ProductOut])
def list_products(
    search:   str | None = Query(None),
    category: str | None = Query(None),
    archived: bool = Query(False),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[ProductOut]:
    q = db.query(Product).filter(
        Product.organization_id == org.id,
        Product.is_active.is_(not archived),
    )
    if search:
        q = q.filter(Product.name.ilike(f"%{search}%") | Product.sku.ilike(f"%{search}%"))
    if category:
        q = q.filter(Product.categories.contains([category]))
    rows = q.order_by(Product.name).all()
    return [_make_product_out(r, org.id) for r in rows]


@router.get("/products/options", response_model=list[ProductOptionOut])
def list_product_options(
    search:   str | None = Query(None),
    archived: bool = Query(False),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[ProductOptionOut]:
    """Compact product catalog for selects and comboboxes."""
    q = db.query(Product).filter(
        Product.organization_id == org.id,
        Product.is_active.is_(not archived),
    )
    if search:
        q = q.filter(Product.name.ilike(f"%{search}%") | Product.sku.ilike(f"%{search}%"))
    rows = q.order_by(Product.name).all()
    return [ProductOptionOut.model_validate(r) for r in rows]


def _check_product_limit(org: Organization, db: Session) -> None:
    from app.models.organization import WAREHOUSE_PRODUCT_LIMIT
    limit = WAREHOUSE_PRODUCT_LIMIT[org.plan]
    if limit is None:
        return
    count = db.query(func.count(Product.id)).filter(Product.organization_id == org.id).scalar() or 0
    if count >= limit:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Досягнуто ліміт {limit} номенклатур для тарифу {org.plan.value}",
        )


@router.post("/products", response_model=ProductOut, status_code=status.HTTP_201_CREATED)
def create_product(
    payload: ProductCreate,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin)),
) -> ProductOut:
    _check_product_limit(org, db)
    data = payload.model_dump()
    if "cell_limit" in data:
        data["box_limit"] = data.pop("cell_limit")
    p = Product(**data, organization_id=org.id, created_by_id=user.id)
    db.add(p)
    db.commit()
    db.refresh(p)
    return ProductOut.model_validate(p)


_IMPORT_COL_MAP = {
    "Назва виробу":         "name",
    "Категорії":            "categories",
    "SKU":                  "sku",
    "Штрих-код":            "barcode",
    "Одиниця виміру":       "unit",
    "Середньозважена ціна": "cost_price",
    "Роздрібна ціна":       "sale_price",
    # "Собівартість виробу" intentionally excluded — Monofarm computes this from specs
    "Опис":                 "description",
}

_EXPORT_HEADERS = [
    "Назва виробу", "Категорії", "SKU", "Штрих-код", "Одиниця виміру",
    "Середньозважена ціна", "Роздрібна ціна", "Опис",
]


def _rows_from_upload(file: UploadFile) -> list[dict]:
    raw = file.file.read()
    fname = (file.filename or "").lower()
    if fname.endswith(".xlsx") or fname.endswith(".xls"):
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = [str(c).strip() if c is not None else "" for c in next(rows_iter, [])]
        return [dict(zip(headers, (str(c).strip() if c is not None else "" for c in row))) for row in rows_iter]
    text = raw.decode("utf-8-sig")
    delimiter = "\t" if "\t" in text[:1024] else ","
    return list(csv.DictReader(io.StringIO(text), delimiter=delimiter))


def _sync_categories(names: list[str], org_id: int, db: Session) -> int:
    """Create any ProductCategory rows that don't yet exist for this org. Returns count created."""
    if not names:
        return 0
    existing = {
        r.name for r in db.query(ProductCategory.name).filter(
            ProductCategory.organization_id == org_id,
            ProductCategory.name.in_(names),
        ).all()
    }
    created = 0
    for name in dict.fromkeys(names):  # preserve order, deduplicate
        if name and name not in existing:
            db.add(ProductCategory(organization_id=org_id, name=name))
            created += 1
    return created


def _parse_product_fields(row: dict) -> dict:
    fields: dict = {}
    for col, field in _IMPORT_COL_MAP.items():
        val = (row.get(col) or "").strip()
        if not val:
            continue
        if field == "categories":
            fields[field] = [c.strip() for c in val.split(",") if c.strip()]
        elif field in ("cost_price", "sale_price", "direct_cost"):
            try:
                fields[field] = Decimal(val.replace(",", ".").replace(" ", ""))
            except Exception:
                pass
        else:
            fields[field] = val
    return fields


_COMPARABLE_FIELDS = ["name", "barcode", "unit", "sale_price", "cost_price", "description", "categories"]


@router.post("/products/import/preview")
def preview_import(
    file: UploadFile = File(...),
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    _:    User         = Depends(require_roles(UserRole.admin)),
) -> dict:
    try:
        rows = _rows_from_upload(file)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Не вдалося прочитати файл: {exc}") from exc

    db_products = {
        p.sku: p
        for p in db.query(Product).filter(Product.organization_id == org.id, Product.is_active).all()
    }
    file_skus: set[str] = set()
    new_items: list[dict] = []
    existing_items: list[dict] = []

    for row in rows:
        fields = _parse_product_fields(row)
        sku = fields.get("sku", "").strip()
        name = fields.get("name", "").strip()
        if not sku and not name:
            continue
        if sku:
            file_skus.add(sku)

        if sku and sku in db_products:
            p = db_products[sku]
            changes: dict = {}
            for field in _COMPARABLE_FIELDS:
                new_val = fields.get(field)
                if new_val is None:
                    continue
                old_val = getattr(p, field, None)
                if field in ("cost_price", "sale_price", "direct_cost"):
                    try:
                        new_dec = Decimal(str(new_val)).normalize()
                        old_dec = Decimal(str(old_val)).normalize() if old_val is not None else None
                        if new_dec == old_dec:
                            continue
                        new_str = str(new_dec)
                        old_str = str(old_dec) if old_dec is not None else ""
                    except Exception:
                        new_str, old_str = str(new_val), str(old_val) if old_val is not None else ""
                elif isinstance(new_val, list):
                    new_str = ", ".join(new_val)
                    old_str = ", ".join(old_val or [])
                else:
                    new_str = str(new_val).strip()
                    old_str = str(old_val).strip() if old_val is not None else ""
                if new_str != old_str:
                    changes[field] = {"from": old_str, "to": new_str}
            existing_items.append({"id": p.id, "name": p.name, "sku": p.sku, "changes": changes})
        else:
            if not name:
                continue
            item: dict = {"name": name, "sku": sku}
            for f in ("unit", "sale_price", "cost_price"):
                if f in fields:
                    item[f] = str(fields[f])
            new_items.append(item)

    missing_items = [
        {"id": p.id, "name": p.name, "sku": p.sku}
        for sku, p in db_products.items()
        if sku not in file_skus
    ]

    headers = list(rows[0].keys()) if rows else []
    raw_rows = [[str(row.get(h) or "") for h in headers] for row in rows]
    mapping = {str(i): _IMPORT_COL_MAP.get(h) for i, h in enumerate(headers)}

    return {
        "headers": headers,
        "rows": raw_rows,
        "mapping": mapping,
        "summary": {
            "new": len(new_items),
            "existing": len(existing_items),
            "missing": len(missing_items),
        },
        "new": new_items,
        "existing": existing_items,
        "missing": missing_items,
    }


@router.post("/products/import")
def import_products(
    file:            UploadFile = File(...),
    action_new:      str = Query("import",  pattern="^(import|skip)$"),
    action_existing: str = Query("update",  pattern="^(update|skip)$"),
    action_missing:  str = Query("nothing", pattern="^(nothing|hide)$"),
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin)),
) -> dict:
    try:
        rows = _rows_from_upload(file)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Не вдалося прочитати файл: {exc}") from exc

    db_products = {
        p.sku: p
        for p in db.query(Product).filter(Product.organization_id == org.id, Product.is_active).all()
    }
    file_skus: set[str] = set()
    all_cats: list[str] = []
    created = updated = skipped = hidden = 0

    for row in rows:
        fields = _parse_product_fields(row)
        sku = fields.get("sku", "").strip()
        name = fields.get("name", "").strip()
        if not sku and not name:
            skipped += 1
            continue
        if sku:
            file_skus.add(sku)
        all_cats.extend(fields.get("categories") or [])

        if sku and sku in db_products:
            if action_existing == "update":
                p = db_products[sku]
                for k, v in fields.items():
                    if k != "sku":
                        setattr(p, k, v)
                p.is_active = True
                updated += 1
            else:
                skipped += 1
        else:
            if action_new == "import":
                if not name:
                    skipped += 1
                    continue
                db.add(Product(organization_id=org.id, created_by_id=user.id, **fields))
                created += 1
            else:
                skipped += 1

    if action_missing == "hide":
        for sku, p in db_products.items():
            if sku not in file_skus:
                p.is_active = False
                hidden += 1

    new_cats = _sync_categories(all_cats, org.id, db)

    db.commit()
    return {"created": created, "updated": updated, "skipped": skipped, "hidden": hidden, "new_categories": new_cats}


def _export_tsv(rows: list, ts: str) -> Response:
    buf = io.StringIO()
    buf.write("﻿")  # UTF-8 BOM for Excel compatibility
    writer = csv.writer(buf, delimiter="\t", lineterminator="\r\n")
    writer.writerow(_EXPORT_HEADERS)
    for p in rows:
        writer.writerow([
            p.name or "",
            ", ".join(p.categories or []),
            p.sku or "",
            p.barcode or "",
            p.unit or "",
            str(p.cost_price) if p.cost_price is not None else "",
            str(p.sale_price) if p.sale_price is not None else "",
            p.description or "",
        ])
    from urllib.parse import quote
    filename = f"номенклатури_{ts}.tsv"
    return Response(
        content=buf.getvalue().encode("utf-8"),
        media_type="text/tab-separated-values; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


def _export_xlsx(rows: list, org_id: int, ts: str) -> Response:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Номенклатури"

    ws.append(list(_EXPORT_HEADERS))

    for p in rows:
        ws.append([
            p.name or "",
            ", ".join(p.categories or []),
            p.sku or "",
            p.barcode or "",
            p.unit or "",
            str(p.cost_price) if p.cost_price is not None else "",
            str(p.sale_price) if p.sale_price is not None else "",
            p.description or "",
        ])

    from urllib.parse import quote
    buf = io.BytesIO()
    wb.save(buf)
    filename = f"номенклатури_{ts}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.get("/products/export")
def export_products(
    ids: str | None = Query(None, description="Comma-separated product IDs; omit for all"),
    fmt: str        = Query("xlsx"),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> Response:
    import traceback as _tb
    q = db.query(Product).filter(Product.organization_id == org.id, Product.is_active)
    if ids:
        id_list = [int(i) for i in ids.split(",") if i.strip().isdigit()]
        q = q.filter(Product.id.in_(id_list))
    rows = q.order_by(Product.name).all()
    ts = datetime.utcnow().strftime("%Y-%m-%d_%H-%M")
    if fmt == "xlsx":
        try:
            return _export_xlsx(rows, org.id, ts)
        except Exception:
            resp = _export_tsv(rows, ts)
            resp.headers["X-Export-Error"] = _tb.format_exc()[-400:]
            return resp
    return _export_tsv(rows, ts)


@router.get("/products/{product_id}", response_model=ProductOut)
def get_product(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> ProductOut:
    return _make_product_out(_get_product(product_id, org, db), org.id)


@router.patch("/products/{product_id}", response_model=ProductOut)
def update_product(
    product_id: int,
    payload:    ProductUpdate,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> ProductOut:
    p = _get_product(product_id, org, db)
    data = payload.model_dump(exclude_unset=True)
    if "cell_limit" in data:
        data["box_limit"] = data.pop("cell_limit")
    for k, v in data.items():
        setattr(p, k, v)
    db.commit()
    db.refresh(p)
    return _make_product_out(p, org.id)


@router.post("/products/{product_id}/archive", response_model=ProductOut)
def archive_product(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> ProductOut:
    p = _get_product(product_id, org, db)
    p.is_active = False
    db.commit()
    db.refresh(p)
    return ProductOut.model_validate(p)


@router.post("/products/{product_id}/restore", response_model=ProductOut)
def restore_product(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> ProductOut:
    p = _get_product(product_id, org, db)
    p.is_active = True
    db.commit()
    db.refresh(p)
    return _make_product_out(p, org.id)


@router.post("/products/{product_id}/copy", response_model=ProductOut, status_code=201)
def copy_product(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> ProductOut:
    p = _get_product(product_id, org, db)
    # Find a unique SKU: base-copy, base-copy-2, base-copy-3 …
    base = p.sku + "-copy"

    taken = {
        row.sku for row in
        db.query(Product.sku).filter(Product.organization_id == org.id, Product.sku.like(base + "%")).all()
    }
    new_sku = base
    n = 2
    while new_sku in taken:
        new_sku = f"{base}-{n}"
        n += 1
    copy = Product(
        organization_id=org.id,
        name=p.name + " (копія)",
        sku=new_sku,
        barcode=None,
        categories=list(p.categories),
        unit=p.unit,
        description=p.description,
        is_active=True,
        sale_price=p.sale_price,
        min_stock=p.min_stock,
        desired_stock=p.desired_stock,
    )
    db.add(copy)
    db.commit()
    db.refresh(copy)
    return _make_product_out(copy, org.id)


# ── Product image ─────────────────────────────────────────────────────────────

@router.post("/products/{product_id}/image", response_model=ProductOut)
async def upload_product_image(
    product_id: int,
    file: UploadFile = File(...),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> ProductOut:
    from app.services import storage as storage_svc
    if file.content_type not in _IMAGE_MIME_EXT:
        raise HTTPException(status_code=400, detail="Підтримуються тільки JPEG, PNG, WebP.")
    data = await file.read()
    if len(data) > _IMAGE_MAX_BYTES:
        raise HTTPException(status_code=400, detail="Файл занадто великий (макс. 8 МБ).")
    p = _get_product(product_id, org, db)
    if p.image_key:
        storage_svc.delete(p.image_key, org.id, prefix=_IMAGE_PREFIX)
    ext = _IMAGE_MIME_EXT[file.content_type]
    key = f"{uuid.uuid4().hex}.{ext}"
    storage_svc.put(key, data, org.id, prefix=_IMAGE_PREFIX)
    p.image_key = key
    db.commit()
    db.refresh(p)
    return _make_product_out(p, org.id)


@router.delete("/products/{product_id}/image", response_model=ProductOut)
def delete_product_image(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> ProductOut:
    from app.services import storage as storage_svc
    p = _get_product(product_id, org, db)
    if p.image_key:
        storage_svc.delete(p.image_key, org.id, prefix=_IMAGE_PREFIX)
        p.image_key = None
        db.commit()
        db.refresh(p)
    return _make_product_out(p, org.id)


@router.get("/products/{product_id}/image")
def serve_product_image(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> Response:
    """Serve image bytes for local-storage mode (S3 mode returns presigned URL instead)."""
    from app.services import storage as storage_svc
    p = _get_product(product_id, org, db)
    if not p.image_key:
        raise HTTPException(status_code=404, detail="No image")
    try:
        data = storage_svc.get_bytes(p.image_key, org.id, prefix=_IMAGE_PREFIX)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="Image file not found")
    ext = p.image_key.rsplit(".", 1)[-1].lower()
    mime = {"jpg": "image/jpeg", "png": "image/png", "webp": "image/webp"}.get(ext, "image/jpeg")
    return Response(content=data, media_type=mime)


# ── Product images (multi-photo) ──────────────────────────────────────────────

@router.get("/products/{product_id}/images", response_model=list[ProductImageOut])
def list_product_images(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[ProductImageOut]:
    _get_product(product_id, org, db)  # 404 guard
    rows = (
        db.query(ProductImage)
        .filter_by(product_id=product_id, organization_id=org.id)
        .order_by(ProductImage.is_primary.desc(), ProductImage.sort_order, ProductImage.id)
        .all()
    )
    return [
        ProductImageOut(
            id=img.id,
            image_url=_image_url(img.image_key, product_id, org.id),
            is_primary=img.is_primary,
            sort_order=img.sort_order,
        )
        for img in rows
    ]


@router.post("/products/{product_id}/images", response_model=list[ProductImageOut], status_code=201)
async def add_product_image(
    product_id: int,
    file: UploadFile = File(...),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> list[ProductImageOut]:
    from app.services import storage as storage_svc
    if file.content_type not in _IMAGE_MIME_EXT:
        raise HTTPException(status_code=400, detail="Підтримуються тільки JPEG, PNG, WebP.")
    data = await file.read()
    if len(data) > _IMAGE_MAX_BYTES:
        raise HTTPException(status_code=400, detail="Файл занадто великий (макс. 8 МБ).")
    p = _get_product(product_id, org, db)
    ext = _IMAGE_MIME_EXT[file.content_type]
    key = f"{uuid.uuid4().hex}.{ext}"
    storage_svc.put(key, data, org.id, prefix=_IMAGE_PREFIX)
    existing_count = db.query(func.count(ProductImage.id)).filter_by(product_id=product_id, organization_id=org.id).scalar() or 0
    is_primary = existing_count == 0
    img = ProductImage(product_id=product_id, organization_id=org.id, image_key=key, is_primary=is_primary, sort_order=existing_count)
    db.add(img)
    if is_primary:
        p.image_key = key
    db.commit()
    return list_product_images(product_id, db, org)


@router.delete("/products/{product_id}/images/{image_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_product_image_by_id(
    product_id: int,
    image_id:   int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> None:
    from app.services import storage as storage_svc
    img = db.query(ProductImage).filter_by(id=image_id, product_id=product_id, organization_id=org.id).first()
    if not img:
        raise HTTPException(status_code=404, detail="Зображення не знайдено")
    was_primary = img.is_primary
    storage_svc.delete(img.image_key, org.id, prefix=_IMAGE_PREFIX)
    db.delete(img)
    db.flush()
    if was_primary:
        next_img = (
            db.query(ProductImage)
            .filter_by(product_id=product_id, organization_id=org.id)
            .order_by(ProductImage.sort_order, ProductImage.id)
            .first()
        )
        p = _get_product(product_id, org, db)
        if next_img:
            next_img.is_primary = True
            p.image_key = next_img.image_key
        else:
            p.image_key = None
    db.commit()


@router.patch("/products/{product_id}/images/{image_id}/set-primary", response_model=list[ProductImageOut])
def set_primary_product_image(
    product_id: int,
    image_id:   int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> list[ProductImageOut]:
    img = db.query(ProductImage).filter_by(id=image_id, product_id=product_id, organization_id=org.id).first()
    if not img:
        raise HTTPException(status_code=404, detail="Зображення не знайдено")
    db.query(ProductImage).filter_by(product_id=product_id, organization_id=org.id).update({"is_primary": False})
    img.is_primary = True
    p = _get_product(product_id, org, db)
    p.image_key = img.image_key
    db.commit()
    return list_product_images(product_id, db, org)


@router.get("/products/{product_id}/images/{image_key}")
def serve_product_image_by_key(
    product_id: int,
    image_key:  str,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> Response:
    """Serve image bytes for local-storage mode."""
    from app.services import storage as storage_svc
    img = db.query(ProductImage).filter_by(image_key=image_key, product_id=product_id, organization_id=org.id).first()
    if not img:
        raise HTTPException(status_code=404, detail="Image not found")
    try:
        data = storage_svc.get_bytes(image_key, org.id, prefix=_IMAGE_PREFIX)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="Image file not found")
    ext = image_key.rsplit(".", 1)[-1].lower()
    mime = {"jpg": "image/jpeg", "png": "image/png", "webp": "image/webp"}.get(ext, "image/jpeg")
    return Response(content=data, media_type=mime)


@router.delete("/products/{product_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_product(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> None:
    p = _get_product(product_id, org, db)
    blockers: list[str] = []
    if db.query(func.count(WarehouseMovement.id)).filter_by(product_id=product_id).scalar():
        blockers.append("рухи складу")
    if db.query(func.count(OrderItem.id)).filter_by(product_id=product_id).scalar():
        blockers.append("замовлення")
    if db.query(func.count(ProductionBatch.id)).filter_by(product_id=product_id).scalar():
        blockers.append("виробничі партії")
    if db.query(CellStock).filter(CellStock.product_id == product_id, CellStock.quantity > 0).first():
        blockers.append("залишки в комірках")
    if blockers:
        raise HTTPException(
            status_code=409,
            detail=f"Товар задіяний у: {', '.join(blockers)}. Заархівуйте замість видалення.",
        )
    try:
        db.delete(p)
        db.commit()
    except Exception:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Неможливо видалити: товар має звʼязані записи. Заархівуйте замість видалення.",
        )


class StockThresholdsUpdate(BaseModel):
    min_stock:     int | None = None
    desired_stock: int | None = None
    cell_limit:    int | None = None


@router.patch("/products/{product_id}/thresholds", response_model=ProductOut)
def update_thresholds(
    product_id: int,
    payload:    StockThresholdsUpdate,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> ProductOut:
    p = _get_product(product_id, org, db)
    data = payload.model_dump(exclude_unset=True)
    if "cell_limit" in data:
        data["box_limit"] = data.pop("cell_limit")
    for k, v in data.items():
        setattr(p, k, v)
    db.commit()
    db.refresh(p)
    return ProductOut.model_validate(p)


# ── Specifications ────────────────────────────────────────────────────────────

@_full.get("/products/{product_id}/specs", response_model=list[SpecOut])
def list_specs(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[SpecOut]:
    _get_product(product_id, org, db)
    specs = db.query(Specification).filter(Specification.product_id == product_id).order_by(Specification.version).all()
    return [_spec_to_out(s, db, org.id) for s in specs]


@_full.post("/products/{product_id}/specs", response_model=SpecOut, status_code=status.HTTP_201_CREATED)
def create_spec(
    product_id: int,
    payload:    SpecCreate,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> SpecOut:
    _get_product(product_id, org, db)
    existing = db.query(Specification).filter(Specification.product_id == product_id).count()
    spec = Specification(product_id=product_id, name=payload.name, notes=payload.notes, version=existing + 1)
    if existing == 0:
        spec.is_default = True
    db.add(spec)
    db.commit()
    db.refresh(spec)
    return _spec_to_out(spec, db, org.id)


_SPEC_TSV_HEADER = "\t".join([
    "Назва виробу", "SKU виробу", "Од. вим. виробу",
    "Пряма собівартість виробу", "Повна собівартість виробу",
    "Назва матеріалу", "SKU матеріалу", "К-сть матеріалу", "Одиниця виміру матеріалу", "Сер.зважена ціна матеріалу",
    "Назва роботи", "К-сть роботи", "Одиниця виміру роботи", "Ціна роботи", "Додаткові витрати", "Вартість витрати",
])


@_full.get("/specs/export")
def export_ordage_specs(
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> Response:
    products = (
        db.query(Product)
        .filter(Product.organization_id == org.id, Product.is_active)
        .order_by(Product.name)
        .all()
    )
    rows = [_SPEC_TSV_HEADER]
    blank16 = [""] * 16

    for p in products:
        spec = db.query(Specification).filter_by(product_id=p.id, is_default=True).first()
        if not spec:
            continue
        components = (
            db.query(SpecComponent).filter_by(specification_id=spec.id)
            .order_by(SpecComponent.sort_order).all()
        )
        component_product_ids = {c.product_id for c in components if c.product_id is not None}
        component_skus: dict[int, str] = {}
        if component_product_ids:
            component_skus = {
                product_id: sku
                for product_id, sku in (
                    db.query(Product.id, Product.sku)
                    .filter(Product.organization_id == org.id, Product.id.in_(component_product_ids))
                    .all()
                )
            }
        operations = (
            db.query(SpecOperation).filter_by(specification_id=spec.id)
            .order_by(SpecOperation.sort_order).all()
        )

        row = blank16.copy()
        row[0]  = p.name
        row[1]  = p.sku
        row[2]  = p.unit
        row[3]  = str(p.direct_cost or "")
        row[4]  = str(p.full_cost   or "")
        rows.append("\t".join(row))

        for c in components:
            r = blank16.copy()
            r[1]  = p.sku
            r[5]  = c.name
            r[6]  = component_skus.get(c.product_id, "") if c.product_id else ""
            r[7]  = str(c.quantity)
            r[8]  = c.unit
            r[9]  = str(c.unit_price or "")
            rows.append("\t".join(r))

        for op in operations:
            total = op.explicit_cost or Decimal("0")
            if op.type.value == "print" and op.print_time_min:
                kwh    = Decimal(op.power_watts or _PRINTER_WATTS) * op.print_time_min / 60 / 1000
                total += kwh * _ELECTRICITY_RATE
            if op.labor_minutes:
                total += (op.labor_minutes / 60) * (op.labor_rate_per_hour or _LABOR_RATE)
            r = blank16.copy()
            r[1]  = p.sku
            r[10] = op.name
            r[11] = "1"
            r[13] = str(total.quantize(Decimal("0.0001")))
            rows.append("\t".join(r))

    content = "\n".join(rows).encode("utf-8-sig")
    return Response(
        content=content,
        media_type="text/tab-separated-values; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="specs.tsv"'},
    )


@_full.get("/specs/defaults", response_model=list[SpecOut])
def list_default_specs(
    db:    Session      = Depends(get_db),
    org:   Organization = Depends(get_current_org),
) -> list[SpecOut]:
    specs = (
        db.query(Specification)
        .join(Product, Product.id == Specification.product_id)
        .filter(
            Product.organization_id == org.id,
            Product.is_active.is_(True),
            Specification.is_default.is_(True),
        )
        .order_by(Product.name)
        .all()
    )
    return _specs_to_out(specs, db, org.id)


@_full.get("/specs/defaults/summary", response_model=list[SpecDefaultSummaryOut])
def list_default_specs_summary(
    db:    Session      = Depends(get_db),
    org:   Organization = Depends(get_current_org),
) -> list[SpecDefaultSummaryOut]:
    """Lightweight default-spec data for the specifications table."""
    specs = (
        db.query(Specification)
        .join(Product, Product.id == Specification.product_id)
        .filter(
            Product.organization_id == org.id,
            Product.is_active.is_(True),
            Specification.is_default.is_(True),
        )
        .order_by(Product.name)
        .all()
    )
    spec_ids = [s.id for s in specs]
    if not spec_ids:
        return []

    components_by_spec: dict[int, list[str]] = {sid: [] for sid in spec_ids}
    operations_by_spec: dict[int, list[SpecOperation]] = {sid: [] for sid in spec_ids}

    components = (
        db.query(SpecComponent)
        .filter(SpecComponent.specification_id.in_(spec_ids))
        .order_by(SpecComponent.specification_id, SpecComponent.sort_order)
        .all()
    )
    for component in components:
        components_by_spec.setdefault(component.specification_id, []).append(
            f"{component.name} {component.quantity} {component.unit}"
        )

    operations = (
        db.query(SpecOperation)
        .filter(SpecOperation.specification_id.in_(spec_ids))
        .order_by(SpecOperation.specification_id, SpecOperation.sort_order)
        .all()
    )
    for operation in operations:
        operations_by_spec.setdefault(operation.specification_id, []).append(operation)

    rows: list[SpecDefaultSummaryOut] = []
    for spec in specs:
        operations = operations_by_spec.get(spec.id, [])
        rows.append(SpecDefaultSummaryOut(
            product_id=spec.product_id,
            material_labels=components_by_spec.get(spec.id, []),
            work_labels=[op.name for op in operations],
            extra_labels=[
                f"{op.name}: {op.explicit_cost}"
                for op in operations
                if op.explicit_cost is not None and op.explicit_cost > 0
            ],
        ))
    return rows


def _parse_spec_rows(rows: list[list[str]]) -> list[dict]:
    """Shared parser for both TSV and XLSX rows (after normalisation to str lists)."""
    products: list[dict] = []
    current: dict | None = None

    for row in rows[1:]:          # skip header
        while len(row) < 16:
            row.append("")
        product_name = row[0].strip()
        sku          = row[1].strip()
        if not sku:
            continue
        if product_name:
            current = {
                "name": product_name,
                "sku": sku,
                "unit": row[2].strip(),
                "components": [],
                "operations": [],
            }
            products.append(current)
        elif current:
            mat_name = row[5].strip()
            op_name  = row[10].strip()
            if mat_name:
                qty   = row[7].strip()
                price = row[9].strip()
                current["components"].append({
                    "name":       mat_name,
                    "sku":        row[6].strip(),
                    "quantity":   Decimal(qty)   if qty   else Decimal("0"),
                    "unit":       row[8].strip(),
                    "unit_price": Decimal(price) if price else None,
                })
            elif op_name:
                qty   = row[11].strip()
                price = row[13].strip()
                q     = Decimal(qty)   if qty   else Decimal("0")
                p     = Decimal(price) if price else Decimal("0")
                current["operations"].append({
                    "name": op_name,
                    "cost": (q * p) if (q and p) else None,
                    "is_print": "друк" in op_name.lower() or "print" in op_name.lower(),
                })
    return products


def _parse_ordage_spec(content: bytes) -> list[dict]:
    """Parse Ordage spec TSV."""
    text   = content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text), delimiter="\t")
    return _parse_spec_rows(list(reader))


def _parse_ordage_xlsx(content: bytes) -> list[dict]:
    """Parse Ordage spec XLSX."""
    wb   = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws   = wb.active
    rows = [
        [str(c).strip() if c is not None else "" for c in row]
        for row in ws.iter_rows(values_only=True)
    ]
    wb.close()
    return _parse_spec_rows(rows)


class OrdageSpecImportResult(BaseModel):
    updated: int
    skipped: int
    errors:  list[dict]


@_full.post("/specs/import", response_model=OrdageSpecImportResult)
def import_ordage_specs(
    file: UploadFile      = File(...),
    db:   Session         = Depends(get_db),
    org:  Organization    = Depends(get_current_org),
    user: User            = Depends(require_roles(UserRole.admin)),
) -> OrdageSpecImportResult:
    content  = file.file.read()
    filename = (file.filename or "").lower()
    parsed   = _parse_ordage_xlsx(content) if filename.endswith(".xlsx") else _parse_ordage_spec(content)
    updated, skipped = 0, 0
    errors: list[dict] = []

    org_products = db.query(Product).filter(Product.organization_id == org.id).all()
    product_by_sku = {_norm_lookup(p.sku): p for p in org_products if p.sku}
    product_by_name = {_norm_lookup(p.name): p for p in org_products if p.name}

    for item in parsed:
        product = product_by_sku.get(_norm_lookup(item["sku"]))
        if not product:
            product = Product(
                organization_id=org.id,
                name=item["name"] or item["sku"],
                sku=item["sku"],
                unit=item["unit"] or "шт",
                created_by_id=user.id,
            )
            db.add(product)
            db.flush()
            product_by_sku[_norm_lookup(product.sku)] = product
            product_by_name[_norm_lookup(product.name)] = product

        spec = db.query(Specification).filter_by(product_id=product.id, is_default=True).first()
        if not spec:
            ver  = db.query(Specification).filter_by(product_id=product.id).count()
            spec = Specification(product_id=product.id, name="Основна", is_default=True, version=ver + 1)
            db.add(spec)
            db.flush()

        db.query(SpecComponent).filter_by(specification_id=spec.id).delete()
        for i, c in enumerate(item["components"]):
            component_product = _ensure_component_product(
                db=db,
                org=org,
                user=user,
                name=c["name"],
                sku=c.get("sku"),
                unit=c["unit"],
                unit_price=c["unit_price"],
                product_by_sku=product_by_sku,
                product_by_name=product_by_name,
            )
            component_name = component_product.name
            component_unit = c["unit"] or (component_product.unit if component_product else "шт")
            component_price = c["unit_price"]
            if component_price is None and component_product and component_product.cost_price is not None:
                component_price = component_product.cost_price

            db.add(SpecComponent(
                specification_id=spec.id,
                product_id=component_product.id,
                name=component_name, quantity=c["quantity"],
                unit=component_unit, unit_price=component_price,
                waste_pct=Decimal("0"), sort_order=i,
            ))

        db.query(SpecOperation).filter_by(specification_id=spec.id).delete()
        for i, op in enumerate(item["operations"]):
            n = op["name"].lower()
            t = (SpecOpType.print       if ("друк" in n or "print" in n) else
                 SpecOpType.postprocess if ("постобр" in n or "post" in n) else
                 SpecOpType.manual)
            db.add(SpecOperation(
                specification_id=spec.id,
                type=t, name=op["name"], sort_order=i,
                explicit_cost=op["cost"],
            ))

        db.flush()
        cost = _calc_cost(spec, db, electricity_rate=org.electricity_rate or _ELECTRICITY_RATE, labor_rate=org.labor_rate or _LABOR_RATE)
        product.direct_cost = cost.material_cost + cost.electricity_cost
        product.full_cost   = cost.total
        db.commit()
        updated += 1

    return OrdageSpecImportResult(updated=updated, skipped=skipped, errors=errors)


@_full.get("/specs/{spec_id}", response_model=SpecOut)
def get_spec(
    spec_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> SpecOut:
    return _spec_to_out(_get_spec(spec_id, org, db), db, org.id)


@_full.post("/specs/{spec_id}/set-default", response_model=SpecOut)
def set_default_spec(
    spec_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> SpecOut:
    spec = _get_spec(spec_id, org, db)
    db.query(Specification).filter(Specification.product_id == spec.product_id).update({"is_default": False})
    spec.is_default = True
    db.commit()
    db.refresh(spec)
    return _spec_to_out(spec, db, org.id)


@_full.post("/specs/{spec_id}/components", response_model=SpecOut, status_code=status.HTTP_201_CREATED)
def add_component(
    spec_id: int,
    payload: SpecComponentCreate,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    user: User        = Depends(require_roles(UserRole.admin)),
) -> SpecOut:
    spec = _get_spec(spec_id, org, db)
    if payload.product_id is not None:
        product = _get_product(payload.product_id, org, db)
    else:
        product = _ensure_component_product(
            db=db,
            org=org,
            user=user,
            name=payload.name,
            sku=None,
            unit=payload.unit,
            unit_price=payload.unit_price,
        )
    data = payload.model_dump()
    data["name"] = data["name"].strip()
    data["product_id"] = product.id
    unit = (data["unit"] or "").strip()
    data["name"] = product.name
    data["unit"] = unit or product.unit
    if data["unit_price"] is None:
        data["unit_price"] = product.cost_price
    comp = SpecComponent(specification_id=spec.id, **data)
    db.add(comp)
    db.commit()
    return _spec_to_out(spec, db, org.id)


@_full.delete("/specs/{spec_id}/components/{comp_id}", status_code=status.HTTP_204_NO_CONTENT)
def remove_component(
    spec_id: int,
    comp_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> None:
    _get_spec(spec_id, org, db)
    comp = db.query(SpecComponent).filter_by(id=comp_id, specification_id=spec_id).first()
    if not comp:
        raise HTTPException(status_code=404, detail="Component not found")
    db.delete(comp)
    db.commit()


@_full.post("/specs/{spec_id}/operations", response_model=SpecOut, status_code=status.HTTP_201_CREATED)
def add_operation(
    spec_id: int,
    payload: SpecOperationCreate,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> SpecOut:
    spec = _get_spec(spec_id, org, db)
    op = SpecOperation(specification_id=spec.id, **payload.model_dump())
    db.add(op)
    db.commit()
    return _spec_to_out(spec, db, org.id)


@_full.delete("/specs/{spec_id}/operations/{op_id}", status_code=status.HTTP_204_NO_CONTENT)
def remove_operation(
    spec_id: int,
    op_id:   int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> None:
    _get_spec(spec_id, org, db)
    op = db.query(SpecOperation).filter_by(id=op_id, specification_id=spec_id).first()
    if not op:
        raise HTTPException(status_code=404, detail="Operation not found")
    db.delete(op)
    db.commit()


@_full.get("/products/{product_id}/cost", response_model=CostBreakdown)
def compute_cost(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> CostBreakdown:
    _get_product(product_id, org, db)
    spec = (
        db.query(Specification)
        .filter(Specification.product_id == product_id, Specification.is_default)
        .first()
    )
    if not spec:
        raise HTTPException(status_code=404, detail="No default specification")
    cost = _calc_cost(spec, db, electricity_rate=org.electricity_rate or _ELECTRICITY_RATE, labor_rate=org.labor_rate or _LABOR_RATE)
    p = db.get(Product, product_id)
    if p:
        p.direct_cost = cost.material_cost + cost.electricity_cost
        p.full_cost   = cost.total
        db.commit()
    return cost


# ── Stock ─────────────────────────────────────────────────────────────────────

@_full.get("/dashboard", response_model=DashboardSummaryOut)
def dashboard_summary(
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> DashboardSummaryOut:
    """Lightweight KPI block for the warehouse overview — aggregates only,
    no per-cell joins or presigned image URLs (unlike /stock)."""
    sku_count = db.query(func.count(func.distinct(StockEntry.product_id))).filter(
        StockEntry.organization_id == org.id
    ).scalar() or 0

    total_units = db.query(func.sum(StockEntry.quantity)).filter(
        StockEntry.organization_id == org.id
    ).scalar() or Decimal("0")

    low_rows = (
        db.query(StockEntry, Product)
        .join(Product, Product.id == StockEntry.product_id)
        .join(Warehouse, Warehouse.id == StockEntry.warehouse_id)
        .filter(
            StockEntry.organization_id == org.id,
            Warehouse.type == WarehouseType.finished,
            (StockEntry.quantity - StockEntry.reserved_qty) < _LOW_STOCK_THRESHOLD,
        )
        .order_by(Product.name)
        .all()
    )
    low_stock = [
        DashboardLowStockOut(
            product_id=e.product_id,
            product_name=p.name,
            available=e.quantity - e.reserved_qty,
        )
        for e, p in low_rows
    ]
    return DashboardSummaryOut(sku_count=sku_count, total_units=total_units, low_stock=low_stock)


@_full.get("/stock", response_model=list[StockEntryOut])
def list_stock(
    warehouse_id: int | None = Query(None),
    product_id:   int | None = Query(None),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[StockEntryOut]:
    q = (
        db.query(StockEntry, Product, Warehouse)
        .join(Product,   Product.id   == StockEntry.product_id)
        .join(Warehouse, Warehouse.id == StockEntry.warehouse_id)
        .filter(StockEntry.organization_id == org.id)
    )
    if warehouse_id:
        q = q.filter(StockEntry.warehouse_id == warehouse_id)
    if product_id:
        q = q.filter(StockEntry.product_id == product_id)
    rows = q.order_by(Product.name).all()

    import collections
    from sqlalchemy import func
    from app.models.warehouse import CellStock, WarehouseCell, WarehouseZone

    product_ids = {p.id for _, p, _ in rows}
    warehouse_ids = {wh.id for _, _, wh in rows}

    total_stock_map = {}
    if product_ids:
        totals = db.query(
            StockEntry.product_id,
            func.sum(StockEntry.quantity).label("total")
        ).filter(
            StockEntry.organization_id == org.id,
            StockEntry.product_id.in_(product_ids)
        ).group_by(StockEntry.product_id).all()
        total_stock_map = {pid: tot for pid, tot in totals}

    cell_stock_map = collections.defaultdict(list)
    if product_ids and warehouse_ids:
        cell_stocks = db.query(
            CellStock.product_id,
            WarehouseZone.warehouse_id,
            WarehouseZone.name.label("zone_name"),
            WarehouseCell.code,
            CellStock.quantity
        ).join(
            WarehouseCell, WarehouseCell.id == CellStock.cell_id
        ).join(
            WarehouseZone, WarehouseZone.id == WarehouseCell.zone_id
        ).filter(
            WarehouseZone.organization_id == org.id,
            CellStock.product_id.in_(product_ids),
            WarehouseZone.warehouse_id.in_(warehouse_ids),
            CellStock.quantity > 0
        ).all()
        for pid, wid, zone_name, cell_code, qty in cell_stocks:
            cell_name = f"{zone_name} {cell_code}"
            cell_stock_map[(pid, wid)].append({"name": cell_name, "quantity": qty})

    # Скільки кожного товару зараз у відкритих виробничих партіях (заплановано/друкується/пауза).
    # Прив'язка партій — на рівні товару, тож значення проставляється всім складським рядкам товару.
    in_production_map: dict[int, int] = {}
    if product_ids:
        open_states = (BatchStatus.draft, BatchStatus.active, BatchStatus.paused)
        prod_rows = db.query(
            ProductionBatch.product_id,
            func.sum(ProductionBatch.target_qty - ProductionBatch.good_qty),
        ).filter(
            ProductionBatch.organization_id == org.id,
            ProductionBatch.product_id.in_(product_ids),
            ProductionBatch.status.in_(open_states),
        ).group_by(ProductionBatch.product_id).all()
        in_production_map = {pid: max(0, int(s or 0)) for pid, s in prod_rows}

    result = []
    for e, p, wh in rows:
        avail = e.quantity - e.reserved_qty
        total_stock = total_stock_map.get(e.product_id, Decimal(0))
        locations = cell_stock_map.get((e.product_id, e.warehouse_id), [])
        assigned = sum((Decimal(str(loc["quantity"])) for loc in locations), Decimal(0))

        result.append(StockEntryOut(
            id=e.id,
            product_id=e.product_id,
            product_name=p.name,
            product_sku=p.sku,
            product_barcode=p.barcode,
            product_categories=p.categories or [],
            image_url=_product_image_url(p, org.id),
            product_unit=p.unit,
            warehouse_id=e.warehouse_id,
            warehouse_name=wh.name,
            locations=locations,
            assigned_qty=assigned,
            unassigned_qty=max(Decimal(0), e.quantity - assigned),
            quantity=e.quantity,
            reserved_qty=e.reserved_qty,
            available=avail,
            total_stock=total_stock,
            full_cost=p.full_cost,
            min_stock=p.min_stock,
            desired_stock=p.desired_stock,
            cell_limit=p.box_limit,
            in_production_qty=in_production_map.get(e.product_id, 0),
            updated_at=e.updated_at,
        ))
    return result


@_full.get("/stock/summary", response_model=list[StockSummaryOut])
def stock_summary(
    warehouse_id: int | None = Query(None),
    product_id:   int | None = Query(None),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[StockSummaryOut]:
    """Lightweight stock totals for list pages.

    Use this when the UI only needs quantities. The full /stock endpoint also
    resolves product data, images, cell locations, and production counters.
    """
    q = db.query(StockEntry).filter(StockEntry.organization_id == org.id)
    if warehouse_id:
        q = q.filter(StockEntry.warehouse_id == warehouse_id)
    if product_id:
        q = q.filter(StockEntry.product_id == product_id)

    rows = q.order_by(StockEntry.warehouse_id, StockEntry.product_id).all()
    return [
        StockSummaryOut(
            product_id=e.product_id,
            warehouse_id=e.warehouse_id,
            quantity=e.quantity,
            reserved_qty=e.reserved_qty,
            available=e.quantity - e.reserved_qty,
        )
        for e in rows
    ]


# ── Replenishment ─────────────────────────────────────────────────────────────

class _ReplenishPreviewItem(BaseModel):
    product_id:       int
    product_name:     str
    product_sku:      str
    unit:             str
    available:        float
    min_stock:        int | None
    desired_stock:    int | None
    qty_needed:       int
    kind:             str         # "batch" | "purchase"
    specification_id: int | None
    warehouse_id:     int | None
    warehouse_name:   str | None


class _ReplenishItem(BaseModel):
    product_id:       int
    qty:              int
    kind:             str         # "batch" | "purchase"
    warehouse_id:     int | None = None
    specification_id: int | None = None
    unit_cost:        Decimal | None = None


class _ReplenishRequest(BaseModel):
    items: list[_ReplenishItem]


@_full.get("/stock/replenish-preview", response_model=list[_ReplenishPreviewItem])
def replenish_preview(
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[_ReplenishPreviewItem]:
    available_expr = StockEntry.quantity - StockEntry.reserved_qty
    rows = (
        db.query(StockEntry, Product, Warehouse)
        .join(Product,   Product.id   == StockEntry.product_id)
        .join(Warehouse, Warehouse.id == StockEntry.warehouse_id)
        .filter(
            StockEntry.organization_id == org.id,
            or_(
                available_expr <= 0,
                Product.min_stock.isnot(None) & (available_expr < Product.min_stock),
            ),
        )
        .order_by(Product.name)
        .all()
    )
    if not rows:
        return []

    product_ids = list({p.id for _, p, _ in rows})
    default_specs = {
        s.product_id: s
        for s in db.query(Specification)
        .join(Product, Product.id == Specification.product_id)
        .filter(
            Product.organization_id == org.id,
            Specification.product_id.in_(product_ids),
            Specification.is_default.is_(True),
        )
        .all()
    }

    seen: set[int] = set()
    result: list[_ReplenishPreviewItem] = []
    for e, p, wh in rows:
        if p.id in seen:
            continue
        seen.add(p.id)
        avail = float(e.quantity - e.reserved_qty)
        target = p.desired_stock if p.desired_stock else (p.min_stock or 1)
        qty_needed = max(1, int(target - avail))
        spec = default_specs.get(p.id)
        result.append(_ReplenishPreviewItem(
            product_id=p.id,
            product_name=p.name,
            product_sku=p.sku,
            unit=p.unit,
            available=avail,
            min_stock=p.min_stock,
            desired_stock=p.desired_stock,
            qty_needed=qty_needed,
            kind="batch",
            specification_id=spec.id if spec else None,
            warehouse_id=wh.id,
            warehouse_name=wh.name,
        ))
    return result


@_full.post("/stock/replenish")
def replenish_stock(
    payload: _ReplenishRequest,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> dict:
    batches_created = 0
    movements_created = 0
    for item in payload.items:
        if item.qty <= 0:
            continue
        product = _get_product(item.product_id, org, db)
        if item.kind == "batch":
            _get_spec_for_product(item.specification_id, item.product_id, org, db)
            b = ProductionBatch(
                organization_id=org.id,
                product_id=item.product_id,
                specification_id=item.specification_id,
                target_qty=item.qty,
                status=BatchStatus.draft,
                created_by_id=user.id,
            )
            db.add(b)
            batches_created += 1
        elif item.kind == "purchase":
            if item.warehouse_id is None:
                raise HTTPException(status_code=400, detail="Purchase replenishment requires warehouse_id")
            _get_warehouse(item.warehouse_id, org, db)
            qty = Decimal(str(item.qty))
            uc  = item.unit_cost
            m = WarehouseMovement(
                organization_id=org.id,
                product_id=item.product_id,
                type=MovementType.PURCHASE_IN,
                warehouse_to_id=item.warehouse_id,
                quantity=qty,
                unit=product.unit,
                unit_cost=uc,
                total_cost=(uc * qty) if uc else None,
                created_by_id=user.id,
            )
            db.add(m)
            db.flush()
            _apply_movement(m, db)
            movements_created += 1
        else:
            raise HTTPException(status_code=400, detail="Unknown replenishment kind")
    db.commit()
    return {"batches": batches_created, "movements": movements_created}


# ── Stock export / import ──────────────────────────────────────────────────────

_STOCK_HEADERS = ("SKU", "Назва", "Склад", "В наявності", "Одиниця")


@_full.get("/stock/export")
def export_stock(
    fmt: str        = Query("xlsx"),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> Response:
    from urllib.parse import quote
    rows = (
        db.query(StockEntry, Product, Warehouse)
        .join(Product,   StockEntry.product_id   == Product.id)
        .join(Warehouse, StockEntry.warehouse_id  == Warehouse.id)
        .filter(StockEntry.organization_id == org.id)
        .order_by(Product.name)
        .all()
    )
    ts = datetime.utcnow().strftime("%Y-%m-%d_%H-%M")

    if fmt == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Залишки"
        ws.append(list(_STOCK_HEADERS))
        for se, p, wh in rows:
            ws.append([p.sku, p.name, wh.name, float(se.quantity), p.unit or "шт"])
        buf = io.BytesIO()
        wb.save(buf)
        filename = f"залишки_{ts}.xlsx"
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
        )

    buf = io.StringIO()
    buf.write("﻿")
    writer = csv.writer(buf, delimiter="\t", lineterminator="\r\n")
    writer.writerow(_STOCK_HEADERS)
    for se, p, wh in rows:
        writer.writerow([p.sku, p.name, wh.name, float(se.quantity), p.unit or "шт"])
    filename = f"залишки_{ts}.tsv"
    return Response(
        content=buf.getvalue().encode("utf-8"),
        media_type="text/tab-separated-values; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


class _StockImportResult(BaseModel):
    updated: int
    skipped: int
    errors:  list[str]


@_full.post("/stock/import", response_model=_StockImportResult)
def import_stock(
    file: UploadFile   = File(...),
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> _StockImportResult:
    """Stocktake import — reads SKU / Warehouse / Quantity and adjusts stock to match."""
    content = file.file.read()
    try:
        if file.filename and file.filename.endswith(".xlsx"):
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            ws = wb.active
            raw_rows = [
                [str(c.value).strip() if c.value is not None else "" for c in row]
                for row in ws.iter_rows(min_row=2)
            ]
        else:
            text = content.decode("utf-8-sig").replace("\r\n", "\n")
            delimiter = "\t" if "\t" in text else ","
            raw_rows = [r for r in csv.reader(text.splitlines(), delimiter=delimiter)][1:]
    except Exception as e:
        raise HTTPException(400, detail=f"Помилка читання файлу: {e}")

    wh_cache:  dict[str, Warehouse] = {}
    sku_cache: dict[str, Product]   = {}
    updated = skipped = 0
    errors: list[str] = []

    for i, row in enumerate(raw_rows, start=2):
        if not any(row):
            continue
        if len(row) < 4:
            errors.append(f"Рядок {i}: замало колонок (потрібно SKU, Назва, Склад, Кількість)")
            skipped += 1
            continue

        sku, wh_name, qty_raw = row[0].strip(), row[2].strip(), row[3].strip()

        # resolve SKU
        if sku not in sku_cache:
            p = db.query(Product).filter_by(sku=sku, organization_id=org.id).first()
            if not p:
                errors.append(f"Рядок {i}: SKU «{sku}» не знайдено")
                skipped += 1
                continue
            sku_cache[sku] = p
        product = sku_cache[sku]

        # resolve warehouse
        if wh_name not in wh_cache:
            wh = db.query(Warehouse).filter(
                Warehouse.organization_id == org.id,
                func.lower(Warehouse.name) == wh_name.lower(),
            ).first()
            if not wh:
                errors.append(f"Рядок {i}: склад «{wh_name}» не знайдено")
                skipped += 1
                continue
            wh_cache[wh_name] = wh
        warehouse = wh_cache[wh_name]

        try:
            target = Decimal(qty_raw.replace(",", "."))
        except Exception:
            errors.append(f"Рядок {i}: невірна кількість «{qty_raw}»")
            skipped += 1
            continue

        entry = db.query(StockEntry).filter_by(
            organization_id=org.id,
            product_id=product.id,
            warehouse_id=warehouse.id,
        ).first()
        current = entry.quantity if entry else Decimal("0")
        delta = target - current
        if delta == 0:
            skipped += 1
            continue

        if delta > 0:
            m = WarehouseMovement(
                organization_id=org.id, type=MovementType.ADJUSTMENT,
                product_id=product.id, warehouse_to_id=warehouse.id,
                quantity=delta, reason="Імпорт залишків",
                created_by_id=user.id,
            )
        else:
            m = WarehouseMovement(
                organization_id=org.id, type=MovementType.ADJUSTMENT,
                product_id=product.id, warehouse_from_id=warehouse.id,
                quantity=-delta, reason="Імпорт залишків",
                created_by_id=user.id,
            )
        db.add(m)
        db.flush()
        _apply_movement(m, db)
        updated += 1

    db.commit()
    return _StockImportResult(updated=updated, skipped=skipped, errors=errors[:50])


# ── Movements ─────────────────────────────────────────────────────────────────

def _encode_cursor(created_at: datetime, row_id: int) -> str:
    raw = f"{created_at.isoformat()}|{row_id}"
    return base64.urlsafe_b64encode(raw.encode()).decode()


def _decode_cursor(cursor: str) -> tuple[datetime, int]:
    raw = base64.urlsafe_b64decode(cursor.encode()).decode()
    created_iso, id_str = raw.rsplit("|", 1)
    return datetime.fromisoformat(created_iso), int(id_str)


@_full.get("/movements", response_model=MovementListOut)
def list_movements(
    movement_type: MovementType | None = Query(None),
    product_id:    int | None          = Query(None),
    batch_id:      int | None          = Query(None),
    order_id:      int | None          = Query(None),
    cursor:        str | None          = Query(None),
    limit:         int                 = Query(50, ge=1, le=200),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> MovementListOut:
    from sqlalchemy import and_, or_, select, func as safunc

    def _base_filter(q):
        q = q.where(WarehouseMovement.organization_id == org.id)
        if movement_type:
            q = q.where(WarehouseMovement.type == movement_type)
        if product_id:
            q = q.where(WarehouseMovement.product_id == product_id)
        if batch_id:
            q = q.where(WarehouseMovement.batch_id == batch_id)
        if order_id:
            q = q.where(WarehouseMovement.order_id == order_id)
        return q

    # total count — only on first page, with same filters (no cursor)
    total: int | None = None
    if cursor is None:
        count_q = _base_filter(select(safunc.count(WarehouseMovement.id)))
        total = db.execute(count_q).scalar_one()

    # data query with keyset cursor
    data_q = _base_filter(select(WarehouseMovement))
    if cursor:
        try:
            c_at, c_id = _decode_cursor(cursor)
            data_q = data_q.where(
                or_(
                    WarehouseMovement.created_at < c_at,
                    and_(
                        WarehouseMovement.created_at == c_at,
                        WarehouseMovement.id < c_id,
                    ),
                )
            )
        except Exception:
            pass  # bad cursor → ignore, return from start

    data_q = data_q.order_by(
        WarehouseMovement.created_at.desc(),
        WarehouseMovement.id.desc(),
    ).limit(limit + 1)

    rows = db.execute(data_q).scalars().all()
    has_more = len(rows) > limit
    rows = list(rows[:limit])

    next_cursor = _encode_cursor(rows[-1].created_at, rows[-1].id) if has_more else None

    # build product name map in one query
    pids = {r.product_id for r in rows}
    products = {
        p.id: p.name
        for p in db.query(Product).filter(Product.organization_id == org.id, Product.id.in_(pids)).all()
    } if pids else {}

    items = [
        MovementOut(
            id=m.id, type=m.type,
            direction=_MOVEMENT_DIRECTION.get(m.type, "in"),
            product_id=m.product_id, product_name=products.get(m.product_id, ""),
            warehouse_from_id=m.warehouse_from_id, warehouse_to_id=m.warehouse_to_id,
            quantity=m.quantity, unit=m.unit,
            unit_cost=m.unit_cost, total_cost=m.total_cost,
            unit_price=m.unit_price, total_revenue=m.total_revenue,
            reason=m.reason, batch_id=m.batch_id, order_id=m.order_id,
            created_at=m.created_at,
        )
        for m in rows
    ]

    return MovementListOut(items=items, next_cursor=next_cursor, has_more=has_more, total=total)


@_full.post("/movements", response_model=MovementOut, status_code=status.HTTP_201_CREATED)
def create_movement(
    payload: MovementCreate,
    bg:   BackgroundTasks,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> MovementOut:
    product = _get_product(payload.product_id, org, db)
    if payload.quantity <= 0:
        raise HTTPException(status_code=400, detail="Quantity must be greater than zero")
    if payload.warehouse_from_id is not None:
        _get_warehouse(payload.warehouse_from_id, org, db)
    if payload.warehouse_to_id is not None:
        _get_warehouse(payload.warehouse_to_id, org, db)
    if payload.batch_id is not None:
        batch = db.query(ProductionBatch).filter_by(id=payload.batch_id, organization_id=org.id).first()
        if not batch:
            raise HTTPException(status_code=404, detail="Batch not found")
        if batch.product_id != payload.product_id and payload.type in (MovementType.PRODUCTION_IN, MovementType.DEFECT):
            raise HTTPException(status_code=400, detail="Movement product does not match batch product")
    if payload.order_id is not None:
        order = db.query(Order).filter_by(id=payload.order_id, organization_id=org.id).first()
        if not order:
            raise HTTPException(status_code=404, detail="Order not found")

    if payload.type in (MovementType.PURCHASE_IN, MovementType.RETURN_IN, MovementType.PRODUCTION_IN) and payload.warehouse_to_id is None:
        raise HTTPException(status_code=400, detail=f"{payload.type.value} requires warehouse_to_id")
    if payload.type in (MovementType.SALE_OUT, MovementType.PRODUCTION_OUT, MovementType.WRITE_OFF) and payload.warehouse_from_id is None:
        raise HTTPException(status_code=400, detail=f"{payload.type.value} requires warehouse_from_id")
    if payload.type == MovementType.TRANSFER:
        if payload.warehouse_from_id is None or payload.warehouse_to_id is None:
            raise HTTPException(status_code=400, detail="TRANSFER requires both warehouses")
        if payload.warehouse_from_id == payload.warehouse_to_id:
            raise HTTPException(status_code=400, detail="Transfer warehouses must differ")
    if payload.type == MovementType.DEFECT and payload.warehouse_from_id is None and payload.warehouse_to_id is None:
        raise HTTPException(status_code=400, detail="DEFECT requires at least one warehouse")
    if payload.type == MovementType.ADJUSTMENT:
        if bool(payload.warehouse_from_id) == bool(payload.warehouse_to_id):
            raise HTTPException(status_code=400, detail="ADJUSTMENT requires exactly one warehouse")

    if payload.type == MovementType.PURCHASE_IN and not payload.unit_cost:
        raise HTTPException(status_code=400, detail="Закупка потребує ціну за одиницю (unit_cost)")
    if payload.type == MovementType.RETURN_IN and payload.order_id:
        item = db.query(OrderItem).filter_by(
            order_id=payload.order_id, product_id=payload.product_id
        ).first()
        if not item:
            raise HTTPException(status_code=400, detail="Товар не входив у це замовлення")
        already_returned: Decimal = db.query(
            func.coalesce(func.sum(WarehouseMovement.quantity), 0)
        ).filter(
            WarehouseMovement.organization_id == org.id,
            WarehouseMovement.order_id == payload.order_id,
            WarehouseMovement.product_id == payload.product_id,
            WarehouseMovement.type == MovementType.RETURN_IN,
        ).scalar() or Decimal("0")
        if already_returned + payload.quantity > item.quantity:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Не можна повернути більше, ніж відвантажено "
                    f"(відвантажено {item.quantity}, вже повернуто {already_returned})"
                ),
            )
    total = (payload.quantity * payload.unit_cost) if payload.unit_cost else None
    data = payload.model_dump()
    cell_from_id = data.pop("cell_from_id", None)   # not WarehouseMovement columns
    cell_to_id   = data.pop("cell_to_id", None)
    m = WarehouseMovement(
        organization_id=org.id,
        created_by_id=user.id,
        total_cost=total,
        **data,
    )
    db.add(m)
    db.flush()

    # Honor an explicit source bin BEFORE the stock change so the FIFO clamp
    # inside _apply_movement leaves it alone (the cell is already drawn down).
    if cell_from_id and m.warehouse_from_id:
        src = _get_cell(cell_from_id, org, db, warehouse_id=m.warehouse_from_id)
        _pick_from_cell(src, m.product_id, m.quantity, org.id, db,
                        movement_id=m.id, created_by_id=user.id)

    _apply_movement(m, db)

    # Honor an explicit target bin AFTER stock increased (cap to what's unassigned).
    if cell_to_id and m.warehouse_to_id:
        dst = _get_cell(cell_to_id, org, db, warehouse_id=m.warehouse_to_id)
        avail = _unassigned_qty(m.product_id, m.warehouse_to_id, org.id, db)
        _putaway(dst, m.product_id, min(m.quantity, avail), org.id, db,
                 movement_id=m.id, created_by_id=user.id)

    if payload.type in (MovementType.SALE_OUT, MovementType.PRODUCTION_OUT, MovementType.DEFECT, MovementType.TRANSFER):
        _check_and_auto_replenish(m.product_id, org.id, db)

    db.commit()
    db.refresh(m)
    bg.add_task(broadcast_warehouse, org.id, "movements")
    bg.add_task(broadcast_warehouse, org.id, "stock")
    return MovementOut(
        id=m.id, type=m.type,
        direction=_MOVEMENT_DIRECTION.get(m.type, "in"),
        product_id=m.product_id, product_name=product.name,
        warehouse_from_id=m.warehouse_from_id, warehouse_to_id=m.warehouse_to_id,
        quantity=m.quantity, unit=m.unit,
        unit_cost=m.unit_cost, total_cost=m.total_cost,
        reason=m.reason, batch_id=m.batch_id, order_id=m.order_id,
        created_at=m.created_at,
    )


# ── ProductionBatch ───────────────────────────────────────────────────────────

@dataclass
class _BatchPrefetch:
    """Bulk-loaded lookups for serializing many batches without N+1 queries."""
    products:         dict[int, Product]
    comps_by_spec:    dict[int, list[SpecComponent]]
    avail_by_product: dict[int, Decimal]
    task_titles:      dict[int, str]
    user_names:       dict[int, str]


def _build_batch_prefetch(batches: list[ProductionBatch], db: Session) -> _BatchPrefetch:
    org_id = batches[0].organization_id if batches else None
    spec_ids = {b.specification_id for b in batches if b.specification_id}
    comps_by_spec: dict[int, list[SpecComponent]] = {}
    if spec_ids:
        comp_rows = (
            db.query(SpecComponent)
            .filter(SpecComponent.specification_id.in_(spec_ids))
            .order_by(SpecComponent.sort_order)
            .all()
        )
        for c in comp_rows:
            comps_by_spec.setdefault(c.specification_id, []).append(c)

    comp_product_ids = {c.product_id for comps in comps_by_spec.values() for c in comps if c.product_id}
    product_ids = {b.product_id for b in batches} | comp_product_ids
    products = {
        p.id: p
        for p in db.query(Product).filter(
            Product.organization_id == org_id,
            Product.id.in_(product_ids),
        ).all()
    } if product_ids else {}

    # Sum per-entry available (clamped at 0) to match the single-batch path exactly.
    avail_by_product: dict[int, Decimal] = {}
    if comp_product_ids:
        for e in db.query(StockEntry).filter(
            StockEntry.organization_id == org_id,
            StockEntry.product_id.in_(comp_product_ids),
        ).all():
            avail_by_product[e.product_id] = (
                avail_by_product.get(e.product_id, Decimal("0"))
                + max(Decimal("0"), e.quantity - e.reserved_qty)
            )

    task_titles: dict[int, str] = {}
    task_ids = {b.print_task_id for b in batches if b.print_task_id}
    if task_ids:
        from app.models.task import PrintTask
        for t in db.query(PrintTask).filter(PrintTask.organization_id == org_id, PrintTask.id.in_(task_ids)).all():
            task_titles[t.id] = t.title

    user_names: dict[int, str] = {}
    user_ids = {b.assigned_to_id for b in batches if b.assigned_to_id}
    if user_ids:
        for u in db.query(User).filter(User.organization_id == org_id, User.id.in_(user_ids)).all():
            user_names[u.id] = u.name or u.email

    return _BatchPrefetch(
        products=products,
        comps_by_spec=comps_by_spec,
        avail_by_product=avail_by_product,
        task_titles=task_titles,
        user_names=user_names,
    )


def _batch_to_out(b: ProductionBatch, db: Session, pf: "_BatchPrefetch | None" = None) -> BatchOut:
    p = pf.products.get(b.product_id) if pf else db.get(Product, b.product_id)

    components: list[BatchComponentOut] = []
    if b.specification_id:
        if pf is not None:
            spec_comps = pf.comps_by_spec.get(b.specification_id, [])
        else:
            spec_comps = (
                db.query(SpecComponent)
                .filter_by(specification_id=b.specification_id)
                .order_by(SpecComponent.sort_order)
                .all()
            )
        for c in spec_comps:
            total_qty = c.quantity * b.target_qty
            available_stock: Decimal | None = None
            product_name: str | None = None
            if c.product_id:
                if pf is not None:
                    cp = pf.products.get(c.product_id)
                    available_stock = pf.avail_by_product.get(c.product_id, Decimal("0"))
                else:
                    cp = db.query(Product).filter_by(id=c.product_id, organization_id=b.organization_id).first()
                    entries = db.query(StockEntry).filter_by(
                        organization_id=b.organization_id,
                        product_id=c.product_id,
                    ).all()
                    available_stock = sum(
                        (max(Decimal("0"), e.quantity - e.reserved_qty) for e in entries),
                        Decimal("0"),
                    )
                product_name = cp.name if cp else None
            components.append(BatchComponentOut(
                id=c.id,
                name=c.name,
                product_id=c.product_id,
                product_name=product_name,
                quantity=c.quantity,
                unit=c.unit,
                total_qty=total_qty,
                available_stock=available_stock,
                is_sufficient=available_stock is None or available_stock >= total_qty,
            ))

    print_task_title: str | None = None
    if b.print_task_id:
        if pf is not None:
            print_task_title = pf.task_titles.get(b.print_task_id)
        else:
            from app.models.task import PrintTask
            pt = db.get(PrintTask, b.print_task_id)
            print_task_title = pt.title if pt else None

    assigned_to_name: str | None = None
    if b.assigned_to_id:
        if pf is not None:
            assigned_to_name = pf.user_names.get(b.assigned_to_id)
        else:
            u = db.get(User, b.assigned_to_id)
            assigned_to_name = (u.name or u.email) if u else None

    return BatchOut(
        id=b.id, product_id=b.product_id,
        product_name=p.name if p else "",  # type: ignore[union-attr]
        specification_id=b.specification_id,
        target_qty=b.target_qty, printed_qty=b.printed_qty,
        good_qty=b.good_qty, defect_qty=b.defect_qty,
        status=b.status, priority=b.priority, due_date=b.due_date,
        order_id=b.order_id,
        print_task_id=b.print_task_id,
        print_task_title=print_task_title,
        notes=b.notes,
        components=components,
        assigned_to_id=b.assigned_to_id,
        assigned_to_name=assigned_to_name,
        created_at=b.created_at, updated_at=b.updated_at,
    )


@_full.get("/batches", response_model=list[BatchOut])
def list_batches(
    batch_status: BatchStatus | None = Query(None),
    product_id:   int | None         = Query(None),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[BatchOut]:
    q = db.query(ProductionBatch).filter(ProductionBatch.organization_id == org.id)
    if batch_status:
        q = q.filter(ProductionBatch.status == batch_status)
    if product_id:
        q = q.filter(ProductionBatch.product_id == product_id)
    rows = q.order_by(ProductionBatch.created_at.desc()).all()
    pf = _build_batch_prefetch(rows, db)
    return [_batch_to_out(b, db, pf) for b in rows]


@_full.post("/batches", response_model=BatchOut, status_code=status.HTTP_201_CREATED)
def create_batch(
    payload: BatchCreate,
    bg:   BackgroundTasks,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> BatchOut:
    _get_product(payload.product_id, org, db)
    _get_spec_for_product(payload.specification_id, payload.product_id, org, db)
    if payload.order_id:
        order = db.query(Order).filter_by(id=payload.order_id, organization_id=org.id).first()
        if not order:
            raise HTTPException(status_code=404, detail="Order not found")
    if payload.print_task_id:
        from app.models.task import PrintTask
        task = db.query(PrintTask).filter_by(id=payload.print_task_id, organization_id=org.id).first()
        if not task:
            raise HTTPException(status_code=404, detail="Print task not found")
    b = ProductionBatch(organization_id=org.id, created_by_id=user.id, **payload.model_dump())
    db.add(b)
    # auto-advance linked order to in_production
    if payload.order_id:
        if order and order.status == OrderStatus.confirmed:
            order.status = OrderStatus.in_production
    db.commit()
    db.refresh(b)
    bg.add_task(broadcast_warehouse, org.id, "production")
    return _batch_to_out(b, db)


@_full.get("/batches/{batch_id}", response_model=BatchOut)
def get_batch(
    batch_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> BatchOut:
    b = db.query(ProductionBatch).filter_by(id=batch_id, organization_id=org.id).first()
    if not b:
        raise HTTPException(status_code=404, detail="Batch not found")
    return _batch_to_out(b, db)


@_full.patch("/batches/{batch_id}", response_model=BatchOut)
def update_batch(
    batch_id: int,
    payload:  BatchUpdate,
    bg:  BackgroundTasks,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> BatchOut:
    b = db.query(ProductionBatch).filter_by(id=batch_id, organization_id=org.id).first()
    if not b:
        raise HTTPException(status_code=404, detail="Batch not found")
    for k, v in payload.model_dump(exclude_unset=True).items():
        setattr(b, k, v)
    db.commit()
    db.refresh(b)
    bg.add_task(broadcast_warehouse, org.id, "production")
    return _batch_to_out(b, db)


@_full.delete("/batches/{batch_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_batch(
    batch_id: int,
    bg:  BackgroundTasks,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> None:
    b = db.query(ProductionBatch).filter_by(id=batch_id, organization_id=org.id).first()
    if not b:
        raise HTTPException(status_code=404, detail="Batch not found")
    if b.status == BatchStatus.done:
        raise HTTPException(
            status_code=400,
            detail="Завершену партію не можна видалити, бо вона вже могла створити складські рухи.",
        )
    db.delete(b)
    db.commit()
    bg.add_task(broadcast_warehouse, org.id, "production")


@_full.patch("/batches/{batch_id}/progress", response_model=BatchOut)
def update_progress(
    batch_id:   int,
    bg:   BackgroundTasks,
    printed_qty: int = Query(..., ge=0),
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    _:    User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> BatchOut:
    b = db.query(ProductionBatch).filter_by(id=batch_id, organization_id=org.id).first()
    if not b:
        raise HTTPException(status_code=404, detail="Batch not found")
    if b.status != BatchStatus.active:
        raise HTTPException(status_code=400, detail="Прогрес можна оновлювати лише в активній партії")
    if printed_qty > b.target_qty:
        raise HTTPException(status_code=400, detail=f"Кількість не може перевищувати ціль ({b.target_qty})")
    b.printed_qty = printed_qty
    db.commit()
    db.refresh(b)
    bg.add_task(broadcast_warehouse, org.id, "production")
    return _batch_to_out(b, db)


@_full.post("/batches/{batch_id}/close", response_model=BatchOut)
def close_batch(
    batch_id: int,
    payload:  BatchClose,
    bg:   BackgroundTasks,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> BatchOut:
    b = db.query(ProductionBatch).filter_by(id=batch_id, organization_id=org.id).first()
    if not b:
        raise HTTPException(status_code=404, detail="Batch not found")
    if b.status == BatchStatus.done:
        raise HTTPException(status_code=400, detail="Batch already closed")
    if payload.good_qty < 0 or payload.defect_qty < 0:
        raise HTTPException(status_code=400, detail="Quantities must be non-negative")
    if payload.good_qty > 0 and payload.finished_warehouse_id is None:
        raise HTTPException(status_code=400, detail="finished_warehouse_id is required when good_qty > 0")
    if payload.defect_qty > 0 and payload.defect_warehouse_id is None:
        raise HTTPException(status_code=400, detail="defect_warehouse_id is required when defect_qty > 0")

    if payload.finished_warehouse_id is not None:
        _get_warehouse(payload.finished_warehouse_id, org, db)
    if payload.defect_warehouse_id is not None:
        _get_warehouse(payload.defect_warehouse_id, org, db)

    batch_product = _get_product(b.product_id, org, db)
    output_unit_cost = batch_product.full_cost or batch_product.direct_cost or batch_product.cost_price

    component_picks: list[tuple[SpecComponent, list[tuple[int, Decimal]]]] = []
    if b.specification_id and payload.good_qty > 0:
        spec_comps = db.query(SpecComponent).filter_by(specification_id=b.specification_id).all()
        shortages: list[str] = []
        for c in spec_comps:
            if not c.product_id:
                continue
            component_product = _get_product(c.product_id, org, db)
            qty_needed = c.quantity * Decimal(payload.good_qty)
            remaining = qty_needed
            picks: list[tuple[int, Decimal]] = []
            entries = (
                db.query(StockEntry)
                .filter_by(organization_id=org.id, product_id=c.product_id)
                .with_for_update()
                .all()
            )
            entries.sort(key=lambda e: float(e.quantity - e.reserved_qty), reverse=True)
            for entry in entries:
                available = max(Decimal("0"), entry.quantity - entry.reserved_qty)
                if available <= 0:
                    continue
                take = min(available, remaining)
                picks.append((entry.warehouse_id, take))
                remaining -= take
                if remaining <= 0:
                    break
            if remaining > 0:
                available = qty_needed - remaining
                shortages.append(f"{component_product.name}: потрібно {qty_needed}, доступно {available}")
            else:
                component_picks.append((c, picks))

        if shortages:
            raise HTTPException(
                status_code=422,
                detail="Недостатньо компонентів для закриття партії:\n" + "\n".join(shortages),
            )

    b.good_qty   = payload.good_qty
    b.defect_qty = payload.defect_qty
    b.status     = BatchStatus.done

    if payload.good_qty > 0 and payload.finished_warehouse_id:
        m_in = WarehouseMovement(
            organization_id=org.id, created_by_id=user.id,
            type=MovementType.PRODUCTION_IN,
            product_id=b.product_id,
            warehouse_to_id=payload.finished_warehouse_id,
            quantity=Decimal(payload.good_qty), unit=batch_product.unit,
            unit_cost=output_unit_cost,
            total_cost=(output_unit_cost * Decimal(payload.good_qty)) if output_unit_cost else None,
            batch_id=b.id,
        )
        db.add(m_in)
        db.flush()
        _apply_movement(m_in, db)

    if payload.defect_qty > 0 and payload.defect_warehouse_id:
        m_def = WarehouseMovement(
            organization_id=org.id, created_by_id=user.id,
            type=MovementType.DEFECT,
            product_id=b.product_id,
            warehouse_to_id=payload.defect_warehouse_id,
            quantity=Decimal(payload.defect_qty), unit=batch_product.unit,
            unit_cost=output_unit_cost,
            total_cost=(output_unit_cost * Decimal(payload.defect_qty)) if output_unit_cost else None,
            batch_id=b.id,
        )
        db.add(m_def)
        db.flush()
        _apply_movement(m_def, db)

    # Deduct spec components from stock for every good unit assembled
    for c, picks in component_picks:
        component_product = _get_product(c.product_id, org, db) if c.product_id else None
        unit_cost = component_product.cost_price if component_product else None
        for warehouse_id, quantity in picks:
            m_out = WarehouseMovement(
                organization_id=org.id, created_by_id=user.id,
                type=MovementType.PRODUCTION_OUT,
                product_id=c.product_id,
                warehouse_from_id=warehouse_id,
                quantity=quantity,
                unit=c.unit,
                unit_cost=unit_cost,
                total_cost=(unit_cost * quantity) if unit_cost else None,
                batch_id=b.id,
            )
            db.add(m_out)
            db.flush()
            _apply_movement(m_out, db)
            _check_and_auto_replenish(c.product_id, org.id, db)

    db.commit()
    db.refresh(b)
    bg.add_task(broadcast_warehouse, org.id, "production")
    bg.add_task(broadcast_warehouse, org.id, "stock")
    bg.add_task(broadcast_warehouse, org.id, "movements")
    return _batch_to_out(b, db)


# ── Assembly sessions ─────────────────────────────────────────────────────────

def _session_to_out(
    s: AssemblySession,
    db: Session,
    *,
    worker_map: dict[int, User] | None = None,
    batch_map: dict[int, ProductionBatch] | None = None,
    product_map: dict[int, Product] | None = None,
) -> AssemblySessionOut:
    worker = worker_map.get(s.worker_id) if worker_map is not None else db.get(User, s.worker_id)
    batch = batch_map.get(s.batch_id) if batch_map is not None else db.get(ProductionBatch, s.batch_id)
    product_name = ""
    if batch:
        p = product_map.get(batch.product_id) if product_map is not None else db.get(Product, batch.product_id)
        product_name = p.name if p else ""
    duration: int | None = None
    if s.closed_at:
        delta = s.closed_at - s.started_at
        duration = int(delta.total_seconds() / 60)
    elif s.started_at:
        delta = datetime.utcnow().replace(tzinfo=s.started_at.tzinfo) - s.started_at
        duration = int(delta.total_seconds() / 60)
    return AssemblySessionOut(
        id=s.id, batch_id=s.batch_id, product_name=product_name,
        worker_id=s.worker_id,
        worker_name=(worker.name or worker.email) if worker else "",
        started_at=s.started_at, closed_at=s.closed_at,
        units_good=s.units_good, units_defective=s.units_defective,
        notes=s.notes, duration_minutes=duration,
    )


@_full.post("/batches/{batch_id}/assign", response_model=BatchOut)
def assign_batch(
    batch_id: int,
    payload:  BatchAssignRequest,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> BatchOut:
    """Assign (or unassign) a batch to a worker."""
    b = db.query(ProductionBatch).filter_by(id=batch_id, organization_id=org.id).first()
    if not b:
        raise HTTPException(status_code=404, detail="Партія не знайдена")
    if payload.assigned_to_id is not None:
        worker = db.query(User).filter_by(id=payload.assigned_to_id, organization_id=org.id).first()
        if not worker:
            raise HTTPException(status_code=404, detail="Робітника не знайдено")
    b.assigned_to_id = payload.assigned_to_id
    db.commit()
    db.refresh(b)
    return _batch_to_out(b, db)


@_full.post("/batches/{batch_id}/sessions", response_model=AssemblySessionOut, status_code=201)
def start_session(
    batch_id: int,
    bg:   BackgroundTasks,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> AssemblySessionOut:
    """Open a new assembly session for the current user on this batch."""
    b = db.query(ProductionBatch).filter_by(id=batch_id, organization_id=org.id).first()
    if not b:
        raise HTTPException(status_code=404, detail="Партія не знайдена")
    if b.status == BatchStatus.done:
        raise HTTPException(status_code=400, detail="Партія вже закрита")
    # open the batch if it's still draft
    if b.status == BatchStatus.draft:
        b.status = BatchStatus.active
    # ensure there's no already-open session for this worker on this batch
    existing = db.query(AssemblySession).filter_by(
        batch_id=batch_id, worker_id=user.id, organization_id=org.id
    ).filter(AssemblySession.closed_at.is_(None)).first()
    if existing:
        raise HTTPException(status_code=400, detail="У вас вже є відкрита сесія для цієї партії")
    s = AssemblySession(
        organization_id=org.id,
        batch_id=batch_id,
        worker_id=user.id,
        units_good=0,
        units_defective=0,
    )
    db.add(s)
    db.commit()
    db.refresh(s)
    bg.add_task(broadcast_warehouse, org.id, "production")
    return _session_to_out(s, db)


@_full.patch("/batches/{batch_id}/sessions/{session_id}", response_model=AssemblySessionOut)
def update_session(
    batch_id:   int,
    session_id: int,
    payload:    AssemblySessionUpdate,
    bg:   BackgroundTasks,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> AssemblySessionOut:
    """Update units_good / units_defective during an active session."""
    s = db.query(AssemblySession).filter_by(
        id=session_id, batch_id=batch_id, organization_id=org.id
    ).first()
    if not s:
        raise HTTPException(status_code=404, detail="Сесію не знайдено")
    if s.closed_at:
        raise HTTPException(status_code=400, detail="Сесія вже закрита")
    # workers can only update their own session; admins/managers can update any
    if user.role == UserRole.operator and s.worker_id != user.id:
        raise HTTPException(status_code=403, detail="Це не ваша сесія")
    if payload.units_good is not None:
        s.units_good = payload.units_good
    if payload.units_defective is not None:
        s.units_defective = payload.units_defective
    if payload.notes is not None:
        s.notes = payload.notes
    db.commit()
    db.refresh(s)
    bg.add_task(broadcast_warehouse, org.id, "production")
    return _session_to_out(s, db)


@_full.post("/batches/{batch_id}/sessions/{session_id}/close", response_model=AssemblySessionOut)
def close_session(
    batch_id:   int,
    session_id: int,
    payload:    AssemblySessionUpdate,
    bg:   BackgroundTasks,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> AssemblySessionOut:
    """Close a session and add its output to the batch totals."""
    s = db.query(AssemblySession).filter_by(
        id=session_id, batch_id=batch_id, organization_id=org.id
    ).first()
    if not s:
        raise HTTPException(status_code=404, detail="Сесію не знайдено")
    if s.closed_at:
        raise HTTPException(status_code=400, detail="Сесія вже закрита")
    if user.role == UserRole.operator and s.worker_id != user.id:
        raise HTTPException(status_code=403, detail="Це не ваша сесія")
    # apply final counts
    if payload.units_good is not None:
        s.units_good = payload.units_good
    if payload.units_defective is not None:
        s.units_defective = payload.units_defective
    if payload.notes is not None:
        s.notes = payload.notes
    s.closed_at = datetime.utcnow()
    # roll up into the batch
    b = db.get(ProductionBatch, batch_id)
    if b:
        b.good_qty    += s.units_good
        b.defect_qty  += s.units_defective
        b.printed_qty += s.units_good + s.units_defective
    db.commit()
    db.refresh(s)
    bg.add_task(broadcast_warehouse, org.id, "production")
    return _session_to_out(s, db)


@_full.get("/assembly/sessions", response_model=list[AssemblySessionOut])
def list_sessions(
    batch_id:  int | None = Query(None),
    worker_id: int | None = Query(None),
    date_from: date | None = Query(None),
    date_to:   date | None = Query(None),
    open_only: bool = Query(False),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    user: User        = Depends(require_roles(UserRole.admin, UserRole.operator, UserRole.manager)),
) -> list[AssemblySessionOut]:
    """List sessions. Workers see only their own; managers/admins see all."""
    q = db.query(AssemblySession).filter(AssemblySession.organization_id == org.id)
    if user.role == UserRole.operator:
        q = q.filter(AssemblySession.worker_id == user.id)
    elif worker_id:
        q = q.filter(AssemblySession.worker_id == worker_id)
    if batch_id:
        q = q.filter(AssemblySession.batch_id == batch_id)
    if open_only:
        q = q.filter(AssemblySession.closed_at.is_(None))
    if date_from:
        q = q.filter(func.date(AssemblySession.started_at) >= date_from)
    if date_to:
        q = q.filter(func.date(AssemblySession.started_at) <= date_to)
    rows = q.order_by(AssemblySession.started_at.desc()).limit(200).all()
    worker_ids = {s.worker_id for s in rows}
    batch_ids = {s.batch_id for s in rows}
    worker_map = {
        u.id: u
        for u in db.query(User).filter(User.organization_id == org.id, User.id.in_(worker_ids)).all()
    } if worker_ids else {}
    batch_map = {
        b.id: b
        for b in db.query(ProductionBatch).filter(ProductionBatch.organization_id == org.id, ProductionBatch.id.in_(batch_ids)).all()
    } if batch_ids else {}
    product_ids = {b.product_id for b in batch_map.values()}
    product_map = {
        p.id: p
        for p in db.query(Product).filter(Product.organization_id == org.id, Product.id.in_(product_ids)).all()
    } if product_ids else {}
    return [
        _session_to_out(s, db, worker_map=worker_map, batch_map=batch_map, product_map=product_map)
        for s in rows
    ]


@_full.get("/assembly/stats", response_model=list[WorkerStatsOut])
def assembly_stats(
    date_from: date | None = Query(None),
    date_to:   date | None = Query(None),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.manager)),
) -> list[WorkerStatsOut]:
    """Per-worker output summary. Admin / manager only."""
    q = db.query(AssemblySession).filter(
        AssemblySession.organization_id == org.id,
        AssemblySession.closed_at.isnot(None),
    )
    if date_from:
        q = q.filter(func.date(AssemblySession.started_at) >= date_from)
    if date_to:
        q = q.filter(func.date(AssemblySession.started_at) <= date_to)
    sessions = q.all()

    from collections import defaultdict
    by_worker: dict[int, list[AssemblySession]] = defaultdict(list)
    for s in sessions:
        by_worker[s.worker_id].append(s)

    worker_ids = set(by_worker)
    worker_map = {
        u.id: u
        for u in db.query(User).filter(User.organization_id == org.id, User.id.in_(worker_ids)).all()
    } if worker_ids else {}
    batch_ids = {s.batch_id for s in sessions}
    batch_map = {
        b.id: b
        for b in db.query(ProductionBatch).filter(ProductionBatch.organization_id == org.id, ProductionBatch.id.in_(batch_ids)).all()
    } if batch_ids else {}
    product_ids = {b.product_id for b in batch_map.values()}
    product_map = {
        p.id: p
        for p in db.query(Product).filter(Product.organization_id == org.id, Product.id.in_(product_ids)).all()
    } if product_ids else {}

    result: list[WorkerStatsOut] = []
    for wid, wss in by_worker.items():
        worker = worker_map.get(wid)
        if not worker:
            continue
        total_min = sum(
            int((s.closed_at - s.started_at).total_seconds() / 60)   # type: ignore[operator]
            for s in wss if s.closed_at
        )
        total_good = sum(s.units_good for s in wss)
        total_def  = sum(s.units_defective for s in wss)
        total      = total_good + total_def
        defect_pct = round(total_def / total * 100, 1) if total else 0.0

        # per-product breakdown
        prod_counts: dict[int, int] = defaultdict(int)
        for s in wss:
            b = batch_map.get(s.batch_id)
            if b:
                prod_counts[b.product_id] += s.units_good

        products = []
        for pid, cnt in prod_counts.items():
            p = product_map.get(pid)
            products.append(WorkerProductStats(
                product_id=pid,
                product_name=p.name if p else f"#{pid}",
                units_good=cnt,
            ))
        products.sort(key=lambda x: x.units_good, reverse=True)

        result.append(WorkerStatsOut(
            worker_id=wid,
            worker_name=worker.name or worker.email,
            worker_email=worker.email,
            total_sessions=len(wss),
            total_minutes=total_min,
            total_good=total_good,
            total_defective=total_def,
            defect_rate_pct=defect_pct,
            products=products,
        ))

    result.sort(key=lambda x: x.total_good, reverse=True)
    return result


# ── Orders ────────────────────────────────────────────────────────────────────

@_full.get("/orders", response_model=list[OrderOut])
def list_orders(
    order_status: OrderStatus | None = Query(None),
    counterparty_id: int | None      = Query(None),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[OrderOut]:
    q = db.query(Order).filter(Order.organization_id == org.id)
    if order_status:
        q = q.filter(Order.status == order_status)
    if counterparty_id:
        q = q.filter(Order.counterparty_id == counterparty_id)
    rows = q.order_by(Order.created_at.desc()).all()

    # Bulk-load items, products and counterparties to avoid N+1 per order.
    order_ids = [o.id for o in rows]
    items_by_order: dict[int, list[OrderItem]] = {}
    if order_ids:
        for it in db.query(OrderItem).filter(OrderItem.order_id.in_(order_ids)).order_by(OrderItem.id).all():
            items_by_order.setdefault(it.order_id, []).append(it)

    product_ids = {it.product_id for order_items in items_by_order.values() for it in order_items}
    product_map = {
        p.id: p for p in db.query(Product).filter(Product.organization_id == org.id, Product.id.in_(product_ids)).all()
    } if product_ids else {}

    cp_ids = {o.counterparty_id for o in rows if o.counterparty_id}
    cp_map = {
        c.id: c for c in db.query(Counterparty).filter(Counterparty.organization_id == org.id, Counterparty.id.in_(cp_ids)).all()
    } if cp_ids else {}

    return [
        _order_to_out(o, db, items=items_by_order.get(o.id, []),
                      product_map=product_map, cp_map=cp_map)
        for o in rows
    ]


@_full.post("/orders", response_model=OrderOut, status_code=status.HTTP_201_CREATED)
def create_order(
    payload: OrderCreate,
    bg:   BackgroundTasks,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> OrderOut:
    if payload.counterparty_id:
        _get_counterparty(payload.counterparty_id, org, db)

    order_number = _next_order_number(org, db)
    total = sum(it.unit_price * it.quantity for it in payload.items) if payload.items else None

    o = Order(
        organization_id=org.id,
        created_by_id=user.id,
        order_number=order_number,
        counterparty_id=payload.counterparty_id,
        customer_name=payload.customer_name,
        source=payload.source,
        due_date=payload.due_date,
        currency=payload.currency,
        notes=payload.notes,
        total_amount=total,
    )
    db.add(o)
    db.flush()

    for it in payload.items:
        _get_product(it.product_id, org, db)
        if it.warehouse_id is not None:
            _get_warehouse(it.warehouse_id, org, db)
        db.add(OrderItem(
            order_id=o.id,
            product_id=it.product_id,
            warehouse_id=it.warehouse_id,
            quantity=it.quantity,
            unit_price=it.unit_price,
            total_price=it.unit_price * it.quantity,
        ))

    db.commit()
    db.refresh(o)
    bg.add_task(broadcast_warehouse, org.id, "orders")
    return _order_to_out(o, db)


@_full.get("/orders/{order_id}", response_model=OrderOut)
def get_order(
    order_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> OrderOut:
    o = db.query(Order).filter_by(id=order_id, organization_id=org.id).first()
    if not o:
        raise HTTPException(status_code=404, detail="Order not found")
    return _order_to_out(o, db)


@_full.patch("/orders/{order_id}", response_model=OrderOut)
def update_order(
    order_id: int,
    payload:  OrderUpdate,
    bg:  BackgroundTasks,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> OrderOut:
    o = db.query(Order).filter_by(id=order_id, organization_id=org.id).first()
    if not o:
        raise HTTPException(status_code=404, detail="Order not found")
    if payload.counterparty_id is not None:
        _get_counterparty(payload.counterparty_id, org, db)
    for k, v in payload.model_dump(exclude_unset=True).items():
        setattr(o, k, v)
    db.commit()
    db.refresh(o)
    bg.add_task(broadcast_warehouse, org.id, "orders")
    return _order_to_out(o, db)


@_full.post("/orders/{order_id}/reserve", response_model=OrderOut)
def reserve_order(
    order_id: int,
    payload:  ReserveRequest,
    bg:   BackgroundTasks,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    _:    User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> OrderOut:
    """Lock stock for all items in the order using SELECT FOR UPDATE.

    Transitions order: new → confirmed.
    Sets item.warehouse_id to the chosen warehouse for each item.
    """
    _get_warehouse(payload.warehouse_id, org, db)

    # Lock the order row first to prevent double-reserve races
    o = (
        db.query(Order)
        .filter_by(id=order_id, organization_id=org.id)
        .with_for_update()
        .first()
    )
    if not o:
        raise HTTPException(status_code=404, detail="Order not found")
    if o.status != OrderStatus.new:
        raise HTTPException(status_code=400, detail=f"Cannot reserve order with status '{o.status}'")

    items = db.query(OrderItem).filter_by(order_id=o.id).all()
    if not items:
        raise HTTPException(status_code=400, detail="Order has no items")

    # Check and lock stock for every item before mutating anything
    shortage: list[str] = []
    for item in items:
        wh_id = item.warehouse_id or payload.warehouse_id
        entry = (
            db.query(StockEntry)
            .filter_by(organization_id=org.id, product_id=item.product_id, warehouse_id=wh_id)
            .with_for_update()
            .first()
        )
        available = (entry.quantity - entry.reserved_qty) if entry else Decimal("0")
        if available < item.quantity:
            product = db.get(Product, item.product_id)
            name = product.name if product else f"#{item.product_id}"
            shortage.append(f"{name}: потрібно {item.quantity}, доступно {available:.0f}")

    if shortage:
        raise HTTPException(
            status_code=422,
            detail="Недостатньо товару на складі:\n" + "\n".join(shortage),
        )

    # All checks passed — apply reservations
    for item in items:
        wh_id = item.warehouse_id or payload.warehouse_id

        entry = db.query(StockEntry).filter_by(
            organization_id=org.id,
            product_id=item.product_id,
            warehouse_id=wh_id,
        ).first()
        if entry:
            entry.reserved_qty += item.quantity

        # Record which warehouse this item is reserved from
        item.warehouse_id = wh_id

    o.status = OrderStatus.confirmed
    db.commit()
    db.refresh(o)
    bg.add_task(broadcast_warehouse, org.id, "orders")
    bg.add_task(broadcast_warehouse, org.id, "stock")
    return _order_to_out(o, db)


@_full.post("/orders/{order_id}/ship", response_model=OrderOut)
def ship_order(
    order_id: int,
    bg:   BackgroundTasks,
    payload: ShipRequest | None = None,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> OrderOut:
    """Deduct stock, create SALE_OUT movements, update counterparty balance.

    Transitions order: confirmed | ready → shipped.

    Optional `picks` says which cells the goods were physically pulled from, so
    bin counts match reality. Anything not covered by picks falls back to FIFO.
    """
    picks_by_product: dict[int, list[ShipPick]] = {}
    for pk in (payload.picks if payload else []):
        picks_by_product.setdefault(pk.product_id, []).append(pk)
    o = (
        db.query(Order)
        .filter_by(id=order_id, organization_id=org.id)
        .with_for_update()
        .first()
    )
    if not o:
        raise HTTPException(status_code=404, detail="Order not found")
    if o.status not in (OrderStatus.confirmed, OrderStatus.ready):
        raise HTTPException(status_code=400, detail=f"Cannot ship order with status '{o.status}'")

    items = db.query(OrderItem).filter_by(order_id=o.id).all()

    for item in items:
        if not item.warehouse_id:
            product = db.get(Product, item.product_id)
            name = product.name if product else f"#{item.product_id}"
            raise HTTPException(
                status_code=400,
                detail=f"Item '{name}' has no warehouse assigned — reserve the order first",
            )

        # Lock and deduct
        entry = (
            db.query(StockEntry)
            .filter_by(organization_id=org.id, product_id=item.product_id, warehouse_id=item.warehouse_id)
            .with_for_update()
            .first()
        )
        if not entry or entry.quantity < Decimal(item.quantity):
            product = db.get(Product, item.product_id)
            name = product.name if product else f"#{item.product_id}"
            available = entry.quantity if entry else Decimal("0")
            raise HTTPException(
                status_code=422,
                detail=f"Недостатньо товару для відвантаження {name}: потрібно {item.quantity}, доступно {available}",
            )

        # Create audit movement — store both cost (COGS) and price (revenue) separately
        prod       = db.get(Product, item.product_id)
        unit_cost  = prod.cost_price if prod else None
        unit_price = Decimal(str(item.unit_price)) if item.unit_price else None
        qty        = Decimal(item.quantity)

        m = WarehouseMovement(
            organization_id=org.id,
            created_by_id=user.id,
            type=MovementType.SALE_OUT,
            product_id=item.product_id,
            warehouse_from_id=item.warehouse_id,
            quantity=qty,
            unit=prod.unit if prod else "шт",
            unit_cost=unit_cost,
            total_cost=(unit_cost * qty) if unit_cost else None,
            unit_price=unit_price,
            total_revenue=(unit_price * qty) if unit_price else None,
            order_id=o.id,
        )
        db.add(m)
        db.flush()
        # Honor operator-specified bins first (no FIFO guessing), then let the
        # clamp draw any remainder FIFO so locations stay within the new total.
        for pk in picks_by_product.get(item.product_id, []):
            cell = _get_cell(pk.cell_id, org, db, warehouse_id=item.warehouse_id)
            _pick_from_cell(cell, item.product_id, pk.quantity, org.id, db,
                            movement_id=m.id, created_by_id=user.id)
        _apply_movement(m, db)
        _check_and_auto_replenish(item.product_id, org.id, db)

    # Update counterparty balance (outstanding debt)
    if o.counterparty_id:
        outstanding = (o.total_amount or Decimal("0")) - (o.paid_amount or Decimal("0"))
        if outstanding > 0:
            cp = db.query(Counterparty).with_for_update().filter_by(
                id=o.counterparty_id,
                organization_id=org.id,
            ).first()
            if cp:
                cp.balance += outstanding

    o.status = OrderStatus.shipped
    db.commit()
    db.refresh(o)
    bg.add_task(broadcast_warehouse, org.id, "orders")
    bg.add_task(broadcast_warehouse, org.id, "stock")
    bg.add_task(broadcast_warehouse, org.id, "movements")
    return _order_to_out(o, db)


@_full.post("/orders/{order_id}/cancel", response_model=OrderOut)
def cancel_order(
    order_id: int,
    bg:   BackgroundTasks,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    _:    User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> OrderOut:
    """Cancel an order and release any reserved stock."""
    o = (
        db.query(Order)
        .filter_by(id=order_id, organization_id=org.id)
        .with_for_update()
        .first()
    )
    if not o:
        raise HTTPException(status_code=404, detail="Order not found")
    if o.status == OrderStatus.shipped:
        raise HTTPException(status_code=400, detail="Cannot cancel a shipped order")

    # Release reservations for any status that had stock locked (confirmed → ready)
    _RESERVED_STATUSES = {OrderStatus.confirmed, OrderStatus.in_production, OrderStatus.ready}
    if o.status in _RESERVED_STATUSES:
        items = db.query(OrderItem).filter_by(order_id=o.id).all()
        for item in items:
            if item.warehouse_id:
                entry = (
                    db.query(StockEntry)
                    .filter_by(organization_id=org.id, product_id=item.product_id, warehouse_id=item.warehouse_id)
                    .with_for_update()
                    .first()
                )
                if entry:
                    entry.reserved_qty -= min(entry.reserved_qty, Decimal(item.quantity))

    o.status = OrderStatus.cancelled
    db.commit()
    db.refresh(o)
    bg.add_task(broadcast_warehouse, org.id, "orders")
    bg.add_task(broadcast_warehouse, org.id, "stock")
    return _order_to_out(o, db)


# ── Order Payments ────────────────────────────────────────────────────────────

@_full.post("/orders/{order_id}/payments", response_model=OrderPaymentOut, status_code=status.HTTP_201_CREATED)
def record_payment(
    order_id: int,
    payload:  OrderPaymentCreate,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> OrderPaymentOut:
    o = db.query(Order).filter_by(id=order_id, organization_id=org.id).first()
    if not o:
        raise HTTPException(status_code=404, detail="Order not found")
    if o.status == OrderStatus.cancelled:
        raise HTTPException(status_code=400, detail="Скасоване замовлення не може приймати оплати")
    if payload.amount <= 0:
        raise HTTPException(status_code=400, detail="Сума оплати має бути > 0")

    # No overpayment (allow prepayment on any status, block if already overpaid)
    total = o.total_amount or Decimal("0")
    if total > 0 and (o.paid_amount + payload.amount) > total:
        raise HTTPException(
            status_code=400,
            detail=f"Сума перевищує залишок до оплати ({total - o.paid_amount} ₴)",
        )

    paid_at = payload.paid_at or date.today()

    # Atomic: cashflow → payment → order.paid_amount → cp.balance
    tx = CashTransaction(
        organization_id=org.id,
        created_by_id=user.id,
        type=CashTxType.income,
        category=CashTxCategory.order_payment,
        amount=payload.amount,
        counterparty_id=o.counterparty_id,
        order_id=o.id,
        description=f"Оплата {o.order_number}",
        transaction_date=paid_at,
    )
    db.add(tx)
    db.flush()  # get tx.id before linking

    payment = OrderPayment(
        organization_id=org.id,
        order_id=o.id,
        amount=payload.amount,
        paid_at=paid_at,
        method=payload.method,
        note=payload.note,
        cashflow_id=tx.id,
    )
    db.add(payment)
    db.flush()

    o.paid_amount = o.paid_amount + payload.amount

    # Reduce counterparty debt (guard: only if counterparty exists)
    if o.counterparty_id:
        cp = db.query(Counterparty).with_for_update().filter_by(
            id=o.counterparty_id,
            organization_id=org.id,
        ).first()
        if cp:
            cp.balance -= payload.amount

    db.commit()
    db.refresh(payment)
    return OrderPaymentOut.model_validate(payment)


@_full.get("/orders/{order_id}/payments", response_model=list[OrderPaymentOut])
def list_payments(
    order_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[OrderPaymentOut]:
    o = db.query(Order).filter_by(id=order_id, organization_id=org.id).first()
    if not o:
        raise HTTPException(status_code=404, detail="Order not found")
    rows = (
        db.query(OrderPayment)
        .filter_by(order_id=order_id, organization_id=org.id)
        .order_by(OrderPayment.paid_at.desc(), OrderPayment.id.desc())
        .all()
    )
    return [OrderPaymentOut.model_validate(p) for p in rows]


@_full.delete("/orders/{order_id}/payments/{payment_id}", status_code=status.HTTP_204_NO_CONTENT, response_model=None)
def delete_payment(
    order_id:   int,
    payment_id: int,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    _:    User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> None:
    o = db.query(Order).filter_by(id=order_id, organization_id=org.id).first()
    if not o:
        raise HTTPException(status_code=404, detail="Order not found")
    p = db.query(OrderPayment).filter_by(id=payment_id, order_id=order_id, organization_id=org.id).first()
    if not p:
        raise HTTPException(status_code=404, detail="Payment not found")

    amount = p.amount

    # Delete linked cashflow (guard: tx may have been deleted manually)
    if p.cashflow_id:
        tx = db.query(CashTransaction).filter_by(id=p.cashflow_id, organization_id=org.id).first()
        if tx:
            db.delete(tx)

    db.delete(p)
    db.flush()

    # Recalculate paid_amount from remaining payments
    from sqlalchemy import select, func as safunc
    new_paid = db.execute(
        select(safunc.coalesce(safunc.sum(OrderPayment.amount), Decimal("0")))
        .where(OrderPayment.organization_id == org.id, OrderPayment.order_id == order_id)
    ).scalar_one()
    o.paid_amount = new_paid

    # Restore counterparty debt
    if o.counterparty_id:
        cp = db.query(Counterparty).with_for_update().filter_by(
            id=o.counterparty_id,
            organization_id=org.id,
        ).first()
        if cp:
            cp.balance += amount

    db.commit()


# ── Cash Flow ─────────────────────────────────────────────────────────────────

def _tx_to_out(tx: CashTransaction, db: Session) -> CashTxOut:
    cp_map = {}
    order_map = {}
    bank_map = {}
    if tx.counterparty_id:
        cp = db.get(Counterparty, tx.counterparty_id)
        cp_map = {cp.id: cp.name} if cp else {}
    if tx.order_id:
        order = db.get(Order, tx.order_id)
        order_map = {order.id: order.order_number} if order else {}
    if tx.bank_account_id:
        bank = db.get(BankAccount, tx.bank_account_id)
        bank_map = {bank.id: bank.name} if bank else {}
    return _tx_to_out_from_maps(tx, cp_map=cp_map, order_map=order_map, bank_map=bank_map)


def _tx_to_out_from_maps(
    tx: CashTransaction,
    *,
    cp_map: dict[int, str],
    order_map: dict[int, str],
    bank_map: dict[int, str],
) -> CashTxOut:
    return CashTxOut(
        id=tx.id,
        type=tx.type,
        category=tx.category,
        amount=tx.amount,
        counterparty_id=tx.counterparty_id,
        counterparty_name=cp_map.get(tx.counterparty_id) if tx.counterparty_id else None,
        order_id=tx.order_id,
        order_number=order_map.get(tx.order_id) if tx.order_id else None,
        bank_account_id=tx.bank_account_id,
        bank_account_name=bank_map.get(tx.bank_account_id) if tx.bank_account_id else None,
        description=tx.description,
        transaction_date=tx.transaction_date,
        created_at=tx.created_at,
    )


@_full.get("/cashflow", response_model=list[CashTxOut])
def list_cashflow(
    tx_type:    str | None = Query(None, alias="type"),
    date_from:  date | None = Query(None),
    date_to:    date | None = Query(None),
    counterparty_id: int | None = Query(None),
    limit:      int = Query(200, le=1000),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[CashTxOut]:
    q = db.query(CashTransaction).filter(CashTransaction.organization_id == org.id)
    if tx_type:
        q = q.filter(CashTransaction.type == tx_type)
    if date_from:
        q = q.filter(CashTransaction.transaction_date >= date_from)
    if date_to:
        q = q.filter(CashTransaction.transaction_date <= date_to)
    if counterparty_id:
        q = q.filter(CashTransaction.counterparty_id == counterparty_id)
    rows = q.order_by(CashTransaction.transaction_date.desc(), CashTransaction.id.desc()).limit(limit).all()

    cp_ids = {tx.counterparty_id for tx in rows if tx.counterparty_id}
    order_ids = {tx.order_id for tx in rows if tx.order_id}
    bank_ids = {tx.bank_account_id for tx in rows if tx.bank_account_id}
    cp_map = {
        c.id: c.name
        for c in db.query(Counterparty).filter(Counterparty.organization_id == org.id, Counterparty.id.in_(cp_ids)).all()
    } if cp_ids else {}
    order_map = {
        o.id: o.order_number
        for o in db.query(Order).filter(Order.organization_id == org.id, Order.id.in_(order_ids)).all()
    } if order_ids else {}
    bank_map = {
        b.id: b.name
        for b in db.query(BankAccount).filter(BankAccount.organization_id == org.id, BankAccount.id.in_(bank_ids)).all()
    } if bank_ids else {}
    return [
        _tx_to_out_from_maps(tx, cp_map=cp_map, order_map=order_map, bank_map=bank_map)
        for tx in rows
    ]


@_full.post("/cashflow", response_model=CashTxOut, status_code=status.HTTP_201_CREATED)
def create_cash_tx(
    payload: CashTxCreate,
    bg:   BackgroundTasks,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> CashTxOut:
    if payload.counterparty_id:
        _get_counterparty(payload.counterparty_id, org, db)
    if payload.order_id:
        o = db.query(Order).filter_by(id=payload.order_id, organization_id=org.id).first()
        if not o:
            raise HTTPException(status_code=404, detail="Order not found")

    tx = CashTransaction(
        organization_id=org.id,
        created_by_id=user.id,
        **payload.model_dump(),
    )
    db.add(tx)
    db.commit()
    db.refresh(tx)
    bg.add_task(broadcast_warehouse, org.id, "cashflow")
    return _tx_to_out(tx, db)


@_full.delete("/cashflow/{tx_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_cash_tx(
    tx_id: int,
    bg:  BackgroundTasks,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> None:
    tx = db.query(CashTransaction).filter_by(id=tx_id, organization_id=org.id).first()
    if not tx:
        raise HTTPException(status_code=404, detail="Transaction not found")
    db.delete(tx)
    db.commit()
    bg.add_task(broadcast_warehouse, org.id, "cashflow")


@_full.get("/cashflow/summary", response_model=CashFlowSummary)
def cashflow_summary(
    date_from: date | None = Query(None),
    date_to:   date | None = Query(None),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> CashFlowSummary:
    q = db.query(CashTransaction).filter(CashTransaction.organization_id == org.id)
    if date_from:
        q = q.filter(CashTransaction.transaction_date >= date_from)
    if date_to:
        q = q.filter(CashTransaction.transaction_date <= date_to)

    grouped = (
        q.with_entities(CashTransaction.category, CashTransaction.type, func.sum(CashTransaction.amount))
        .group_by(CashTransaction.category, CashTransaction.type)
        .all()
    )
    total_income = sum((amount or Decimal("0")) for _, tx_type, amount in grouped if tx_type == CashTxType.income)
    total_expense = sum((amount or Decimal("0")) for _, tx_type, amount in grouped if tx_type == CashTxType.expense)

    by_category = [
        {"category": category.value, "type": tx_type.value, "total": float(amount or Decimal("0"))}
        for category, tx_type, amount in sorted(grouped, key=lambda row: row[2] or Decimal("0"), reverse=True)
    ]

    return CashFlowSummary(
        total_income=total_income,
        total_expense=total_expense,
        net=total_income - total_expense,
        by_category=by_category,
    )


# ── Analytics ─────────────────────────────────────────────────────────────────

class TopProduct(BaseModel):
    product_id:   int
    product_name: str
    revenue:      Decimal
    units:        int

class MaterialCost(BaseModel):
    name: str
    cost: Decimal

class CashFlowBucket(BaseModel):
    label:   str
    inflow:  Decimal
    outflow: Decimal

class WarehouseAnalytics(BaseModel):
    revenue:        Decimal
    cogs:           Decimal
    gross_profit:   Decimal
    margin_pct:     Decimal
    units_produced: int
    units_sold:     int
    defect_rate:    Decimal
    top_products:   list[TopProduct]
    material_costs: list[MaterialCost]
    cash_flow:      list[CashFlowBucket]


def _period_range(period: str) -> tuple[date, date]:
    today = date.today()
    if period == "month":
        start = today.replace(day=1)
    elif period == "quarter":
        q_start_month = ((today.month - 1) // 3) * 3 + 1
        start = today.replace(month=q_start_month, day=1)
    else:
        start = today.replace(month=1, day=1)
    return start, today


@_full.get("/analytics", response_model=WarehouseAnalytics)
def get_analytics(
    period: str          = Query("month", pattern="^(month|quarter|year)$"),
    db:     Session      = Depends(get_db),
    org:    Organization = Depends(get_current_org),
) -> WarehouseAnalytics:
    start, end = _period_range(period)
    start_dt = datetime.combine(start, datetime.min.time())
    end_dt = datetime.combine(end + timedelta(days=1), datetime.min.time())

    money_zero = Decimal("0")
    revenue_expr = func.coalesce(WarehouseMovement.total_revenue, WarehouseMovement.total_cost, money_zero)
    cost_expr = func.coalesce(WarehouseMovement.total_cost, money_zero)

    sales_rows = (
        db.query(
            WarehouseMovement.product_id,
            func.sum(WarehouseMovement.quantity).label("units"),
            func.sum(revenue_expr).label("revenue"),
            func.sum(cost_expr).label("cogs"),
        )
        .filter(
            WarehouseMovement.organization_id == org.id,
            WarehouseMovement.type == MovementType.SALE_OUT,
            WarehouseMovement.created_at >= start_dt,
            WarehouseMovement.created_at < end_dt,
        )
        .group_by(WarehouseMovement.product_id)
        .all()
    )

    revenue = sum((row.revenue or money_zero) for row in sales_rows)
    cogs = sum((row.cogs or money_zero) for row in sales_rows)
    gross_profit = revenue - cogs
    margin_pct   = (gross_profit / revenue * 100) if revenue > 0 else Decimal("0")

    produced_good, produced_defect = (
        db.query(
            func.coalesce(func.sum(ProductionBatch.good_qty), 0),
            func.coalesce(func.sum(ProductionBatch.defect_qty), 0),
        )
        .filter(
            ProductionBatch.organization_id == org.id,
            ProductionBatch.status == BatchStatus.done,
            ProductionBatch.updated_at >= start_dt,
            ProductionBatch.updated_at < end_dt,
        )
        .one()
    )
    units_produced = int(produced_good or 0)
    total_printed = int((produced_good or 0) + (produced_defect or 0))
    units_sold = int(sum((row.units or 0) for row in sales_rows))
    defect_rate    = (
        Decimal(produced_defect or 0) / Decimal(total_printed) * 100
        if total_printed > 0 else Decimal("0")
    )

    top_sales = sorted(sales_rows, key=lambda row: row.revenue or money_zero, reverse=True)[:5]
    top_product_ids = {row.product_id for row in top_sales}
    product_names = {
        p.id: p.name
        for p in db.query(Product).filter(Product.organization_id == org.id, Product.id.in_(top_product_ids)).all()
    } if top_product_ids else {}

    top_products = [
        TopProduct(
            product_id=row.product_id,
            product_name=product_names.get(row.product_id, f"#{row.product_id}"),
            revenue=row.revenue or money_zero,
            units=int(row.units or 0),
        )
        for row in top_sales
    ]

    purchase_rows = (
        db.query(
            WarehouseMovement.product_id,
            func.sum(cost_expr).label("cost"),
        )
        .filter(
            WarehouseMovement.organization_id == org.id,
            WarehouseMovement.type == MovementType.PURCHASE_IN,
            WarehouseMovement.created_at >= start_dt,
            WarehouseMovement.created_at < end_dt,
        )
        .group_by(WarehouseMovement.product_id)
        .order_by(func.sum(cost_expr).desc())
        .limit(5)
        .all()
    )
    material_product_ids = {row.product_id for row in purchase_rows}
    material_names = {
        p.id: p.name
        for p in db.query(Product).filter(Product.organization_id == org.id, Product.id.in_(material_product_ids)).all()
    } if material_product_ids else {}
    material_costs = [
        MaterialCost(name=material_names.get(row.product_id, f"#{row.product_id}"), cost=row.cost or money_zero)
        for row in purchase_rows
    ]

    # Explicit cash-flow direction per movement type:
    #   inflow  = money received  (SALE_OUT: use total_revenue)
    #   outflow = money spent     (PURCHASE_IN: cost of goods; PRODUCTION_OUT: component consumption)
    #   ignore  = internal moves  (PRODUCTION_IN, TRANSFER, ADJUSTMENT, DEFECT)
    _CF_DIRECTION: dict[MovementType, str] = {
        MovementType.SALE_OUT:       "inflow",
        MovementType.PURCHASE_IN:    "outflow",
        MovementType.PRODUCTION_OUT: "outflow",
        MovementType.PRODUCTION_IN:  "ignore",
        MovementType.TRANSFER:       "ignore",
        MovementType.ADJUSTMENT:     "ignore",
        MovementType.DEFECT:         "ignore",
    }

    cf_rows = (
        db.query(
            func.date(WarehouseMovement.created_at).label("bucket_date"),
            WarehouseMovement.type,
            func.sum(revenue_expr).label("revenue"),
            func.sum(cost_expr).label("cost"),
        )
        .filter(
            WarehouseMovement.organization_id == org.id,
            WarehouseMovement.type.in_([MovementType.SALE_OUT, MovementType.PURCHASE_IN, MovementType.PRODUCTION_OUT]),
            WarehouseMovement.created_at >= start_dt,
            WarehouseMovement.created_at < end_dt,
        )
        .group_by(func.date(WarehouseMovement.created_at), WarehouseMovement.type)
        .all()
    )

    month_names = ["Січ","Лют","Бер","Квіт","Трав","Черв","Лип","Серп","Вер","Жовт","Лист","Груд"]
    cash_flow: list[CashFlowBucket] = []

    if period == "month":
        week = start
        while week <= end:
            w_end = min(week + timedelta(days=6), end)
            label = f"{week.day}–{w_end.day} {week.strftime('%b')}"
            inflow = outflow = Decimal("0")
            for row in cf_rows:
                bucket_date = row.bucket_date
                if not (week <= bucket_date <= w_end):
                    continue
                direction = _CF_DIRECTION.get(row.type, "ignore")
                if direction == "inflow":
                    inflow += row.revenue or money_zero
                elif direction == "outflow":
                    outflow += row.cost or money_zero
            cash_flow.append(CashFlowBucket(label=label, inflow=inflow, outflow=outflow))
            week = w_end + timedelta(days=1)
    else:
        buckets: dict[str, tuple[Decimal, Decimal]] = {}
        for row in cf_rows:
            direction = _CF_DIRECTION.get(row.type, "ignore")
            if direction == "ignore":
                continue
            k = row.bucket_date.strftime("%Y-%m")
            i, o_val = buckets.get(k, (Decimal("0"), Decimal("0")))
            if direction == "inflow":
                buckets[k] = (i + (row.revenue or money_zero), o_val)
            else:
                buckets[k] = (i, o_val + (row.cost or money_zero))
        for k in sorted(buckets):
            mo = int(k[5:])
            i, o_val = buckets[k]
            cash_flow.append(CashFlowBucket(label=month_names[mo - 1], inflow=i, outflow=o_val))

    return WarehouseAnalytics(
        revenue=revenue, cogs=cogs, gross_profit=gross_profit,
        margin_pct=margin_pct.quantize(Decimal("0.1")),
        units_produced=units_produced, units_sold=units_sold,
        defect_rate=defect_rate.quantize(Decimal("0.1")),
        top_products=top_products,
        material_costs=material_costs,
        cash_flow=cash_flow,
    )


# ── Label Templates ───────────────────────────────────────────────────────────

class LabelTemplateCreate(BaseModel):
    name: str
    item_type: str = "universal"
    width_mm: float = 57.0
    height_mm: float = 32.0
    elements: list = []
    is_default: bool = False

class LabelTemplateUpdate(BaseModel):
    name: str | None = None
    item_type: str | None = None
    width_mm: float | None = None
    height_mm: float | None = None
    elements: list | None = None
    is_default: bool | None = None

class LabelTemplateOut(BaseModel):
    id: int
    name: str
    item_type: str
    width_mm: float
    height_mm: float
    elements: list
    is_default: bool
    is_builtin: bool = False
    model_config = {"from_attributes": True}


@_full.get("/label-templates", response_model=list[LabelTemplateOut])
def list_label_templates(
    item_type: str | None = None,
    org=Depends(get_current_org),
    db: Session = Depends(get_db),
):
    rows = db.query(LabelTemplate).filter(LabelTemplate.organization_id == org.id).all()
    custom = [
        LabelTemplateOut(
            id=r.id, name=r.name, item_type=r.item_type,
            width_mm=float(r.width_mm), height_mm=float(r.height_mm),
            elements=r.elements, is_default=r.is_default, is_builtin=False,
        )
        for r in rows
    ]
    builtins = [
        LabelTemplateOut(**{**t, "is_default": False, "is_builtin": True})
        for t in BUILTIN_TEMPLATES
        if not item_type or t["item_type"] in (item_type, "universal")
    ]
    result = builtins + custom
    if item_type:
        result = [t for t in result if t.item_type in (item_type, "universal")]
    return result


@_full.post("/label-templates", response_model=LabelTemplateOut, status_code=201)
def create_label_template(
    body: LabelTemplateCreate,
    org=Depends(get_current_org),
    db: Session = Depends(get_db),
):
    tpl = LabelTemplate(
        organization_id=org.id,
        name=body.name, item_type=body.item_type,
        width_mm=body.width_mm, height_mm=body.height_mm,
        elements=body.elements, is_default=body.is_default,
    )
    db.add(tpl)
    db.commit()
    db.refresh(tpl)
    return LabelTemplateOut(
        id=tpl.id, name=tpl.name, item_type=tpl.item_type,
        width_mm=float(tpl.width_mm), height_mm=float(tpl.height_mm),
        elements=tpl.elements, is_default=tpl.is_default, is_builtin=False,
    )


@_full.put("/label-templates/{tpl_id}", response_model=LabelTemplateOut)
def update_label_template(
    tpl_id: int,
    body: LabelTemplateUpdate,
    org=Depends(get_current_org),
    db: Session = Depends(get_db),
):
    tpl = db.query(LabelTemplate).filter(
        LabelTemplate.id == tpl_id, LabelTemplate.organization_id == org.id
    ).first()
    if not tpl:
        raise HTTPException(404, "Template not found")
    if body.name is not None:
        tpl.name = body.name
    if body.item_type is not None:
        tpl.item_type = body.item_type
    if body.width_mm is not None:
        tpl.width_mm = body.width_mm
    if body.height_mm is not None:
        tpl.height_mm = body.height_mm
    if body.elements is not None:
        tpl.elements = body.elements
    if body.is_default is not None:
        tpl.is_default = body.is_default
    db.commit()
    db.refresh(tpl)
    return LabelTemplateOut(
        id=tpl.id, name=tpl.name, item_type=tpl.item_type,
        width_mm=float(tpl.width_mm), height_mm=float(tpl.height_mm),
        elements=tpl.elements, is_default=tpl.is_default, is_builtin=False,
    )


@_full.delete("/label-templates/{tpl_id}", status_code=204)
def delete_label_template(
    tpl_id: int,
    org=Depends(get_current_org),
    db: Session = Depends(get_db),
):
    tpl = db.query(LabelTemplate).filter(
        LabelTemplate.id == tpl_id, LabelTemplate.organization_id == org.id
    ).first()
    if not tpl:
        raise HTTPException(404, "Template not found")
    db.delete(tpl)
    db.commit()


# ── Bank Accounts ─────────────────────────────────────────────────────────────

@_full.get("/bank-accounts", response_model=list[BankAccountOut])
def list_bank_accounts(
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[BankAccountOut]:
    rows = (
        db.query(BankAccount)
        .filter_by(organization_id=org.id)
        .order_by(BankAccount.sort_order, BankAccount.id)
        .all()
    )
    return [BankAccountOut.model_validate(r) for r in rows]


@_full.post("/bank-accounts", response_model=BankAccountOut, status_code=status.HTTP_201_CREATED)
def create_bank_account(
    payload: BankAccountCreate,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> BankAccountOut:
    ba = BankAccount(
        organization_id=org.id,
        name=payload.name,
        balance=payload.initial_balance,
        initial_balance=payload.initial_balance,
        initial_balance_date=payload.initial_balance_date,
        sort_order=payload.sort_order,
    )
    db.add(ba)
    db.commit()
    db.refresh(ba)
    return BankAccountOut.model_validate(ba)


@_full.patch("/bank-accounts/{ba_id}", response_model=BankAccountOut)
def update_bank_account(
    ba_id:   int,
    payload: BankAccountUpdate,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> BankAccountOut:
    ba = db.query(BankAccount).filter_by(id=ba_id, organization_id=org.id).first()
    if not ba:
        raise HTTPException(status_code=404, detail="Bank account not found")
    if payload.name is not None:
        ba.name = payload.name
    if payload.status is not None:
        ba.status = BankAccountStatus(payload.status)
    if payload.sort_order is not None:
        ba.sort_order = payload.sort_order
    db.commit()
    db.refresh(ba)
    return BankAccountOut.model_validate(ba)


@_full.delete("/bank-accounts/{ba_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_bank_account(
    ba_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> None:
    ba = db.query(BankAccount).filter_by(id=ba_id, organization_id=org.id).first()
    if not ba:
        raise HTTPException(status_code=404, detail="Bank account not found")
    if abs(ba.balance) > Decimal("0"):
        raise HTTPException(status_code=422, detail="Не можна видалити рахунок з ненульовим балансом")
    db.delete(ba)
    db.commit()


# ── Cash Register (POS) ───────────────────────────────────────────────────────

@_full.get("/cashregister/products")
def cashregister_products(
    q:            str | None = Query(None),
    warehouse_id: int | None = Query(None),
    page:         int        = Query(1, ge=1),
    limit:        int        = Query(30, le=100),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> dict:
    query = (
        db.query(Product)
        .filter(Product.organization_id == org.id, Product.is_active.is_(True))
    )
    if q:
        term = f"%{q.strip()}%"
        query = query.filter(
            or_(
                Product.name.ilike(term),
                Product.sku.ilike(term),
                Product.barcode.ilike(term),
            )
        )
    total = query.count()
    products = query.order_by(Product.name).offset((page - 1) * limit).limit(limit).all()

    # Fetch stock per product for requested warehouse (or sum across all)
    stock_map: dict[int, Decimal] = {}
    if warehouse_id:
        entries = (
            db.query(StockEntry.product_id, StockEntry.quantity)
            .filter(
                StockEntry.organization_id == org.id,
                StockEntry.warehouse_id == warehouse_id,
                StockEntry.product_id.in_([p.id for p in products]),
            )
            .all()
        )
        stock_map = {pid: qty for pid, qty in entries}
    else:
        from sqlalchemy import func as sqlfunc
        entries = (
            db.query(StockEntry.product_id, sqlfunc.sum(StockEntry.quantity))
            .filter(
                StockEntry.organization_id == org.id,
                StockEntry.product_id.in_([p.id for p in products]),
            )
            .group_by(StockEntry.product_id)
            .all()
        )
        stock_map = {pid: qty for pid, qty in entries}

    items = [
        {
            "id":         p.id,
            "name":       p.name,
            "sku":        p.sku,
            "barcode":    p.barcode,
            "sale_price": p.sale_price,
            "unit":       p.unit,
            "image_url":  _product_image_url(p, org.id),
            "stock_qty":  float(stock_map.get(p.id, Decimal("0"))),
        }
        for p in products
    ]
    return {"total": total, "items": items}


@_full.post("/cashregister/sell", response_model=CashRegisterReceipt, status_code=status.HTTP_201_CREATED)
def cashregister_sell(
    payload: CashRegisterSell,
    bg:   BackgroundTasks,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> CashRegisterReceipt:
    if not payload.items:
        raise HTTPException(status_code=422, detail="Кошик порожній")

    wh = db.query(Warehouse).filter_by(id=payload.warehouse_id, organization_id=org.id).first()
    if not wh:
        raise HTTPException(status_code=404, detail="Warehouse not found")

    ba: BankAccount | None = None
    if payload.bank_account_id:
        ba = db.query(BankAccount).filter_by(id=payload.bank_account_id, organization_id=org.id).first()
        if not ba:
            raise HTTPException(status_code=404, detail="Bank account not found")
        if ba.status == BankAccountStatus.closed:
            raise HTTPException(status_code=422, detail="Рахунок закритий")

    receipt_items = []
    total = Decimal("0")

    for item in payload.items:
        product = db.query(Product).filter_by(id=item.product_id, organization_id=org.id).first()
        if not product:
            raise HTTPException(status_code=404, detail=f"Product {item.product_id} not found")

        line_total = (item.qty * item.unit_price).quantize(Decimal("0.01"))
        total += line_total

        movement = WarehouseMovement(
            organization_id=org.id,
            type=MovementType.SALE_OUT,
            product_id=item.product_id,
            warehouse_from_id=payload.warehouse_id,
            quantity=item.qty,
            unit=product.unit,
            unit_price=item.unit_price,
            total_revenue=line_total,
            reason=payload.note or "Каса",
            created_by_id=user.id,
        )
        db.add(movement)
        db.flush()
        _apply_movement(movement, db)
        _check_and_auto_replenish(item.product_id, org.id, db)

        receipt_items.append({
            "product_id":   product.id,
            "product_name": product.name,
            "qty":          item.qty,
            "unit_price":   item.unit_price,
            "total":        line_total,
        })

    # Create cash transaction
    tx = CashTransaction(
        organization_id=org.id,
        type=CashTxType.income,
        category=CashTxCategory.order_payment,
        amount=total,
        bank_account_id=payload.bank_account_id,
        description=payload.note or "Каса",
        transaction_date=date.today(),
        created_by_id=user.id,
    )
    db.add(tx)
    db.flush()

    # Update bank account balance
    if ba is not None:
        ba.balance += total

    created_at = datetime.now()
    db.commit()

    bg.add_task(broadcast_warehouse, org.id, "cashflow")
    bg.add_task(broadcast_warehouse, org.id, "movements")

    return CashRegisterReceipt(
        total=total,
        items=[CashRegisterReceiptItem(**item) for item in receipt_items],
        cash_tx_id=tx.id,
        created_at=created_at,
    )


# ── Purchase Orders ───────────────────────────────────────────────────────────

def _po_to_out(po: PurchaseOrder, product_map: dict, cp_map: dict, wh_map: dict) -> PurchaseOrderOut:
    from decimal import Decimal
    items_out = []
    total = Decimal("0")
    for it in po.items:
        p = product_map.get(it.product_id)
        line_total = it.quantity * it.unit_cost
        total += line_total
        items_out.append(PurchaseOrderItemOut(
            id=it.id,
            product_id=it.product_id,
            product_name=p.name if p else str(it.product_id),
            quantity=it.quantity,
            unit_cost=it.unit_cost,
            total_cost=line_total,
        ))
    cp = cp_map.get(po.counterparty_id) if po.counterparty_id else None
    wh = wh_map.get(po.warehouse_id) if po.warehouse_id else None
    return PurchaseOrderOut(
        id=po.id,
        counterparty_id=po.counterparty_id,
        counterparty_name=cp.name if cp else None,
        warehouse_id=po.warehouse_id,
        warehouse_name=wh.name if wh else None,
        status=po.status.value,
        notes=po.notes,
        total_cost=total,
        items=items_out,
        received_at=po.received_at,
        created_at=po.created_at,
    )


def _po_lookup_maps(po: PurchaseOrder, org_id: int, db: Session) -> tuple[dict[int, Product], dict[int, Counterparty], dict[int, Warehouse]]:
    product_ids = {it.product_id for it in po.items}
    product_map = {
        p.id: p
        for p in db.query(Product).filter(Product.organization_id == org_id, Product.id.in_(product_ids)).all()
    } if product_ids else {}
    cp_map = {}
    if po.counterparty_id:
        cp = db.query(Counterparty).filter_by(id=po.counterparty_id, organization_id=org_id).first()
        cp_map = {cp.id: cp} if cp else {}
    wh_map = {}
    if po.warehouse_id:
        wh = db.query(Warehouse).filter_by(id=po.warehouse_id, organization_id=org_id).first()
        wh_map = {wh.id: wh} if wh else {}
    return product_map, cp_map, wh_map


@_full.get("/purchases", response_model=list[PurchaseOrderOut])
def list_purchases(
    po_status: PurchaseOrderStatus | None = Query(None),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[PurchaseOrderOut]:
    q = db.query(PurchaseOrder).filter(PurchaseOrder.organization_id == org.id)
    if po_status:
        q = q.filter(PurchaseOrder.status == po_status)
    rows = q.order_by(PurchaseOrder.created_at.desc()).all()

    po_ids = [po.id for po in rows]
    items_by_po: dict[int, list[PurchaseOrderItem]] = {}
    if po_ids:
        for it in db.query(PurchaseOrderItem).filter(PurchaseOrderItem.purchase_order_id.in_(po_ids)).all():
            items_by_po.setdefault(it.purchase_order_id, []).append(it)
        for po in rows:
            po.items = items_by_po.get(po.id, [])

    product_ids = {it.product_id for po in rows for it in po.items}
    product_map = {p.id: p for p in db.query(Product).filter(
        Product.organization_id == org.id, Product.id.in_(product_ids)
    ).all()} if product_ids else {}

    cp_ids = {po.counterparty_id for po in rows if po.counterparty_id}
    cp_map = {c.id: c for c in db.query(Counterparty).filter(
        Counterparty.organization_id == org.id, Counterparty.id.in_(cp_ids)
    ).all()} if cp_ids else {}

    wh_ids = {po.warehouse_id for po in rows if po.warehouse_id}
    wh_map = {w.id: w for w in db.query(Warehouse).filter(
        Warehouse.organization_id == org.id, Warehouse.id.in_(wh_ids)
    ).all()} if wh_ids else {}

    return [_po_to_out(po, product_map, cp_map, wh_map) for po in rows]


@_full.post("/purchases", response_model=PurchaseOrderOut, status_code=status.HTTP_201_CREATED)
def create_purchase(
    payload: PurchaseOrderCreate,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> PurchaseOrderOut:
    if payload.counterparty_id:
        cp = db.query(Counterparty).filter_by(organization_id=org.id, id=payload.counterparty_id).first()
        if not cp:
            raise HTTPException(status_code=404, detail="Counterparty not found")
    if payload.warehouse_id:
        wh = db.query(Warehouse).filter_by(organization_id=org.id, id=payload.warehouse_id).first()
        if not wh:
            raise HTTPException(status_code=404, detail="Warehouse not found")

    po = PurchaseOrder(
        organization_id=org.id,
        counterparty_id=payload.counterparty_id,
        warehouse_id=payload.warehouse_id,
        notes=payload.notes,
        status=PurchaseOrderStatus.draft,
        created_by_id=user.id,
    )
    db.add(po)
    db.flush()

    for item in payload.items:
        prod = db.query(Product).filter_by(organization_id=org.id, id=item.product_id).first()
        if not prod:
            raise HTTPException(status_code=404, detail=f"Product {item.product_id} not found")
        db.add(PurchaseOrderItem(
            purchase_order_id=po.id,
            product_id=item.product_id,
            quantity=item.quantity,
            unit_cost=item.unit_cost,
        ))

    db.commit()
    db.refresh(po)

    product_map, cp_map, wh_map = _po_lookup_maps(po, org.id, db)
    return _po_to_out(po, product_map, cp_map, wh_map)


@_full.get("/purchases/{po_id}", response_model=PurchaseOrderOut)
def get_purchase(
    po_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> PurchaseOrderOut:
    po = db.query(PurchaseOrder).filter_by(organization_id=org.id, id=po_id).first()
    if not po:
        raise HTTPException(status_code=404, detail="Purchase order not found")

    product_map, cp_map, wh_map = _po_lookup_maps(po, org.id, db)
    return _po_to_out(po, product_map, cp_map, wh_map)


@_full.post("/purchases/{po_id}/receive", response_model=PurchaseOrderOut)
def receive_purchase(
    po_id: int,
    bg:   BackgroundTasks,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> PurchaseOrderOut:
    po = db.query(PurchaseOrder).filter_by(organization_id=org.id, id=po_id).first()
    if not po:
        raise HTTPException(status_code=404, detail="Purchase order not found")
    if po.status != PurchaseOrderStatus.draft:
        raise HTTPException(status_code=400, detail="Only draft purchase orders can be received")
    if not po.warehouse_id:
        raise HTTPException(status_code=400, detail="Purchase order must have a destination warehouse")
    if not po.items:
        raise HTTPException(status_code=400, detail="Purchase order has no items")

    from datetime import datetime as dt
    for item in po.items:
        mv = WarehouseMovement(
            organization_id=org.id,
            type=MovementType.PURCHASE_IN,
            product_id=item.product_id,
            quantity=item.quantity,
            unit_cost=item.unit_cost,
            warehouse_to_id=po.warehouse_id,
            created_by_id=user.id,
        )
        db.add(mv)
        db.flush()
        _apply_movement(mv, db)

    total_cost = sum(it.quantity * it.unit_cost for it in po.items)
    if po.counterparty_id:
        cp = db.query(Counterparty).filter_by(organization_id=org.id, id=po.counterparty_id).first()
        if cp:
            cp.balance -= total_cost

    po.status = PurchaseOrderStatus.received
    po.received_at = dt.utcnow()
    db.commit()
    db.refresh(po)

    bg.add_task(broadcast_warehouse, org.id, "stock")
    bg.add_task(broadcast_warehouse, org.id, "movements")

    product_map, cp_map, wh_map = _po_lookup_maps(po, org.id, db)
    return _po_to_out(po, product_map, cp_map, wh_map)


@_full.delete("/purchases/{po_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_purchase(
    po_id: int,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    _user: User        = Depends(require_roles(UserRole.admin)),
) -> None:
    po = db.query(PurchaseOrder).filter_by(organization_id=org.id, id=po_id).first()
    if not po:
        raise HTTPException(status_code=404, detail="Purchase order not found")
    if po.status == PurchaseOrderStatus.received:
        raise HTTPException(status_code=400, detail="Cannot delete a received purchase order")
    db.delete(po)
    db.commit()


# ── Order Invoice PDF ─────────────────────────────────────────────────────────

@_full.post("/orders/{order_id}/invoice")
def order_invoice(
    order_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _user: User       = Depends(require_roles(UserRole.admin, UserRole.operator, UserRole.manager)),
):
    from io import BytesIO
    from fastapi.responses import Response
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm

    order = db.query(Order).filter_by(organization_id=org.id, id=order_id).first()
    if not order:
        raise HTTPException(status_code=404, detail="Order not found")

    items = db.query(OrderItem).filter_by(order_id=order.id).all()
    product_ids = {it.product_id for it in items}
    product_map = {p.id: p for p in db.query(Product).filter(
        Product.organization_id == org.id, Product.id.in_(product_ids)
    ).all()}
    cp = db.query(Counterparty).filter_by(organization_id=org.id, id=order.counterparty_id).first() if order.counterparty_id else None

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=20*mm, rightMargin=20*mm,
                            topMargin=20*mm, bottomMargin=20*mm)
    styles = getSampleStyleSheet()
    elems = []

    elems.append(Paragraph(f"Рахунок-фактура №{order.id}", styles["Title"]))
    elems.append(Spacer(1, 6*mm))

    meta = [
        ["Контрагент:", cp.name if cp else "—"],
        ["Статус:", order.status.value],
        ["Дата:", order.created_at.strftime("%d.%m.%Y") if order.created_at else "—"],
    ]
    t = Table(meta, colWidths=[50*mm, 120*mm])
    t.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.grey),
    ]))
    elems.append(t)
    elems.append(Spacer(1, 8*mm))

    header = ["Товар", "Кількість", "Ціна", "Сума"]
    rows = []
    total = 0
    for it in items:
        p = product_map.get(it.product_id)
        name = p.name if p else str(it.product_id)
        price = float(it.unit_price or 0)
        qty = float(it.quantity)
        line = float(it.total_price or qty * price)
        total += line
        rows.append([name, f"{qty:.0f}", f"{price:.2f}", f"{line:.2f}"])
    rows.append(["", "", "Разом:", f"{total:.2f}"])

    tbl = Table([header] + rows, colWidths=[80*mm, 30*mm, 30*mm, 30*mm])
    tbl.setStyle(TableStyle([
        ("BACKGROUND",   (0, 0), (-1, 0), colors.HexColor("#1a1a2e")),
        ("TEXTCOLOR",    (0, 0), (-1, 0), colors.white),
        ("FONTSIZE",     (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#f5f5f5")]),
        ("FONTNAME",     (0, -1), (-1, -1), "Helvetica-Bold"),
        ("LINEBELOW",    (0, -1), (-1, -1), 1, colors.black),
        ("ALIGN",        (1, 0), (-1, -1), "RIGHT"),
    ]))
    elems.append(tbl)

    doc.build(elems)
    buf.seek(0)
    filename = f"invoice_{order.id}.pdf"
    return Response(
        content=buf.read(),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


router.include_router(_full)
