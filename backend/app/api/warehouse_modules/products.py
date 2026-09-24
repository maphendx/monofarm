"""Product catalog, import/export, archive, and image endpoints."""
from datetime import datetime
from decimal import Decimal

import csv
import io
import uuid

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile, status
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.api.deps import get_current_org, require_roles, require_warehouse_full
from app.core.db import get_db
from app.models.gcode_file_output import GcodeFileOutput
from app.models.organization import Organization
from app.models.user import User, UserRole
from app.models.warehouse import (
    CellStock, OrderItem, ProductCategory, ProductImage, ProductionBatch, WarehouseMovement,
    Product,
)
from app.schemas.filament import FilamentOut
from app.schemas.warehouse import (
    ProductCreate, ProductImageOut, ProductOptionOut, ProductOut, ProductUpdate,
    StockThresholdsUpdate,
)

public_router = APIRouter(tags=["warehouse"])

_IMAGE_PREFIX = "product-images"
_IMAGE_MIME_EXT = {"image/jpeg": "jpg", "image/png": "png", "image/webp": "webp"}
_IMAGE_MAX_BYTES = 8 * 1024 * 1024

# All routes that require Starter plan or above (full warehouse access).
# Free plan can only access /products and /categories.
full_router = APIRouter(dependencies=[Depends(require_warehouse_full)])

from app.api.warehouse_modules.common import (_check_product_limit, _get_product, _image_url, _make_product_out)
from app.api.warehouse_modules.pagination import DEFAULT_PAGE_SIZE, PageLimit, PageOffset

# ── Products ──────────────────────────────────────────────────────────────────

@public_router.get("/products", response_model=list[ProductOut])
def list_products(
    search:   str | None = Query(None),
    category: str | None = Query(None),
    archived: bool = Query(False),
    skip: PageOffset = 0,
    limit: PageLimit = DEFAULT_PAGE_SIZE,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[ProductOut]:
    q = db.query(Product).filter(
        Product.organization_id == org.id,
        Product.is_active.is_(not archived),
    )
    if search:
        q = q.filter(Product.name.ilike(f"%{search}%") | Product.sku.ilike(f"%{search}%"))
    if category:
        q = q.filter(Product.categories.contains([category]))
    rows = q.order_by(Product.name, Product.id).offset(skip).limit(limit).all()
    return [_make_product_out(r, org.id) for r in rows]


@public_router.get("/products/options", response_model=list[ProductOptionOut])
def list_product_options(
    search:   str | None = Query(None),
    archived: bool = Query(False),
    skip: PageOffset = 0,
    limit: PageLimit = DEFAULT_PAGE_SIZE,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[ProductOptionOut]:
    """Compact product catalog for selects and comboboxes."""
    q = db.query(Product).filter(
        Product.organization_id == org.id,
        Product.is_active.is_(not archived),
    )
    if search:
        q = q.filter(Product.name.ilike(f"%{search}%") | Product.sku.ilike(f"%{search}%"))
    rows = q.order_by(Product.name, Product.id).offset(skip).limit(limit).all()
    return [ProductOptionOut.model_validate(r) for r in rows]



@public_router.post("/products", response_model=ProductOut, status_code=status.HTTP_201_CREATED)
def create_product(
    payload: ProductCreate,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin)),
) -> ProductOut:
    _check_product_limit(org, db)
    data = payload.model_dump()
    if "cell_limit" in data:
        data["box_limit"] = data.pop("cell_limit")
    p = Product(**data, organization_id=org.id, created_by_id=user.id)
    db.add(p)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"Номенклатура з таким SKU вже існує: {data.get('sku')}")
    db.refresh(p)
    return ProductOut.model_validate(p)


_IMPORT_COL_MAP = {
    "Назва виробу":         "name",
    "Категорії":            "categories",
    "SKU":                  "sku",
    "Штрих-код":            "barcode",
    "Одиниця виміру":       "unit",
    "Середньозважена ціна": "cost_price",
    "Роздрібна ціна":       "sale_price",
    # "Собівартість виробу" intentionally excluded — Monofarm computes this from specs
    "Опис":                 "description",
}

_EXPORT_HEADERS = [
    "Назва виробу", "Категорії", "SKU", "Штрих-код", "Одиниця виміру",
    "Середньозважена ціна", "Роздрібна ціна", "Опис",
]


def _rows_from_upload(file: UploadFile) -> list[dict]:
    raw = file.file.read()
    fname = (file.filename or "").lower()
    if fname.endswith(".xlsx") or fname.endswith(".xls"):
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        headers = [str(c).strip() if c is not None else "" for c in next(rows_iter, [])]
        return [dict(zip(headers, (str(c).strip() if c is not None else "" for c in row))) for row in rows_iter]
    text = raw.decode("utf-8-sig")
    delimiter = "\t" if "\t" in text[:1024] else ","
    return list(csv.DictReader(io.StringIO(text), delimiter=delimiter))


def _sync_categories(names: list[str], org_id: int, db: Session) -> int:
    """Create any ProductCategory rows that don't yet exist for this org. Returns count created."""
    if not names:
        return 0
    existing = {
        r.name for r in db.query(ProductCategory.name).filter(
            ProductCategory.organization_id == org_id,
            ProductCategory.name.in_(names),
        ).all()
    }
    created = 0
    for name in dict.fromkeys(names):  # preserve order, deduplicate
        if name and name not in existing:
            db.add(ProductCategory(organization_id=org_id, name=name))
            created += 1
    return created


def _parse_product_fields(row: dict) -> dict:
    fields: dict = {}
    for col, field in _IMPORT_COL_MAP.items():
        val = (row.get(col) or "").strip()
        if not val:
            continue
        if field == "categories":
            fields[field] = [c.strip() for c in val.split(",") if c.strip()]
        elif field in ("cost_price", "sale_price", "direct_cost"):
            try:
                fields[field] = Decimal(val.replace(",", ".").replace(" ", ""))
            except Exception:
                pass
        else:
            fields[field] = val
    return fields


_COMPARABLE_FIELDS = ["name", "barcode", "unit", "sale_price", "cost_price", "description", "categories"]


@public_router.post("/products/import/preview")
def preview_import(
    file: UploadFile = File(...),
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    _:    User         = Depends(require_roles(UserRole.admin)),
) -> dict:
    try:
        rows = _rows_from_upload(file)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Не вдалося прочитати файл: {exc}") from exc

    db_products = {
        p.sku: p
        for p in db.query(Product).filter(Product.organization_id == org.id, Product.is_active).all()
    }
    file_skus: set[str] = set()
    new_items: list[dict] = []
    existing_items: list[dict] = []

    for row in rows:
        fields = _parse_product_fields(row)
        sku = fields.get("sku", "").strip()
        name = fields.get("name", "").strip()
        if not sku and not name:
            continue
        if sku:
            file_skus.add(sku)

        if sku and sku in db_products:
            p = db_products[sku]
            changes: dict = {}
            for field in _COMPARABLE_FIELDS:
                new_val = fields.get(field)
                if new_val is None:
                    continue
                old_val = getattr(p, field, None)
                if field in ("cost_price", "sale_price", "direct_cost"):
                    try:
                        new_dec = Decimal(str(new_val)).normalize()
                        old_dec = Decimal(str(old_val)).normalize() if old_val is not None else None
                        if new_dec == old_dec:
                            continue
                        new_str = str(new_dec)
                        old_str = str(old_dec) if old_dec is not None else ""
                    except Exception:
                        new_str, old_str = str(new_val), str(old_val) if old_val is not None else ""
                elif isinstance(new_val, list):
                    new_str = ", ".join(new_val)
                    old_str = ", ".join(old_val or [])
                else:
                    new_str = str(new_val).strip()
                    old_str = str(old_val).strip() if old_val is not None else ""
                if new_str != old_str:
                    changes[field] = {"from": old_str, "to": new_str}
            existing_items.append({"id": p.id, "name": p.name, "sku": p.sku, "changes": changes})
        else:
            if not name:
                continue
            item: dict = {"name": name, "sku": sku}
            for f in ("unit", "sale_price", "cost_price"):
                if f in fields:
                    item[f] = str(fields[f])
            new_items.append(item)

    missing_items = [
        {"id": p.id, "name": p.name, "sku": p.sku}
        for sku, p in db_products.items()
        if sku not in file_skus
    ]

    headers = list(rows[0].keys()) if rows else []
    raw_rows = [[str(row.get(h) or "") for h in headers] for row in rows]
    mapping = {str(i): _IMPORT_COL_MAP.get(h) for i, h in enumerate(headers)}

    return {
        "headers": headers,
        "rows": raw_rows,
        "mapping": mapping,
        "summary": {
            "new": len(new_items),
            "existing": len(existing_items),
            "missing": len(missing_items),
        },
        "new": new_items,
        "existing": existing_items,
        "missing": missing_items,
    }


@public_router.post("/products/import")
def import_products(
    file:            UploadFile = File(...),
    action_new:      str = Query("import",  pattern="^(import|skip)$"),
    action_existing: str = Query("update",  pattern="^(update|skip)$"),
    action_missing:  str = Query("nothing", pattern="^(nothing|hide)$"),
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin)),
) -> dict:
    try:
        rows = _rows_from_upload(file)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Не вдалося прочитати файл: {exc}") from exc

    db_products = {
        p.sku: p
        for p in db.query(Product).filter(Product.organization_id == org.id, Product.is_active).all()
    }
    file_skus: set[str] = set()
    all_cats: list[str] = []
    created = updated = skipped = hidden = 0

    for row in rows:
        fields = _parse_product_fields(row)
        sku = fields.get("sku", "").strip()
        name = fields.get("name", "").strip()
        if not sku and not name:
            skipped += 1
            continue
        if sku:
            file_skus.add(sku)
        all_cats.extend(fields.get("categories") or [])

        if sku and sku in db_products:
            if action_existing == "update":
                p = db_products[sku]
                for k, v in fields.items():
                    if k != "sku":
                        setattr(p, k, v)
                p.is_active = True
                updated += 1
            else:
                skipped += 1
        else:
            if action_new == "import":
                if not name:
                    skipped += 1
                    continue
                db.add(Product(organization_id=org.id, created_by_id=user.id, **fields))
                created += 1
            else:
                skipped += 1

    if action_missing == "hide":
        for sku, p in db_products.items():
            if sku not in file_skus:
                p.is_active = False
                hidden += 1

    new_cats = _sync_categories(all_cats, org.id, db)

    db.commit()
    return {"created": created, "updated": updated, "skipped": skipped, "hidden": hidden, "new_categories": new_cats}


def _export_tsv(rows: list, ts: str) -> Response:
    buf = io.StringIO()
    buf.write("﻿")  # UTF-8 BOM for Excel compatibility
    writer = csv.writer(buf, delimiter="\t", lineterminator="\r\n")
    writer.writerow(_EXPORT_HEADERS)
    for p in rows:
        writer.writerow([
            p.name or "",
            ", ".join(p.categories or []),
            p.sku or "",
            p.barcode or "",
            p.unit or "",
            str(p.cost_price) if p.cost_price is not None else "",
            str(p.sale_price) if p.sale_price is not None else "",
            p.description or "",
        ])
    from urllib.parse import quote
    filename = f"номенклатури_{ts}.tsv"
    return Response(
        content=buf.getvalue().encode("utf-8"),
        media_type="text/tab-separated-values; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


def _export_xlsx(rows: list, org_id: int, ts: str) -> Response:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Номенклатури"

    ws.append(list(_EXPORT_HEADERS))

    for p in rows:
        ws.append([
            p.name or "",
            ", ".join(p.categories or []),
            p.sku or "",
            p.barcode or "",
            p.unit or "",
            str(p.cost_price) if p.cost_price is not None else "",
            str(p.sale_price) if p.sale_price is not None else "",
            p.description or "",
        ])

    from urllib.parse import quote
    buf = io.BytesIO()
    wb.save(buf)
    filename = f"номенклатури_{ts}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@public_router.get("/products/export")
def export_products(
    ids: str | None = Query(None, description="Comma-separated product IDs; omit for all"),
    fmt: str        = Query("xlsx"),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> Response:
    import traceback as _tb
    q = db.query(Product).filter(Product.organization_id == org.id, Product.is_active)
    if ids:
        id_list = [int(i) for i in ids.split(",") if i.strip().isdigit()]
        q = q.filter(Product.id.in_(id_list))
    rows = q.order_by(Product.name).all()
    ts = datetime.utcnow().strftime("%Y-%m-%d_%H-%M")
    if fmt == "xlsx":
        try:
            return _export_xlsx(rows, org.id, ts)
        except Exception:
            resp = _export_tsv(rows, ts)
            resp.headers["X-Export-Error"] = _tb.format_exc()[-400:]
            return resp
    return _export_tsv(rows, ts)


@public_router.get("/products/{product_id}", response_model=ProductOut)
def get_product(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> ProductOut:
    return _make_product_out(_get_product(product_id, org, db), org.id)


@public_router.patch("/products/{product_id}", response_model=ProductOut)
def update_product(
    product_id: int,
    payload:    ProductUpdate,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> ProductOut:
    p = _get_product(product_id, org, db)
    data = payload.model_dump(exclude_unset=True)
    if "cell_limit" in data:
        data["box_limit"] = data.pop("cell_limit")
    for k, v in data.items():
        setattr(p, k, v)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"Номенклатура з таким SKU вже існує: {p.sku}")
    db.refresh(p)
    return _make_product_out(p, org.id)


@public_router.post("/products/{product_id}/archive", response_model=ProductOut)
def archive_product(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> ProductOut:
    p = _get_product(product_id, org, db)
    p.is_active = False
    db.commit()
    db.refresh(p)
    return ProductOut.model_validate(p)


@public_router.post("/products/{product_id}/restore", response_model=ProductOut)
def restore_product(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> ProductOut:
    p = _get_product(product_id, org, db)
    p.is_active = True
    db.commit()
    db.refresh(p)
    return _make_product_out(p, org.id)


@public_router.post("/products/{product_id}/copy", response_model=ProductOut, status_code=201)
def copy_product(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> ProductOut:
    p = _get_product(product_id, org, db)
    # Find a unique SKU: base-copy, base-copy-2, base-copy-3 …
    base = p.sku + "-copy"

    taken = {
        row.sku for row in
        db.query(Product.sku).filter(Product.organization_id == org.id, Product.sku.like(base + "%")).all()
    }
    new_sku = base
    n = 2
    while new_sku in taken:
        new_sku = f"{base}-{n}"
        n += 1
    copy = Product(
        organization_id=org.id,
        name=p.name + " (копія)",
        sku=new_sku,
        barcode=None,
        categories=list(p.categories),
        unit=p.unit,
        description=p.description,
        is_active=True,
        sale_price=p.sale_price,
        min_stock=p.min_stock,
        desired_stock=p.desired_stock,
    )
    db.add(copy)
    db.commit()
    db.refresh(copy)
    return _make_product_out(copy, org.id)


# ── Product image ─────────────────────────────────────────────────────────────

@public_router.post("/products/{product_id}/image", response_model=ProductOut)
async def upload_product_image(
    product_id: int,
    file: UploadFile = File(...),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> ProductOut:
    from app.services import storage as storage_svc
    if file.content_type not in _IMAGE_MIME_EXT:
        raise HTTPException(status_code=400, detail="Підтримуються тільки JPEG, PNG, WebP.")
    data = await file.read()
    if len(data) > _IMAGE_MAX_BYTES:
        raise HTTPException(status_code=400, detail="Файл занадто великий (макс. 8 МБ).")
    p = _get_product(product_id, org, db)
    if p.image_key:
        storage_svc.delete(p.image_key, org.id, prefix=_IMAGE_PREFIX)
    ext = _IMAGE_MIME_EXT[file.content_type]
    key = f"{uuid.uuid4().hex}.{ext}"
    storage_svc.put(key, data, org.id, prefix=_IMAGE_PREFIX)
    p.image_key = key
    db.commit()
    db.refresh(p)
    return _make_product_out(p, org.id)


@public_router.delete("/products/{product_id}/image", response_model=ProductOut)
def delete_product_image(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> ProductOut:
    from app.services import storage as storage_svc
    p = _get_product(product_id, org, db)
    if p.image_key:
        storage_svc.delete(p.image_key, org.id, prefix=_IMAGE_PREFIX)
        p.image_key = None
        db.commit()
        db.refresh(p)
    return _make_product_out(p, org.id)


@public_router.get("/products/{product_id}/image")
def serve_product_image(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> Response:
    """Serve image bytes for local-storage mode (S3 mode returns presigned URL instead)."""
    from app.services import storage as storage_svc
    p = _get_product(product_id, org, db)
    if not p.image_key:
        raise HTTPException(status_code=404, detail="No image")
    try:
        data = storage_svc.get_bytes(p.image_key, org.id, prefix=_IMAGE_PREFIX)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="Image file not found")
    ext = p.image_key.rsplit(".", 1)[-1].lower()
    mime = {"jpg": "image/jpeg", "png": "image/png", "webp": "image/webp"}.get(ext, "image/jpeg")
    return Response(content=data, media_type=mime)


# ── Product images (multi-photo) ──────────────────────────────────────────────

@public_router.get("/products/{product_id}/images", response_model=list[ProductImageOut])
def list_product_images(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[ProductImageOut]:
    _get_product(product_id, org, db)  # 404 guard
    rows = (
        db.query(ProductImage)
        .filter_by(product_id=product_id, organization_id=org.id)
        .order_by(ProductImage.is_primary.desc(), ProductImage.sort_order, ProductImage.id)
        .all()
    )
    return [
        ProductImageOut(
            id=img.id,
            image_url=_image_url(img.image_key, product_id, org.id),
            is_primary=img.is_primary,
            sort_order=img.sort_order,
        )
        for img in rows
    ]


@public_router.post("/products/{product_id}/images", response_model=list[ProductImageOut], status_code=201)
async def add_product_image(
    product_id: int,
    file: UploadFile = File(...),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> list[ProductImageOut]:
    from app.services import storage as storage_svc
    if file.content_type not in _IMAGE_MIME_EXT:
        raise HTTPException(status_code=400, detail="Підтримуються тільки JPEG, PNG, WebP.")
    data = await file.read()
    if len(data) > _IMAGE_MAX_BYTES:
        raise HTTPException(status_code=400, detail="Файл занадто великий (макс. 8 МБ).")
    p = _get_product(product_id, org, db)
    ext = _IMAGE_MIME_EXT[file.content_type]
    key = f"{uuid.uuid4().hex}.{ext}"
    storage_svc.put(key, data, org.id, prefix=_IMAGE_PREFIX)
    existing_count = db.query(func.count(ProductImage.id)).filter_by(product_id=product_id, organization_id=org.id).scalar() or 0
    is_primary = existing_count == 0
    img = ProductImage(product_id=product_id, organization_id=org.id, image_key=key, is_primary=is_primary, sort_order=existing_count)
    db.add(img)
    if is_primary:
        p.image_key = key
    db.commit()
    return list_product_images(product_id, db, org)


@public_router.delete("/products/{product_id}/images/{image_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_product_image_by_id(
    product_id: int,
    image_id:   int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> None:
    from app.services import storage as storage_svc
    img = db.query(ProductImage).filter_by(id=image_id, product_id=product_id, organization_id=org.id).first()
    if not img:
        raise HTTPException(status_code=404, detail="Зображення не знайдено")
    was_primary = img.is_primary
    storage_svc.delete(img.image_key, org.id, prefix=_IMAGE_PREFIX)
    db.delete(img)
    db.flush()
    if was_primary:
        next_img = (
            db.query(ProductImage)
            .filter_by(product_id=product_id, organization_id=org.id)
            .order_by(ProductImage.sort_order, ProductImage.id)
            .first()
        )
        p = _get_product(product_id, org, db)
        if next_img:
            next_img.is_primary = True
            p.image_key = next_img.image_key
        else:
            p.image_key = None
    db.commit()


@public_router.patch("/products/{product_id}/images/{image_id}/set-primary", response_model=list[ProductImageOut])
def set_primary_product_image(
    product_id: int,
    image_id:   int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> list[ProductImageOut]:
    img = db.query(ProductImage).filter_by(id=image_id, product_id=product_id, organization_id=org.id).first()
    if not img:
        raise HTTPException(status_code=404, detail="Зображення не знайдено")
    db.query(ProductImage).filter_by(product_id=product_id, organization_id=org.id).update({"is_primary": False})
    img.is_primary = True
    p = _get_product(product_id, org, db)
    p.image_key = img.image_key
    db.commit()
    return list_product_images(product_id, db, org)


@public_router.get("/products/{product_id}/images/{image_key}")
def serve_product_image_by_key(
    product_id: int,
    image_key:  str,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> Response:
    """Serve image bytes for local-storage mode."""
    from app.services import storage as storage_svc
    img = db.query(ProductImage).filter_by(image_key=image_key, product_id=product_id, organization_id=org.id).first()
    if not img:
        raise HTTPException(status_code=404, detail="Image not found")
    try:
        data = storage_svc.get_bytes(image_key, org.id, prefix=_IMAGE_PREFIX)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="Image file not found")
    ext = image_key.rsplit(".", 1)[-1].lower()
    mime = {"jpg": "image/jpeg", "png": "image/png", "webp": "image/webp"}.get(ext, "image/jpeg")
    return Response(content=data, media_type=mime)


@public_router.delete("/products/{product_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_product(
    product_id: int,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin)),
) -> None:
    p = _get_product(product_id, org, db)
    blockers: list[str] = []
    if db.query(func.count(WarehouseMovement.id)).filter_by(product_id=product_id).scalar():
        blockers.append("рухи складу")
    if db.query(func.count(OrderItem.id)).filter_by(product_id=product_id).scalar():
        blockers.append("замовлення")
    if db.query(func.count(ProductionBatch.id)).filter_by(product_id=product_id).scalar():
        blockers.append("виробничі партії")
    if db.query(func.count(GcodeFileOutput.id)).filter_by(product_id=product_id).scalar():
        blockers.append("виробничі прив'язки файлів")
    if db.query(CellStock).filter(CellStock.product_id == product_id, CellStock.quantity > 0).first():
        blockers.append("залишки в комірках")
    if blockers:
        raise HTTPException(
            status_code=409,
            detail=f"Товар задіяний у: {', '.join(blockers)}. Заархівуйте замість видалення.",
        )
    try:
        db.delete(p)
        db.commit()
    except Exception:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Неможливо видалити: товар має звʼязані записи. Заархівуйте замість видалення.",
        )


@public_router.patch("/products/{product_id}/thresholds", response_model=ProductOut)
def update_thresholds(
    product_id: int,
    payload:    StockThresholdsUpdate,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
    _:   User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> ProductOut:
    p = _get_product(product_id, org, db)
    data = payload.model_dump(exclude_unset=True)
    if "cell_limit" in data:
        data["box_limit"] = data.pop("cell_limit")
    for k, v in data.items():
        setattr(p, k, v)
    db.commit()
    db.refresh(p)
    return ProductOut.model_validate(p)


@public_router.get("/products/{product_id}/spools", response_model=list[FilamentOut])
def product_spools(
    product_id: int, skip: PageOffset = 0, limit: PageLimit = DEFAULT_PAGE_SIZE,
    db: Session = Depends(get_db), org: Organization = Depends(get_current_org),
):
    from app.models.filament import Filament
    from app.api.filaments import _to_out
    _get_product(product_id, org, db)
    rows = db.query(Filament).filter_by(organization_id=org.id, warehouse_product_id=product_id).order_by(
        Filament.status, Filament.id,
    ).offset(skip).limit(limit).all()
    return [_to_out(f, db) for f in rows]
