"""Stock views, replenishment, and stock import/export endpoints."""
from datetime import datetime
from decimal import Decimal

import csv
import io

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app.api.deps import get_current_org, require_roles, require_warehouse_full
from app.core.db import get_db
from app.models.organization import Organization
from app.models.user import User, UserRole
from app.models.warehouse import (
    BatchStatus, CellStock, MovementType, ProductionBatch, Specification, StockEntry, Warehouse, WarehouseCell, WarehouseMovement,
    WarehouseType, WarehouseZone, Product,
)
from app.schemas.warehouse import (
    DashboardLowStockOut, DashboardSummaryOut,
    ReplenishPreviewItem, ReplenishRequest,
    StockImportResult,
    StockEntryOut, StockSummaryOut,
)

public_router = APIRouter(tags=["warehouse"])

_LOW_STOCK_THRESHOLD = Decimal("10")

# All routes that require Starter plan or above (full warehouse access).
# Free plan can only access /products and /categories.
full_router = APIRouter(dependencies=[Depends(require_warehouse_full)])

from app.api.warehouse_modules.common import (_apply_movement, _get_product, _get_spec_for_product, _get_warehouse, _product_image_url)
from app.api.warehouse_modules.pagination import DEFAULT_PAGE_SIZE, PageLimit, PageOffset

# ── Stock ─────────────────────────────────────────────────────────────────────

@full_router.get("/dashboard", response_model=DashboardSummaryOut)
def dashboard_summary(
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> DashboardSummaryOut:
    """Lightweight KPI block for the warehouse overview — aggregates only,
    no per-cell joins or presigned image URLs (unlike /stock)."""
    sku_count = db.query(func.count(func.distinct(StockEntry.product_id))).filter(
        StockEntry.organization_id == org.id
    ).scalar() or 0

    total_units = db.query(func.sum(StockEntry.quantity)).filter(
        StockEntry.organization_id == org.id
    ).scalar() or Decimal("0")

    low_rows = (
        db.query(StockEntry, Product)
        .join(Product, Product.id == StockEntry.product_id)
        .join(Warehouse, Warehouse.id == StockEntry.warehouse_id)
        .filter(
            StockEntry.organization_id == org.id,
            Warehouse.type == WarehouseType.finished,
            (StockEntry.quantity - StockEntry.reserved_qty) < _LOW_STOCK_THRESHOLD,
        )
        .order_by(Product.name)
        .all()
    )
    low_stock = [
        DashboardLowStockOut(
            product_id=e.product_id,
            product_name=p.name,
            available=e.quantity - e.reserved_qty,
        )
        for e, p in low_rows
    ]
    return DashboardSummaryOut(sku_count=sku_count, total_units=total_units, low_stock=low_stock)


@full_router.get("/stock", response_model=list[StockEntryOut])
def list_stock(
    warehouse_id: int | None = Query(None),
    product_id:   int | None = Query(None),
    skip: PageOffset = 0,
    limit: PageLimit = DEFAULT_PAGE_SIZE,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[StockEntryOut]:
    q = (
        db.query(StockEntry, Product, Warehouse)
        .join(Product,   Product.id   == StockEntry.product_id)
        .join(Warehouse, Warehouse.id == StockEntry.warehouse_id)
        .filter(StockEntry.organization_id == org.id)
    )
    if warehouse_id:
        q = q.filter(StockEntry.warehouse_id == warehouse_id)
    if product_id:
        q = q.filter(StockEntry.product_id == product_id)
    rows = q.order_by(Product.name, StockEntry.id).offset(skip).limit(limit).all()

    import collections
    from sqlalchemy import func

    product_ids = {p.id for _, p, _ in rows}
    warehouse_ids = {wh.id for _, _, wh in rows}

    total_stock_map = {}
    if product_ids:
        totals = db.query(
            StockEntry.product_id,
            func.sum(StockEntry.quantity).label("total")
        ).filter(
            StockEntry.organization_id == org.id,
            StockEntry.product_id.in_(product_ids)
        ).group_by(StockEntry.product_id).all()
        total_stock_map = {pid: tot for pid, tot in totals}

    cell_stock_map = collections.defaultdict(list)
    if product_ids and warehouse_ids:
        cell_stocks = db.query(
            CellStock.product_id,
            WarehouseZone.warehouse_id,
            WarehouseZone.name.label("zone_name"),
            WarehouseCell.code,
            CellStock.quantity
        ).join(
            WarehouseCell, WarehouseCell.id == CellStock.cell_id
        ).join(
            WarehouseZone, WarehouseZone.id == WarehouseCell.zone_id
        ).filter(
            WarehouseZone.organization_id == org.id,
            CellStock.product_id.in_(product_ids),
            WarehouseZone.warehouse_id.in_(warehouse_ids),
            CellStock.quantity > 0
        ).all()
        for pid, wid, zone_name, cell_code, qty in cell_stocks:
            cell_name = f"{zone_name} {cell_code}"
            cell_stock_map[(pid, wid)].append({"name": cell_name, "quantity": qty})

    # Скільки кожного товару зараз у відкритих виробничих партіях (заплановано/друкується/пауза).
    # Прив'язка партій — на рівні товару, тож значення проставляється всім складським рядкам товару.
    in_production_map: dict[int, int] = {}
    if product_ids:
        open_states = (BatchStatus.draft, BatchStatus.active, BatchStatus.paused)
        prod_rows = db.query(
            ProductionBatch.product_id,
            func.sum(ProductionBatch.target_qty - ProductionBatch.good_qty),
        ).filter(
            ProductionBatch.organization_id == org.id,
            ProductionBatch.product_id.in_(product_ids),
            ProductionBatch.status.in_(open_states),
        ).group_by(ProductionBatch.product_id).all()
        in_production_map = {pid: max(0, int(s or 0)) for pid, s in prod_rows}

    result = []
    for e, p, wh in rows:
        avail = e.quantity - e.reserved_qty
        total_stock = total_stock_map.get(e.product_id, Decimal(0))
        locations = cell_stock_map.get((e.product_id, e.warehouse_id), [])
        assigned = sum((Decimal(str(loc["quantity"])) for loc in locations), Decimal(0))

        result.append(StockEntryOut(
            id=e.id,
            product_id=e.product_id,
            product_name=p.name,
            product_sku=p.sku,
            product_barcode=p.barcode,
            product_categories=p.categories or [],
            image_url=_product_image_url(p, org.id),
            product_unit=p.unit,
            warehouse_id=e.warehouse_id,
            warehouse_name=wh.name,
            locations=locations,
            assigned_qty=assigned,
            unassigned_qty=max(Decimal(0), e.quantity - assigned),
            quantity=e.quantity,
            reserved_qty=e.reserved_qty,
            available=avail,
            total_stock=total_stock,
            full_cost=p.full_cost,
            min_stock=p.min_stock,
            desired_stock=p.desired_stock,
            cell_limit=p.box_limit,
            in_production_qty=in_production_map.get(e.product_id, 0),
            updated_at=e.updated_at,
        ))
    return result


@full_router.get("/stock/summary", response_model=list[StockSummaryOut])
def stock_summary(
    warehouse_id: int | None = Query(None),
    product_id:   int | None = Query(None),
    skip: PageOffset = 0,
    limit: PageLimit = DEFAULT_PAGE_SIZE,
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[StockSummaryOut]:
    """Lightweight stock totals for list pages.

    Use this when the UI only needs quantities. The full /stock endpoint also
    resolves product data, images, cell locations, and production counters.
    """
    q = db.query(StockEntry).filter(StockEntry.organization_id == org.id)
    if warehouse_id:
        q = q.filter(StockEntry.warehouse_id == warehouse_id)
    if product_id:
        q = q.filter(StockEntry.product_id == product_id)

    rows = (
        q.order_by(StockEntry.warehouse_id, StockEntry.product_id)
        .offset(skip)
        .limit(limit)
        .all()
    )
    return [
        StockSummaryOut(
            product_id=e.product_id,
            warehouse_id=e.warehouse_id,
            quantity=e.quantity,
            reserved_qty=e.reserved_qty,
            available=e.quantity - e.reserved_qty,
        )
        for e in rows
    ]


# ── Replenishment ─────────────────────────────────────────────────────────────

@full_router.get("/stock/replenish-preview", response_model=list[ReplenishPreviewItem])
def replenish_preview(
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> list[ReplenishPreviewItem]:
    available_expr = StockEntry.quantity - StockEntry.reserved_qty
    rows = (
        db.query(StockEntry, Product, Warehouse)
        .join(Product,   Product.id   == StockEntry.product_id)
        .join(Warehouse, Warehouse.id == StockEntry.warehouse_id)
        .filter(
            StockEntry.organization_id == org.id,
            or_(
                available_expr <= 0,
                Product.min_stock.isnot(None) & (available_expr < Product.min_stock),
            ),
        )
        .order_by(Product.name)
        .all()
    )
    if not rows:
        return []

    product_ids = list({p.id for _, p, _ in rows})
    default_specs = {
        s.product_id: s
        for s in db.query(Specification)
        .join(Product, Product.id == Specification.product_id)
        .filter(
            Product.organization_id == org.id,
            Specification.product_id.in_(product_ids),
            Specification.is_default.is_(True),
        )
        .all()
    }

    seen: set[int] = set()
    result: list[ReplenishPreviewItem] = []
    for e, p, wh in rows:
        if p.id in seen:
            continue
        seen.add(p.id)
        avail = float(e.quantity - e.reserved_qty)
        target = p.desired_stock if p.desired_stock else (p.min_stock or 1)
        qty_needed = max(1, int(target - avail))
        spec = default_specs.get(p.id)
        result.append(ReplenishPreviewItem(
            product_id=p.id,
            product_name=p.name,
            product_sku=p.sku,
            unit=p.unit,
            available=avail,
            min_stock=p.min_stock,
            desired_stock=p.desired_stock,
            qty_needed=qty_needed,
            kind="batch",
            specification_id=spec.id if spec else None,
            warehouse_id=wh.id,
            warehouse_name=wh.name,
        ))
    return result


@full_router.post("/stock/replenish")
def replenish_stock(
    payload: ReplenishRequest,
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> dict:
    batches_created = 0
    movements_created = 0
    for item in payload.items:
        if item.qty <= 0:
            continue
        product = _get_product(item.product_id, org, db)
        if item.kind == "batch":
            _get_spec_for_product(item.specification_id, item.product_id, org, db)
            b = ProductionBatch(
                organization_id=org.id,
                product_id=item.product_id,
                specification_id=item.specification_id,
                target_qty=item.qty,
                status=BatchStatus.draft,
                created_by_id=user.id,
            )
            db.add(b)
            batches_created += 1
        elif item.kind == "purchase":
            if item.warehouse_id is None:
                raise HTTPException(status_code=400, detail="Purchase replenishment requires warehouse_id")
            _get_warehouse(item.warehouse_id, org, db)
            qty = Decimal(str(item.qty))
            uc  = item.unit_cost
            m = WarehouseMovement(
                organization_id=org.id,
                product_id=item.product_id,
                type=MovementType.PURCHASE_IN,
                warehouse_to_id=item.warehouse_id,
                quantity=qty,
                unit=product.unit,
                unit_cost=uc,
                total_cost=(uc * qty) if uc else None,
                created_by_id=user.id,
            )
            db.add(m)
            db.flush()
            _apply_movement(m, db)
            movements_created += 1
        else:
            raise HTTPException(status_code=400, detail="Unknown replenishment kind")
    db.commit()
    return {"batches": batches_created, "movements": movements_created}


# ── Stock export / import ──────────────────────────────────────────────────────

_STOCK_HEADERS = ("SKU", "Назва", "Склад", "В наявності", "Одиниця")


@full_router.get("/stock/export")
def export_stock(
    fmt: str        = Query("xlsx"),
    db:  Session      = Depends(get_db),
    org: Organization = Depends(get_current_org),
) -> Response:
    from urllib.parse import quote
    rows = (
        db.query(StockEntry, Product, Warehouse)
        .join(Product,   StockEntry.product_id   == Product.id)
        .join(Warehouse, StockEntry.warehouse_id  == Warehouse.id)
        .filter(StockEntry.organization_id == org.id)
        .order_by(Product.name)
        .all()
    )
    ts = datetime.utcnow().strftime("%Y-%m-%d_%H-%M")

    if fmt == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Залишки"
        ws.append(list(_STOCK_HEADERS))
        for se, p, wh in rows:
            ws.append([p.sku, p.name, wh.name, float(se.quantity), p.unit or "шт"])
        buf = io.BytesIO()
        wb.save(buf)
        filename = f"залишки_{ts}.xlsx"
        return Response(
            content=buf.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
        )

    buf = io.StringIO()
    buf.write("﻿")
    writer = csv.writer(buf, delimiter="\t", lineterminator="\r\n")
    writer.writerow(_STOCK_HEADERS)
    for se, p, wh in rows:
        writer.writerow([p.sku, p.name, wh.name, float(se.quantity), p.unit or "шт"])
    filename = f"залишки_{ts}.tsv"
    return Response(
        content=buf.getvalue().encode("utf-8"),
        media_type="text/tab-separated-values; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@full_router.post("/stock/import", response_model=StockImportResult)
def import_stock(
    file: UploadFile   = File(...),
    db:   Session      = Depends(get_db),
    org:  Organization = Depends(get_current_org),
    user: User         = Depends(require_roles(UserRole.admin, UserRole.operator)),
) -> StockImportResult:
    """Stocktake import — reads SKU / Warehouse / Quantity and adjusts stock to match."""
    content = file.file.read()
    try:
        if file.filename and file.filename.endswith(".xlsx"):
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            ws = wb.active
            raw_rows = [
                [str(c.value).strip() if c.value is not None else "" for c in row]
                for row in ws.iter_rows(min_row=2)
            ]
        else:
            text = content.decode("utf-8-sig").replace("\r\n", "\n")
            delimiter = "\t" if "\t" in text else ","
            raw_rows = [r for r in csv.reader(text.splitlines(), delimiter=delimiter)][1:]
    except Exception as e:
        raise HTTPException(400, detail=f"Помилка читання файлу: {e}")

    wh_cache:  dict[str, Warehouse] = {}
    sku_cache: dict[str, Product]   = {}
    updated = skipped = 0
    errors: list[str] = []

    for i, row in enumerate(raw_rows, start=2):
        if not any(row):
            continue
        if len(row) < 4:
            errors.append(f"Рядок {i}: замало колонок (потрібно SKU, Назва, Склад, Кількість)")
            skipped += 1
            continue

        sku, wh_name, qty_raw = row[0].strip(), row[2].strip(), row[3].strip()

        # resolve SKU
        if sku not in sku_cache:
            p = db.query(Product).filter_by(sku=sku, organization_id=org.id).first()
            if not p:
                errors.append(f"Рядок {i}: SKU «{sku}» не знайдено")
                skipped += 1
                continue
            sku_cache[sku] = p
        product = sku_cache[sku]

        # resolve warehouse
        if wh_name not in wh_cache:
            wh = db.query(Warehouse).filter(
                Warehouse.organization_id == org.id,
                func.lower(Warehouse.name) == wh_name.lower(),
            ).first()
            if not wh:
                errors.append(f"Рядок {i}: склад «{wh_name}» не знайдено")
                skipped += 1
                continue
            wh_cache[wh_name] = wh
        warehouse = wh_cache[wh_name]

        try:
            target = Decimal(qty_raw.replace(",", "."))
        except Exception:
            errors.append(f"Рядок {i}: невірна кількість «{qty_raw}»")
            skipped += 1
            continue

        entry = db.query(StockEntry).filter_by(
            organization_id=org.id,
            product_id=product.id,
            warehouse_id=warehouse.id,
        ).first()
        current = entry.quantity if entry else Decimal("0")
        delta = target - current
        if delta == 0:
            skipped += 1
            continue

        if delta > 0:
            m = WarehouseMovement(
                organization_id=org.id, type=MovementType.ADJUSTMENT,
                product_id=product.id, warehouse_to_id=warehouse.id,
                quantity=delta, reason="Імпорт залишків",
                created_by_id=user.id,
            )
        else:
            m = WarehouseMovement(
                organization_id=org.id, type=MovementType.ADJUSTMENT,
                product_id=product.id, warehouse_from_id=warehouse.id,
                quantity=-delta, reason="Імпорт залишків",
                created_by_id=user.id,
            )
        db.add(m)
        db.flush()
        _apply_movement(m, db)
        updated += 1

    db.commit()
    return StockImportResult(updated=updated, skipped=skipped, errors=errors[:50])
